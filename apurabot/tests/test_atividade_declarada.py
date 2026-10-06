"""A atividade declarada na linha — a exceção do caso, não da regra.

Matéria-prima comprada por transmissão de propriedade com a mercadoria já no
armazém volta por retorno de armazenagem, CFOP que o mapa de MS põe no
comercial. O time fiscal marca as linhas dessa operação na BASE TRATADA, com
aprovação, e a atividade delas passa a ser a declarada.

Roda sobre o Livro sintético de Rio Brilhante, sem dado real.
"""
from __future__ import annotations

from types import SimpleNamespace

import openpyxl
import pytest

from apurabot.apuracao import apurar, ler_ajustes
from apurabot.base_tratada import tratar
from apurabot.nucleo import atividade as ativ
from apurabot.saida import escrever
from sintetico import RB
from sintetico import livro as _livro

APROVADA = {
    "atividade_ajustada": "industrial",
    "atividade_motivo": "MP comprada por transmissão de propriedade em armazém",
    "atividade_responsavel": "Fulano",
    "atividade_aprovador": "Ciclano",
}


@pytest.fixture(scope="module")
def mapa(parametros):
    return ativ.mapa_da_uf("MS", parametros)


def _retorno(**dados):
    """Um retorno de armazenagem (2906) — comercial pelo CFOP."""
    return SimpleNamespace(cfop_int=2906,
                           dados={"entrada_saida": "Entrada", **dados})


def test_sem_declaracao_vale_o_cfop(mapa):
    assert ativ.classificar(_retorno(), mapa).atividade == ativ.COMERCIAL


def test_a_declarada_vence_o_cfop(mapa):
    resultado = ativ.classificar(_retorno(**APROVADA), mapa)
    assert resultado.atividade == ativ.INDUSTRIAL
    assert "transmissão de propriedade" in resultado.regra


@pytest.mark.parametrize("falta", [
    "atividade_motivo", "atividade_responsavel", "atividade_aprovador"])
def test_sem_aprovacao_nao_vale(mapa, falta):
    """Não cai no CFOP calada: sai SEM REGRA e diz o que falta."""
    resultado = ativ.classificar(_retorno(**{**APROVADA, falta: None}), mapa)
    assert resultado.atividade == ativ.SEM_REGRA
    assert "preencha para valer" in resultado.regra


def test_atividade_que_nao_existe_nao_vale(mapa):
    resultado = ativ.classificar(
        _retorno(**{**APROVADA, "atividade_ajustada": "fábrica"}), mapa)
    assert resultado.atividade == ativ.SEM_REGRA
    assert "Produção" in resultado.regra


@pytest.mark.parametrize("escrito, esperado", [
    (" Industrial ", ativ.INDUSTRIAL),
    ("Produção", ativ.INDUSTRIAL),
    ("producao", ativ.INDUSTRIAL),
    ("Comercial", ativ.COMERCIAL),
    ("Prestacional / Outras", ativ.PRESTACIONAL),
    ("prestacional_outras", ativ.PRESTACIONAL),
])
def test_vale_o_nome_que_a_planilha_mostra(mapa, escrito, esperado):
    """A conferência escreve "Produção"; o parâmetro, "industrial". Os dois valem."""
    resultado = ativ.classificar(
        _retorno(**{**APROVADA, "atividade_ajustada": escrito}), mapa)
    assert resultado.atividade == esperado


def test_a_marca_volta_no_arquivo_e_muda_a_apuracao(parametros, tmp_path):
    """A venda do Livro sintético marcada como comercial: o benefício some."""
    base = tratar(_livro(tmp_path / "livro.xlsx"), parametros=parametros)
    antes = apurar(base, parametros)
    assert antes.filiais[RB].credito_presumido > 0

    wb = openpyxl.load_workbook(escrever(base, tmp_path / "saida.xlsx", antes))
    aba = wb["BASE TRATADA"]
    coluna = {c.value: c.column for c in aba[1]}
    venda = next(linha[0].row for linha in aba.iter_rows(min_row=2)
                 if linha[coluna["cfop"] - 1].value == 5101)
    for campo, valor in {**APROVADA, "atividade_ajustada": "comercial"}.items():
        aba.cell(row=venda, column=coluna[campo], value=valor)
    devolvido = tmp_path / "devolvido.xlsx"
    wb.save(devolvido)

    base = tratar(devolvido, parametros=parametros)
    depois = apurar(base, parametros, ajustes=ler_ajustes(devolvido))
    filial = depois.filiais[RB]
    assert filial.atividade(ativ.INDUSTRIAL).debito == 0
    assert filial.atividade(ativ.COMERCIAL).debito == pytest.approx(17_000.00)
    assert filial.credito_presumido == 0

    # E a marca sobrevive a mais uma volta.
    wb = openpyxl.load_workbook(escrever(base, tmp_path / "de_novo.xlsx", depois))
    aba = wb["BASE TRATADA"]
    assert aba.cell(row=venda, column=coluna["atividade_ajustada"]).value == "comercial"
    assert aba.cell(row=1, column=coluna["atividade_ajustada"]).fill.fgColor.rgb \
        == aba.cell(row=1, column=coluna["ajuste_linha"]).fill.fgColor.rgb
