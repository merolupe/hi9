"""AJUSTES NO SANKHYA e REGISTRO 1200 — o que se digita no Sankhya e na EFD.

A maior parte dos testes roda sobre um Livro Fiscal sintético de Rio Brilhante,
sem dado real: uma venda que gera débito e benefício, e a NF-e de recebimento
do crédito transferido (CFOP 1601). O último confere julho/2026 contra o
Registro do ERP e é pulado quando o arquivo real não está na máquina.
"""
from __future__ import annotations

import dataclasses
from types import SimpleNamespace

import openpyxl
import pytest

from apurabot import ajustes as aj
from apurabot import controle_de_creditos as cc
from apurabot import lancamentos as lanc
from apurabot.apuracao import apurar, ler_ajustes
from apurabot.base_tratada import tratar
from apurabot.saida import ORDEM_DAS_ABAS, escrever
from sintetico import CHAVE, RB, RECEBIDO
from sintetico import livro as _livro

CENTAVO = 0.005
ABERTURA_DE_AGOSTO = 108_426.33         # parametros/saldos.yaml
UTILIZADO = 108_858.68                  # Ajuste de Apuração do Sankhya, 08/2026


def _linhas(aba) -> list[list]:
    return [[c for c in linha if c is not None]
            for linha in aba.iter_rows(values_only=True)
            if any(c is not None for c in linha)]


@pytest.fixture(scope="module")
def agosto(parametros, tmp_path_factory):
    base = tratar(_livro(tmp_path_factory.mktemp("livro") / "livro.xlsx"),
                  parametros=parametros)
    return base, apurar(base, parametros)


def _devolver(base, apuracao, pasta, valor=UTILIZADO, codigo="87"):
    """A planilha devolvida com o uso do crédito declarado na aba AJUSTES MANUAIS."""
    destino = escrever(base, pasta / "saida.xlsx", apuracao)
    wb = openpyxl.load_workbook(destino)
    aba = wb[aj.ABA]
    titulo = next(celulas[0].row for celulas in aba.iter_rows(max_col=1)
                  if str(celulas[0].value or "").upper() == aj.TITULO_PARCELAS)
    for coluna, v in enumerate(
        [RB, "Produção", "006", valor, "crédito do art. 68 do RICMS/MS",
         "Fulano", "Ciclano", codigo], start=1,
    ):
        aba.cell(row=titulo + 2, column=coluna, value=v)
    devolvido = pasta / "devolvido.xlsx"
    wb.save(devolvido)
    return devolvido


@pytest.fixture(scope="module")
def reapurado(agosto, parametros, tmp_path_factory):
    base, apuracao = agosto
    devolvido = _devolver(base, apuracao, tmp_path_factory.mktemp("volta"))
    base = tratar(devolvido, parametros=parametros)
    return base, apurar(base, parametros, ajustes=ler_ajustes(devolvido))


# -- AJUSTES NO SANKHYA -----------------------------------------------------

def test_ajustes_no_sankhya_fica_por_ultimo(agosto, tmp_path):
    """Só se usa depois de fechados os ajustes manuais: é a última visível."""
    base, apuracao = agosto
    planilha = openpyxl.load_workbook(escrever(base, tmp_path / "x.xlsx", apuracao))
    nomes = planilha.sheetnames
    visiveis = [a.title for a in planilha.worksheets if a.sheet_state == "visible"]
    assert visiveis[-1] == "AJUSTES NO SANKHYA"
    assert [n for n in ORDEM_DAS_ABAS if n in nomes] == nomes


def test_o_calculado_entra_sem_ninguem_declarar(agosto, parametros):
    """O benefício é conta da apuração: vira lançamento com o código 37."""
    base, apuracao = agosto
    roteiro = lanc.montar(apuracao, parametros)
    beneficio = [x for x in roteiro.lancamentos if x.parcela == "beneficio"]
    assert len(beneficio) == 1
    assert beneficio[0].tipo == "Deduções do imposto apurado"
    assert beneficio[0].observacao_padrao == "37"
    assert "1190/2018" in beneficio[0].texto
    assert beneficio[0].valor == pytest.approx(
        apuracao.filiais[RB].credito_presumido, abs=CENTAVO)


def test_o_declarado_entra_com_o_codigo_que_veio_na_aba(reapurado, parametros):
    _, apuracao = reapurado
    roteiro = lanc.montar(apuracao, parametros)
    declarado = [x for x in roteiro.lancamentos if x.parcela == "declarado"]
    assert len(declarado) == 1
    x = declarado[0]
    assert (x.linha, x.tipo, x.observacao_padrao) == (6, "Outros créditos", "87")
    assert x.valor == pytest.approx(UTILIZADO, abs=CENTAVO)
    assert "ART. 68" in x.texto
    assert x.complemento == "crédito do art. 68 do RICMS/MS"


def test_os_lancamentos_fecham_com_o_registro(reapurado, parametros):
    _, apuracao = reapurado
    roteiro = lanc.montar(apuracao, parametros)
    assert roteiro.conferencia
    assert roteiro.divergencias == []


def test_o_codigo_da_linha_volta_na_planilha(reapurado, tmp_path):
    """Quem reenvia a planilha não perde o código que digitou."""
    base, apuracao = reapurado
    planilha = openpyxl.load_workbook(escrever(base, tmp_path / "y.xlsx", apuracao))
    parcelas = [linha for linha in planilha[aj.ABA].iter_rows(values_only=True)
                if linha[0] == RB and linha[2] == "006"]
    assert len(parcelas) == 1
    assert str(parcelas[0][7]) == "87"


@pytest.mark.parametrize(
    "parcela, regime, uf, beneficio, esperado",
    [
        ("estorno", "sp_equilibrio_fiscal", "SP", "", "83"),
        ("estorno", "ms_beneficio_rio_brilhante", "MS", "ms_rio_brilhante", "62"),
        ("estorno", "ms_estorno_proporcional", "MS", "", "61"),
        ("estorno", "mt_diferimento", "MT", "", "29"),
        ("credito_indevido", "ms_estorno_proporcional", "MS", "", "34"),
        ("difal", "sp_equilibrio_fiscal", "SP", "", "33"),
        ("centralizacao_recebe_debito", "sp_equilibrio_fiscal", "SP", "", "82"),
        ("centralizacao_transfere_debito", "ms_estorno_proporcional", "MS", "", "79"),
        ("beneficio", "ms_beneficio_rio_brilhante", "MS", "ms_rio_brilhante", "37"),
        # Sem regra cadastrada: sai sem código, não com um parecido.
        ("estorno", "pr_diferimento", "PR", "", ""),
        ("difal", "ms_estorno_proporcional", "MS", "", ""),
        ("credito_indevido", "sp_equilibrio_fiscal", "SP", "", ""),
    ],
)
def test_o_codigo_de_cada_parcela_vem_do_parametro(
    parametros, parcela, regime, uf, beneficio, esperado
):
    """Os códigos do relatório de Ajuste de Apuração do Sankhya de 08/2026."""
    filial = SimpleNamespace(regime=regime, uf=uf)
    assert lanc._codigo_da_parcela(
        parametros, parcela, filial, beneficio, "2026-08") == esperado


def test_o_codigo_tem_vigencia(parametros):
    """Os códigos foram vistos a partir de 08/2026: julho não os herda."""
    filial = SimpleNamespace(regime="sp_equilibrio_fiscal", uf="SP")
    assert lanc._codigo_da_parcela(parametros, "estorno", filial, "", "2026-07") == ""


# -- REGISTRO 1200 ----------------------------------------------------------

def test_o_recebido_vem_da_nota_do_livro(agosto, parametros):
    _, apuracao = agosto
    (controle,) = cc.montar(apuracao, parametros)
    assert (controle.estabelecimento, controle.cod_aj_apur) == (RB, "MS090004")
    assert controle.recebido == pytest.approx(RECEBIDO)
    assert [n.chave for n in controle.notas] == [CHAVE]
    assert controle.saldo_inicial == pytest.approx(ABERTURA_DE_AGOSTO)


def test_sem_uso_declarado_o_estoque_so_cresce(agosto, parametros):
    """Quanto usar é decisão do fiscal: sem declaração, nada é utilizado."""
    _, apuracao = agosto
    (controle,) = cc.montar(apuracao, parametros)
    assert controle.utilizado == 0
    assert controle.saldo_final == pytest.approx(ABERTURA_DE_AGOSTO + RECEBIDO)
    assert controle.linha_1210() is None


def test_agosto_fecha_como_o_resumo_do_sped(reapurado, parametros):
    """108.426,33 + 62.720,00 − 108.858,68 = 62.287,65 — a tela do 1200 de 08/2026."""
    _, apuracao = reapurado
    (controle,) = cc.montar(apuracao, parametros)
    assert controle.utilizado == pytest.approx(UTILIZADO)
    assert controle.saldo_final == pytest.approx(62_287.65, abs=CENTAVO)
    # O Livro sintético tem saldo devedor de 17.000,00: o uso real de agosto
    # passa do teto e do benefício dele, e só isso fica pendente.
    assert [p for p in controle.pendencias
            if "teto" not in p and "benefício" not in p] == []
    assert controle.linha_1200() == (
        "|1200|MS090004|108426,33|0|62720|108858,68|62287,65|")
    assert controle.linha_1210() == f"|1210|MS03||108858,68|{CHAVE}|"


def test_uso_acima_do_estoque_fica_pendente(agosto, parametros, tmp_path):
    base, apuracao = agosto
    devolvido = _devolver(base, apuracao, tmp_path, valor=200_000.00)
    base = tratar(devolvido, parametros=parametros)
    apuracao = apurar(base, parametros, ajustes=ler_ajustes(devolvido))
    (controle,) = cc.montar(apuracao, parametros)
    assert controle.saldo_final < 0
    assert any("negativo" in p for p in controle.pendencias)


def test_competencia_sem_saldo_declarado_fica_pendente(parametros, tmp_path):
    """Outubro ainda não tem abertura em saldos.yaml: não é zero, é aberto."""
    base = tratar(_livro(tmp_path / "outubro.xlsx", mes=10), parametros=parametros)
    (controle,) = cc.montar(apurar(base, parametros), parametros)
    assert controle.saldo_inicial is None
    assert any("não declarado" in p for p in controle.pendencias)


def test_a_aba_traz_as_linhas_da_efd(reapurado, tmp_path):
    base, apuracao = reapurado
    planilha = openpyxl.load_workbook(escrever(base, tmp_path / "z.xlsx", apuracao))
    texto = [c for linha in _linhas(planilha["REGISTRO 1200"]) for c in linha]
    assert "|1200|MS090004|108426,33|0|62720|108858,68|62287,65|" in texto
    assert f"|1210|MS03||108858,68|{CHAVE}|" in texto


def _sem(params, **trocas):
    """Os parâmetros com blocos de `saldos.yaml` trocados — para o teste."""
    return dataclasses.replace(params, saldos={**params.saldos, **trocas})


def test_o_teto_e_30_do_saldo_devedor_antes_do_uso(reapurado, parametros):
    """Saldo devedor do Livro sintético: 17.000,00 de débito, sem crédito.

    O uso declarado já abateu a linha 011; o teto se mede antes dele.
    """
    _, apuracao = reapurado
    (controle,) = cc.montar(apuracao, parametros)
    assert controle.percentual_do_teto == 30
    assert controle.saldo_devedor == pytest.approx(17_000.00)
    assert controle.teto == pytest.approx(5_100.00)
    assert controle.disponivel == pytest.approx(ABERTURA_DE_AGOSTO + RECEBIDO)
    assert controle.a_utilizar == pytest.approx(5_100.00)
    assert any("teto de 30%" in p for p in controle.pendencias)


def test_dentro_do_teto_nao_ha_pendencia(agosto, parametros, tmp_path):
    base, apuracao = agosto
    devolvido = _devolver(base, apuracao, tmp_path, valor=5_000.00)
    base = tratar(devolvido, parametros=parametros)
    apuracao = apurar(base, parametros, ajustes=ler_ajustes(devolvido))
    (controle,) = cc.montar(apuracao, parametros)
    assert controle.saldo_devedor == pytest.approx(17_000.00)
    assert controle.pendencias == []


def test_o_teto_limita_so_quando_o_disponivel_passa_dele(agosto, parametros):
    """Sem saldo transportado nem recebido, a utilizar é o disponível."""
    _, apuracao = agosto
    (controle,) = cc.montar(apuracao, parametros)
    controle.saldo_inicial, controle.notas = 0.0, []
    assert controle.disponivel == 0
    assert controle.a_utilizar == 0


def test_setembro_abre_no_fim_de_agosto_transmitido(parametros, tmp_path):
    """Sem saldo declarado no mês, o 1200 abre no SLD_CRED_FIM do anterior."""
    base = tratar(_livro(tmp_path / "setembro.xlsx", mes=9), parametros=parametros)
    sem_declaracao = _sem(parametros, creditos_controlados=[])
    (controle,) = cc.montar(apurar(base, sem_declaracao), sem_declaracao)
    assert controle.saldo_inicial == pytest.approx(62_287.65, abs=CENTAVO)
    assert "transmitido" in controle.origem_do_saldo_inicial
    assert not any("não declarado" in p for p in controle.pendencias)


def test_abertura_diferente_do_fim_anterior_fica_pendente(parametros, tmp_path):
    base = tratar(_livro(tmp_path / "setembro.xlsx", mes=9), parametros=parametros)
    outro = _sem(parametros, creditos_controlados=[{
        "competencia": "2026-09",
        "por_estabelecimento": {2: {"MS090004": 50_000.00}},
    }])
    (controle,) = cc.montar(apurar(base, outro), outro)
    assert controle.saldo_inicial == pytest.approx(50_000.00)
    assert any("mês anterior" in p for p in controle.pendencias)


def test_o_recebido_declarado_vale_sobre_o_livro(parametros, tmp_path):
    """A NF-e vem com desconto: o Livro tem o documento, não o crédito."""
    base = tratar(_livro(tmp_path / "setembro.xlsx", mes=9), parametros=parametros)
    outro = _sem(parametros, creditos_controlados=[{
        "competencia": "2026-09",
        "por_estabelecimento": {2: {"MS090004": 62_287.65}},
        "recebido_por_estabelecimento": {2: {"MS090004": 70_000.00}},
    }])
    (controle,) = cc.montar(apurar(base, outro), outro)
    assert controle.recebido_no_livro == pytest.approx(RECEBIDO)
    assert controle.recebido == pytest.approx(70_000.00)
    assert controle.linha_1200().split("|")[5] == "70000"


def _setembro(**campos) -> cc.Controle:
    """O 1200 de Rio Brilhante de 09/2026, com os números da simulação."""
    controle = cc.Controle(
        estabelecimento=RB, empresa=2, cod_aj_apur="MS090004", descricao="",
        tipo_util="MS03", competencia="2026-09", saldo_inicial=62_287.65,
        recebido_declarado=62_720.00, saldo_devedor=226_295.20,
        percentual_do_teto=30.0, beneficio=173_659.00, percentual_fadefe=2.0,
    )
    for campo, valor in campos.items():
        setattr(controle, campo, valor)
    return controle


def test_o_beneficio_limita_o_uso_abaixo_do_teto():
    """Setembro: o teto é 67.888,56, mas o benefício cobre tudo menos 52.636,20."""
    c = _setembro()
    assert c.teto == pytest.approx(67_888.56)
    assert c.nao_coberto_pelo_beneficio == pytest.approx(52_636.20)
    assert c.a_utilizar == pytest.approx(52_636.20)
    assert c.a_recolher(c.a_utilizar) == 0
    assert c.beneficio_perdido(c.a_utilizar) == 0


def test_usar_acima_do_que_o_beneficio_nao_cobre_perde_beneficio():
    """62.720,00 zera o imposto igual, mas 10.083,80 de benefício não entram."""
    c = _setembro(utilizacoes=[SimpleNamespace(valor=62_720.00)])
    assert c.a_recolher(c.utilizado) == 0
    assert c.beneficio_deduzido(c.utilizado) == pytest.approx(163_575.20)
    assert c.beneficio_perdido(c.utilizado) == pytest.approx(10_083.80)
    assert any("10.083,80 do benefício" in p for p in c.pendencias)


def test_uso_dentro_dos_tres_limites_nao_tem_pendencia():
    c = _setembro(utilizacoes=[SimpleNamespace(valor=52_636.20)],
                  notas=[cc.NotaRecebida("247431", "23/09/2026", CHAVE, "", 57_702.40)])
    assert c.pendencias == []


def test_sem_beneficio_o_limite_e_o_teto():
    c = _setembro(beneficio=0.0)
    assert c.nao_coberto_pelo_beneficio == pytest.approx(226_295.20)
    assert c.a_utilizar == pytest.approx(67_888.56)


def test_o_historico_transmitido_fecha_mes_a_mes(parametros):
    """O SLD_CRED de cada mês é o SLD_CRED_FIM do anterior — a EFD entregue."""
    meses = sorted(parametros.transmitidos_do_controle(2, "MS090004").items())
    assert [m for m, _ in meses][:1] == ["2026-01"]
    for (_, antes), (mes, depois) in zip(meses, meses[1:]):
        fim = (antes["sld_cred"] + antes["cred_apr"] + antes["cred_receb"]
               - antes["cred_util"])
        assert depois["sld_cred"] == pytest.approx(fim, abs=CENTAVO), mes


def test_a_aba_mostra_o_ano_e_a_conta_do_uso(parametros, tmp_path):
    base = tratar(_livro(tmp_path / "setembro.xlsx", mes=9), parametros=parametros)
    planilha = openpyxl.load_workbook(
        escrever(base, tmp_path / "s.xlsx", apurar(base, parametros)))
    aba = planilha["REGISTRO 1200"]
    coluna_a = [linha[0] for linha in aba.iter_rows(values_only=True)]
    meses = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
             "Julho", "Agosto", "Setembro"]
    inicio = coluna_a.index("Janeiro")
    assert coluna_a[inicio:inicio + 9] == meses
    for rotulo in ("TXT do SPED", "Valor recebido por transf. de crédito",
                   "Valor transportado", "Total de crédito disponível",
                   "Saldo devedor de Rio Brilhante (antes do uso)",
                   "(a) 30% do saldo devedor",
                   "Benefício fiscal do mês (dedução da linha 012)",
                   "(b) Saldo devedor que o benefício não cobre",
                   "(c) Total de crédito disponível",
                   "Total de crédito a utilizar", "Benefício não aproveitado",
                   "Saldo a transportar para o mês seguinte"):
        assert rotulo in coluna_a, rotulo
    linha = coluna_a.index("Total de crédito a utilizar") + 1
    assert aba.cell(row=linha, column=2).value.startswith("=MAX(MIN(")
    setembro = inicio + 9
    assert aba.cell(row=setembro, column=7).value == (
        f"=C{setembro}+D{setembro}+E{setembro}-F{setembro}")


@pytest.mark.parametrize("valor, esperado", [
    (0.0, "0"), (62720.0, "62720"), (46138.68, "46138,68"), (0.1, "0,10"),
])
def test_numero_no_formato_da_efd(valor, esperado):
    assert cc._efd(valor) == esperado


# -- julho/2026, contra o Registro do ERP ------------------------------------

def test_julho_desmonta_o_registro_de_rio_brilhante(base_julho, parametros):
    """Linha 003 = 331.236,11 do Livro; linha 002 = 99.412,10 recebidos."""
    apuracao = apurar(base_julho, parametros)
    roteiro = lanc.montar(apuracao, parametros)
    assert roteiro.divergencias == []
    rb = {(x.parcela, x.linha): x.valor for x in roteiro.lancamentos
          if x.estabelecimento == RB}
    assert rb[("estorno", 3)] == pytest.approx(331_236.11, abs=CENTAVO)
    assert rb[("centralizacao_recebe_debito", 2)] == pytest.approx(
        99_412.10, abs=CENTAVO)
