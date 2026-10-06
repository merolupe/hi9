"""A planilha de resultado — o arquivo que o time fiscal abre.

Roda sobre o Livro sintético de Rio Brilhante (`sintetico.py`): uma venda que
gera débito e benefício e a nota do crédito recebido. Não precisa de dado real.

O que se confere aqui é a forma, pedida na rodada de setembro/2026: a ordem
das abas pela prioridade de leitura, nada congelado, borda em toda tabela, as
abas que saem ocultas, e as memórias de cálculo com o valor em célula própria
e a conta em fórmula — fórmula que tem que chegar ao número do motor.
"""
from __future__ import annotations

import re

import openpyxl
import pytest

from apurabot import ajustes as aj
from apurabot.apuracao import apurar
from apurabot.base_tratada import tratar
from apurabot.saida import ABAS_OCULTAS, ORDEM_DAS_ABAS, escrever
from sintetico import RB
from sintetico import livro as _livro

CENTAVO = 0.005


@pytest.fixture(scope="module")
def gerada(parametros, tmp_path_factory):
    pasta = tmp_path_factory.mktemp("planilha")
    base = tratar(_livro(pasta / "livro.xlsx"), parametros=parametros)
    apuracao = apurar(base, parametros)
    return apuracao, escrever(base, pasta / "saida.xlsx", apuracao)


@pytest.fixture(scope="module")
def planilha(gerada):
    return openpyxl.load_workbook(gerada[1])


# -- um avaliador mínimo, só para as fórmulas que a planilha escreve ----------

_REFERENCIA = re.compile(r"\b([A-Z]{1,2})(\d+)\b")
_INTERVALO = re.compile(r"SUM\(([A-Z]{1,2})(\d+):([A-Z]{1,2})(\d+)\)")


def avaliar(aba, coordenada: str) -> float:
    """Resolve SUM, MAX, MIN, ROUND, as quatro operações e referências."""
    valor = aba[coordenada].value
    if not (isinstance(valor, str) and valor.startswith("=")):
        return float(valor or 0.0)
    expressao = valor[1:]

    def soma(m) -> str:
        coluna, de, _, ate = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        return repr(sum(avaliar(aba, f"{coluna}{n}") for n in range(de, ate + 1)))

    expressao = _INTERVALO.sub(soma, expressao)
    expressao = _REFERENCIA.sub(lambda m: repr(avaliar(aba, m.group(0))), expressao)
    for funcao in ("MAX", "MIN", "ROUND"):
        expressao = expressao.replace(funcao, funcao.lower())
    return float(eval(expressao, {"__builtins__": {}},
                      {"max": max, "min": min, "round": round}))


def _linha(aba, rotulo: str, coluna: int = 1) -> int:
    return next(c[0].row for c in aba.iter_rows(min_col=coluna, max_col=coluna)
                if c[0].value == rotulo)


# -- as abas -----------------------------------------------------------------

def test_a_ordem_das_abas_segue_a_prioridade(planilha):
    nomes = planilha.sheetnames
    assert nomes == [n for n in ORDEM_DAS_ABAS if n in nomes]
    visiveis = [a.title for a in planilha.worksheets if a.sheet_state == "visible"]
    assert visiveis[:2] == ["REGISTRO DE APURAÇÃO", "APURAÇÃO EFETIVA"]
    assert visiveis[2] == "RESUMO E DETALHES"
    assert visiveis[-1] == "AJUSTES NO SANKHYA"


def test_resumo_e_visao_por_carga_saem_ocultas(planilha):
    assert ABAS_OCULTAS == {"RESUMO", "POR ESTABELECIMENTO E CARGA"}
    for nome in ABAS_OCULTAS:
        assert planilha[nome].sheet_state == "hidden", nome
    # O arquivo abre na primeira aba visível, não numa oculta.
    assert planilha.active.title == "REGISTRO DE APURAÇÃO"


def test_os_nomes_antigos_sumiram(planilha):
    for antigo in ("REGISTRO", "AJUSTES", "AJUSTES A LANÇAR", "APURAÇÃO POR FILIAL"):
        assert antigo not in planilha.sheetnames, antigo
    assert aj.ABA == "AJUSTES MANUAIS" and aj.ABA in planilha.sheetnames


def test_nada_fica_congelado(planilha):
    for aba in planilha.worksheets:
        assert aba.freeze_panes is None, aba.title


def test_toda_tabela_tem_borda_cinza(planilha):
    from apurabot.conferencia import ABAS_COM_FILIAL

    for aba in planilha.worksheets:
        # A coluna A dessas abas é a chave do filtro por filial, não tabela.
        filial = aba.title in ABAS_COM_FILIAL
        for linha in aba.iter_rows():
            ocupadas = [c for c in linha if c.value not in (None, "")
                        and not (filial and c.column == 1)]
            if len(ocupadas) < 2:
                continue                # texto solto: título, nota, instrução
            for celula in ocupadas:
                lado = celula.border.left
                assert lado.style == "thin", (aba.title, celula.coordinate)
                assert lado.color.rgb.endswith("808080"), (aba.title, celula.coordinate)


def test_sem_pendencia_nao_tem_aba_de_pendencias(gerada, planilha):
    apuracao, _ = gerada
    assert not apuracao.base.com_pendencia
    assert "PENDÊNCIAS" not in planilha.sheetnames


def test_o_registro_nao_traz_a_marca_de_ajuste(planilha):
    textos = [c.value for linha in planilha["REGISTRO DE APURAÇÃO"].iter_rows()
              for c in linha if isinstance(c.value, str)]
    assert not any("AGUARDA AJUSTE" in t for t in textos)
    assert "Situação" not in textos


# -- AJUSTES NO SANKHYA ------------------------------------------------------

def test_ajustes_no_sankhya_e_titulo_e_tabela(planilha):
    aba = planilha["AJUSTES NO SANKHYA"]
    assert aba["A1"].value == (
        "AJUSTES PARA GUIAR AJUSTES DO SANKHYA - Utilizar após finalizar "
        "ajustes manuais"
    )
    assert aba["A2"].value == "Nome Fantasia"
    # Nenhuma linha em branco no meio: atrapalha o filtro.
    corpo = [linha[0].value for linha in aba.iter_rows(min_row=3)]
    assert corpo and all(corpo)
    assert aba.auto_filter.ref == f"A2:H{aba.max_row}"
    textos = [c.value for linha in aba.iter_rows() for c in linha if c.value]
    assert "CONFERÊNCIA COM O REGISTRO" not in textos


# -- REGISTRO 1200 -----------------------------------------------------------

def test_o_txt_do_sped_traz_as_linhas_da_efd(planilha):
    aba = planilha["REGISTRO 1200"]
    n = _linha(aba, "TXT do SPED") + 1
    assert str(aba[f"A{n}"].value).startswith("|1200|")
    textos = [str(c.value) for linha in aba.iter_rows() for c in linha if c.value]
    for sumido in ("Crédito recebido", "Crédito utilizado", "Saldo final a cadastrar"):
        assert not any(t.startswith(sumido) for t in textos), sumido


def test_o_saldo_final_do_1200_e_formula(gerada, planilha, parametros):
    from apurabot import controle_de_creditos as cc

    apuracao, _ = gerada
    (controle,) = cc.montar(apuracao, parametros)
    aba = planilha["REGISTRO 1200"]
    n = _linha(aba, "Agosto")           # o mês em apuração, abaixo do histórico
    assert _linha(aba, "Janeiro") == _linha(aba, "Registro 1200") + 1
    assert aba[f"G{n}"].value == f"=C{n}+D{n}+E{n}-F{n}"
    assert avaliar(aba, f"G{n}") == pytest.approx(controle.saldo_final, abs=CENTAVO)


def test_a_conta_do_uso_chega_ao_numero_do_motor(gerada, planilha, parametros):
    from apurabot import controle_de_creditos as cc

    apuracao, _ = gerada
    (controle,) = cc.montar(apuracao, parametros)
    aba = planilha["REGISTRO 1200"]
    for rotulo, esperado in (
        ("Total de crédito disponível", controle.disponivel),
        ("(a) 30% do saldo devedor", controle.teto),
        ("(b) Saldo devedor que o benefício não cobre",
         controle.nao_coberto_pelo_beneficio),
        ("Total de crédito a utilizar", controle.a_utilizar),
    ):
        n = _linha(aba, rotulo)
        assert str(aba[f"B{n}"].value).startswith("="), rotulo
        assert avaliar(aba, f"B{n}") == pytest.approx(esperado, abs=CENTAVO), rotulo

    # O resultado: a coluna C é o uso recomendado, que não perde benefício.
    uso = controle.a_utilizar
    for rotulo, esperado in (
        ("Benefício deduzido (linha 012)", controle.beneficio_deduzido(uso)),
        ("Benefício não aproveitado", 0.0),
        ("ICMS a recolher", controle.a_recolher(uso)),
        ("Saldo a transportar para o mês seguinte", controle.disponivel - uso),
    ):
        n = _linha(aba, rotulo)
        assert avaliar(aba, f"C{n}") == pytest.approx(esperado, abs=CENTAVO), rotulo


# -- RESUMO E DETALHES: as memórias em fórmula --------------------------------

def test_a_memoria_do_beneficio_chega_ao_numero_do_motor(gerada, planilha):
    apuracao, _ = gerada
    b = apuracao.filiais[RB].beneficio
    assert b.credito_presumido, "o Livro sintético deixou de gerar benefício"
    aba = planilha["RESUMO E DETALHES"]

    credito = _linha(aba, "(=) crédito da parcela incentivada")
    assert aba[f"B{credito}"].value.startswith("=MAX(")
    assert avaliar(aba, f"B{credito}") == pytest.approx(
        b.credito_da_parcela_incentivada, abs=CENTAVO)

    for nome, parcela in (("intraestadual", b.intra), ("interestadual", b.inter)):
        n = _linha(aba, nome)
        assert avaliar(aba, f"D{n}") == pytest.approx(
            parcela.base_do_incentivo, abs=CENTAVO), nome
        assert avaliar(aba, f"F{n}") == pytest.approx(
            parcela.credito_presumido, abs=CENTAVO), nome

    total = _linha(aba, "Crédito presumido total")
    assert aba[f"F{total}"].value.startswith("=SUM(")
    assert avaliar(aba, f"F{total}") == pytest.approx(b.credito_presumido, abs=CENTAVO)

    # O FADEFE parte do total da memória, não de um número colado.
    linhas_rb = [c[0].row for c in aba.iter_rows(min_col=1, max_col=1)
                 if c[0].value == RB]
    fadefe = next(n for n in linhas_rb if aba[f"B{n}"].value == f"=F{total}")
    assert avaliar(aba, f"D{fadefe}") == pytest.approx(b.fadefe, abs=CENTAVO)


def test_a_memoria_nao_mistura_valor_com_texto(planilha):
    aba = planilha["RESUMO E DETALHES"]
    for rotulo in ("Crédito industrial bruto", "(−) estorno industrial",
                   "Saldo próprio da centralizadora", "Saldo final do grupo"):
        n = _linha(aba, rotulo)
        valor = aba[f"B{n}"].value
        assert isinstance(valor, (int, float)) or str(valor).startswith("="), rotulo


def test_a_centralizacao_fecha_em_formula(gerada, planilha):
    apuracao, _ = gerada
    aba = planilha["RESUMO E DETALHES"]
    finais = [c[0].row for c in aba.iter_rows(min_col=1, max_col=1)
              if c[0].value == "Saldo final do grupo"]
    assert len(finais) == len(apuracao.centralizacao)
    for n, grupo in zip(finais, apuracao.centralizacao):
        assert avaliar(aba, f"B{n}") == pytest.approx(grupo.saldo_final, abs=CENTAVO)


def test_o_texto_da_convencao_de_caixa_saiu(planilha):
    textos = [str(c.value) for linha in planilha["RESUMO E DETALHES"].iter_rows()
              for c in linha if c.value]
    assert not any(t.startswith("Saldo na convenção de caixa") for t in textos)


def test_a_apuracao_efetiva_nao_repete_a_memoria_do_beneficio(planilha):
    textos = [c.value for linha in planilha["APURAÇÃO EFETIVA"].iter_rows()
              for c in linha if c.value]
    assert "BENEFÍCIO FISCAL" not in textos


# -- a aba de ajustes ainda lê arquivo antigo --------------------------------

def test_arquivo_com_a_aba_de_nome_antigo_continua_sendo_lido(tmp_path):
    wb = openpyxl.Workbook()
    aba = wb.active
    aba.title = "AJUSTES"
    aba.append([aj.TITULO_PARCELAS])
    aba.append(["estabelecimento", "atividade", "linha", "valor", "motivo",
                "responsável", "aprovador", "observação padrão"])
    aba.append([RB, "Produção", "006", 100.0, "teste", "Fulano", "Ciclano", None])
    destino = tmp_path / "antigo.xlsx"
    wb.save(destino)
    lidos = aj.ler_aba(destino)
    assert [p.valor for p in lidos.parcelas] == [100.0]


def test_registro_e_apuracao_efetiva_filtram_por_filial(planilha):
    """A2 é o filtro; toda linha abaixo diz de que filial ela é."""
    from apurabot.conferencia import ABAS_COM_FILIAL

    for nome in ABAS_COM_FILIAL:
        aba = planilha[nome]
        assert aba["A2"].value == "Filial", nome
        assert aba.auto_filter.ref == f"A2:A{aba.max_row}", nome
        vazias = [n for n in range(3, aba.max_row + 1)
                  if not aba.cell(row=n, column=1).value]
        assert vazias == [], (nome, vazias[:5])
