"""A APURAÇÃO EFETIVA de MS mostra de onde a carga efetiva veio.

Alíquota da nota → redução de base → carga efetiva → % do crédito estornado.
Livro sintético de Rio Brilhante, sem dado real.
"""
from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from apurabot.apuracao import apurar
from apurabot.base_tratada import tratar
from apurabot.ingestao import COLUNAS
from apurabot.saida import escrever

RB = "HINOVE (RIO BRILHANTE)"


@pytest.fixture(scope="module")
def planilha(parametros, tmp_path_factory):
    pasta = tmp_path_factory.mktemp("reducao")
    cabecalho = list(COLUNAS)
    wb = openpyxl.Workbook()
    aba = wb.active
    aba.append(cabecalho)
    notas = [
        # produto, código, cfop, valor, base, alíquota
        ("ZINCO 15 GR", 110000001, 2101, 82_500.00, 47_140.50, 7),
        ("SACO SOLD. 50KG", 201000001, 2101, 47_880.00, 47_880.00, 7),
        ("CLORETO DE AMONIO", 110000002, 2906, 416_010.43, 138_656.02, 12),
    ]
    for n, (desc, prod, cfop, vc, base, aliq) in enumerate(notas, start=1):
        valores = {
            "Nro único Nota": n, "Nro. da nota": n, "Empresa": 2,
            "Nome Fantasia (Empresa)": RB,
            "Dt. do movimento": dt.datetime(2026, 9, 10),
            "Dt. do documento": dt.datetime(2026, 9, 10),
            "CFOP": cfop, "Entrada/Saída": "Entrada",
            "Cód. de tributação": "20" if base < vc else "00",
            "Vlr. contábil": vc, "Base do ICMS": base, "Alíquota ICMS": aliq,
            "Vlr. do ICMS": round(base * aliq / 100, 2),
            "UF de Origem": "SP", "UF de Destino": "MS",
            "Espécie do documento": "NFE", "Modelo do Documento": "55",
            "Produto": prod, "Descrição (Produto)": desc,
        }
        aba.append([valores.get(c) for c in cabecalho])
    livro = pasta / "livro.xlsx"
    wb.save(livro)
    base = tratar(livro, parametros=parametros)
    destino = escrever(base, pasta / "saida.xlsx", apurar(base, parametros))
    return openpyxl.load_workbook(destino)


def _linhas_de_produto(planilha):
    aba = planilha["APURAÇÃO EFETIVA"]
    return {
        linha[3]: linha for linha in aba.iter_rows(values_only=True)
        if linha[3] and linha[0] in (None, "") and linha[2]
    }


def test_ms_mostra_aliquota_e_reducao_no_fim_da_tabela(planilha):
    aba = planilha["APURAÇÃO EFETIVA"]
    cabecalho = next(linha for linha in aba.iter_rows(values_only=True)
                     if linha[0] == "CFOP")
    assert list(cabecalho[11:13]) == ["Alíquota da nota", "Redução de base"]


@pytest.mark.parametrize("produto, carga, aliquota, reducao", [
    ("ZINCO 15 GR", "4%", "7%", 0.4286),
    ("SACO SOLD. 50KG", "7%", "7%", 0.0),
    ("CLORETO DE AMONIO", "4%", "12%", 0.6667),
])
def test_a_reducao_explica_a_carga(planilha, produto, carga, aliquota, reducao):
    linha = _linhas_de_produto(planilha)[produto]
    assert linha[2] == carga
    assert linha[11] == aliquota
    assert linha[12] == pytest.approx(reducao, abs=0.00005)


def test_quem_ja_veio_reduzido_nao_estorna_e_quem_veio_cheio_estorna(planilha):
    linhas = _linhas_de_produto(planilha)
    assert linhas["ZINCO 15 GR"][9] == pytest.approx(0.0)
    assert linhas["CLORETO DE AMONIO"][9] == pytest.approx(0.0)
    assert linhas["SACO SOLD. 50KG"][9] == pytest.approx(3_351.60 * 0.4286, abs=0.01)
