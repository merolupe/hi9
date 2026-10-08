"""Só segrega por atividade o estabelecimento com benefício fiscal.

Corumbá e Rio Brilhante são MS e usam o mesmo mapa de atividades, mas o
benefício é só de Rio Brilhante. Em Corumbá a segregação não alimentava conta
nem declaração, e mostrá-la fazia parecer que alimentava. Decisão nº 6.

Roda sobre um Livro sintético, sem dado real.
"""
from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from apurabot import ajustes as aj
from apurabot.apuracao import _segrega_por_atividade, apurar, de_declarados
from apurabot.base_tratada import tratar
from apurabot.ingestao import COLUNAS
from apurabot.nucleo import atividade as ativ
from apurabot.saida import escrever

RB = "HINOVE (RIO BRILHANTE)"
CORUMBA = "HINOVE (CORUMBÁ- MS)"


def _livro(destino):
    cabecalho = list(COLUNAS)

    def linha(**campos):
        valores = {
            "Dt. do movimento": dt.datetime(2026, 9, 10),
            "Dt. do documento": dt.datetime(2026, 9, 10),
            "Cód. de tributação": "00", "UF de Origem": "MS",
            "UF de Destino": "MS", "Espécie do documento": "NFE",
            "Produto": 100000001, "Modelo do Documento": "55",
        }
        valores.update(campos)
        return [valores.get(c) for c in cabecalho]

    wb = openpyxl.Workbook()
    aba = wb.active
    aba.append(cabecalho)
    for n, (codigo, nome) in enumerate([(2, RB), (9, CORUMBA)], start=1):
        aba.append(linha(**{
            "Nome Fantasia (Empresa)": nome, "Empresa": codigo,
            "Nro único Nota": n, "Nro. da nota": n, "CFOP": 5101,
            "Descrição da CFOP": "Vda prod do estab", "Entrada/Saída": "Saída",
            "Vlr. contábil": 100_000, "Base do ICMS": 100_000,
            "Alíquota ICMS": 17, "Vlr. do ICMS": 17_000,
        }))
    wb.save(destino)
    return destino


@pytest.fixture(scope="module")
def apurado(parametros, tmp_path_factory):
    pasta = tmp_path_factory.mktemp("segregacao")
    base = tratar(_livro(pasta / "livro.xlsx"), parametros=parametros)
    apuracao = apurar(base, parametros)
    return base, apuracao, pasta


def test_so_quem_tem_beneficio_segrega(apurado):
    _, apuracao, _ = apurado
    assert apuracao.filiais[RB].segrega_por_atividade
    corumba = apuracao.filiais[CORUMBA]
    assert not corumba.segrega_por_atividade
    assert corumba.por_atividade == {}
    assert all(not a.atividade for a in corumba.apuradas)


def test_a_conta_de_corumba_nao_muda(apurado):
    """O débito continua o do Livro — só a divisão deixa de sair."""
    _, apuracao, _ = apurado
    assert apuracao.filiais[CORUMBA].debito == pytest.approx(17_000.00)


def test_ajuste_de_corumba_nao_pede_atividade(apurado, parametros):
    base, _, _ = apurado

    def com_ajuste_em(estabelecimento):
        declarados = aj.Declarados(parcelas=[
            aj.Ajuste(estabelecimento=estabelecimento, linha=3, valor=100.0,
                      motivo="m", responsavel="r", aprovador="a",
                      onde="aba AJUSTES, linha 9")
        ])
        return apurar(base, parametros, ajustes=de_declarados(declarados))

    assert not com_ajuste_em(CORUMBA).bloqueios_de_ajuste
    assert com_ajuste_em(RB).bloqueios_de_ajuste


def test_a_planilha_nao_mostra_atividade_em_corumba(apurado):
    base, apuracao, pasta = apurado
    wb = openpyxl.load_workbook(escrever(base, pasta / "saida.xlsx", apuracao))

    resumo = wb["RESUMO E DETALHES"]
    titulo = next(
        c.row for c in resumo["A"]
        if c.value == "Segregação por atividade — base do benefício fiscal"
    )
    bloco = []
    for linha in resumo.iter_rows(min_row=titulo + 2, values_only=True):
        if not linha[0]:
            break
        bloco.append(linha[0])
    assert bloco and set(bloco) == {RB}

    rotulos = {"Produção", "Comercial"}
    efetiva = wb["APURAÇÃO EFETIVA"]
    de_corumba = {
        v for linha in efetiva.iter_rows(values_only=True)
        if linha[0] == CORUMBA for v in linha if isinstance(v, str)
    }
    assert not de_corumba & rotulos
    de_rb = {
        v for linha in efetiva.iter_rows(values_only=True)
        if linha[0] == RB for v in linha if isinstance(v, str)
    }
    assert "Produção" in de_rb


def test_beneficio_sem_mapa_de_atividade_nao_passa_calado(parametros):
    """Sem o mapa, o benefício sairia zerado sem ninguém perceber."""
    ficha = {"nome": "FILIAL NOVA", "beneficio_fiscal": "qualquer"}
    with pytest.raises(ativ.MapaDeAtividadeAusente, match="FILIAL NOVA"):
        _segrega_por_atividade(ficha, "GO", parametros)
    assert not _segrega_por_atividade({"nome": "X"}, "MS", parametros)
