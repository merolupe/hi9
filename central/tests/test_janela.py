"""A janela da central: o menu, a execução de uma ferramenta e as travas."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import openpyxl

from central import ferramentas

RAIZ = Path(__file__).resolve().parents[2]


# -- o menu ----------------------------------------------------------------

def test_o_menu_traz_todas_as_ferramentas_do_setor(janela):
    codigo, dados = janela.pedir("/ferramentas")
    assert codigo == 200
    por_id = {f["id"]: f for f in dados["ferramentas"]}
    assert set(por_id) == {
        "apurabot", "dixml", "fiscalbot", "gerarpendentes",
        "gerarservpend", "resumoexecutivo", "conhecimento", "faturabot",
    }


def test_ferramenta_ainda_nao_importada_aparece_apagada_em_vez_de_sumir(janela):
    """O time enxerga o que falta, em vez de descobrir quando precisar."""
    _, dados = janela.pedir("/ferramentas")
    por_id = {f["id"]: f for f in dados["ferramentas"]}
    assert por_id["faturabot"]["estado"] == ferramentas.A_IMPORTAR
    assert por_id["faturabot"]["resumo"]
    assert por_id["dixml"]["estado"] == ferramentas.DISPONIVEL
    assert por_id["fiscalbot"]["estado"] == ferramentas.DISPONIVEL
    assert por_id["gerarservpend"]["estado"] == ferramentas.DISPONIVEL
    assert por_id["gerarpendentes"]["estado"] == ferramentas.DISPONIVEL
    assert por_id["resumoexecutivo"]["estado"] == ferramentas.DISPONIVEL
    assert por_id["conhecimento"]["estado"] == ferramentas.DISPONIVEL
    assert por_id["apurabot"]["estado"] == ferramentas.JANELA_PROPRIA


def test_a_ferramenta_de_servicos_pede_varios_arquivos_em_qualquer_ordem(janela):
    """Quatro relatórios, uma caixa só: cada um é reconhecido pelo cabeçalho."""
    _, dados = janela.pedir("/ferramentas")
    servpend = next(f for f in dados["ferramentas"] if f["id"] == "gerarservpend")
    assert servpend["entrada"]["varios"] is True
    assert servpend["entrada"]["extensoes"] == [".xls", ".xlsx", ".xlsm"]
    assert servpend["verbo"] and servpend["detalhe"]


def test_so_quem_guarda_estado_declara_configuracao(janela):
    """O Fiscalbot e a base de conhecimento guardam estado; o DiXML não."""
    _, dados = janela.pedir("/ferramentas")
    por_id = {f["id"]: f for f in dados["ferramentas"]}
    assert por_id["fiscalbot"]["tem_configuracao"] is True
    assert por_id["conhecimento"]["tem_configuracao"] is True
    assert por_id["dixml"]["tem_configuracao"] is False


def test_a_base_de_conhecimento_abre_na_tela_mesmo_sem_base_no_disco(janela):
    """Máquina nova: a tela tem que abrir vazia e dizer o que fazer, e não cair."""
    codigo, dados = janela.pedir("/configuracao", ferramenta="conhecimento")

    assert codigo == 200, dados
    assert [s["id"] for s in dados["secoes"]] == [
        "parceiros", "gestores", "operacoes"]
    for secao in dados["secoes"]:
        assert secao["explicacao"] and secao["campos"]
        assert secao["fixa"] is False


def test_a_tela_da_base_marca_a_evidencia_como_campo_de_leitura(janela):
    """Evidência e carimbo se leem; proposta importada não se rasura."""
    _, dados = janela.pedir("/configuracao", ferramenta="conhecimento")
    parceiros = next(s for s in dados["secoes"] if s["id"] == "parceiros")
    por_chave = {c["chave"]: c["tipo"] for c in parceiros["campos"]}

    assert por_chave["guardiao_evidencia"] == "leitura"
    assert por_chave["guardiao_proposto"] == "leitura"
    assert por_chave["trilha"] == "leitura"
    assert por_chave["guardiao"] == "texto"
    assert por_chave["categoria"] == "texto"


def test_quem_roda_na_janela_declara_o_que_pede(janela):
    _, dados = janela.pedir("/ferramentas")
    dixml = next(f for f in dados["ferramentas"] if f["id"] == "dixml")
    assert dixml["entrada"]["extensoes"] == [".zip"]
    assert dixml["entrada"]["varios"] is True


def test_a_pagina_e_servida(janela):
    codigo, corpo = janela.pedir("/")
    assert codigo == 200
    assert b"Central" in corpo


# -- a trava de sessão -----------------------------------------------------

def test_sem_a_chave_da_sessao_nada_responde(janela):
    """Outra pessoa logada na mesma máquina também alcança 127.0.0.1."""
    for rota in ("/", "/ferramentas", "/baixar"):
        codigo, _ = janela.pedir(rota, chave="chave-errada")
        assert codigo == 403, rota


# -- o DiXML de ponta a ponta ---------------------------------------------

def _enviar_e_executar(janela, lote, nome="lote.zip"):
    codigo, envio = janela.pedir("/enviar", corpo=lote, ferramenta="dixml", nome=nome)
    assert codigo == 200, envio
    return janela.postar("/executar", ferramenta="dixml")


def test_o_lote_entra_pela_janela_e_sai_planilha(janela, lote):
    codigo, dados = _enviar_e_executar(janela, lote)
    assert codigo == 200, dados

    fichas = {f["rotulo"]: f["valor"] for f in dados["fichas"]}
    assert fichas["Notas (NF-e)"] == "1"
    assert fichas["Linhas de item"] == "1"
    assert dados["planilha"].startswith("DiXML_")

    codigo, corpo = janela.pedir("/baixar")
    assert codigo == 200
    livro = openpyxl.load_workbook(io.BytesIO(corpo))
    assert livro.sheetnames == ["NFe", "CTe"]
    aba = livro["NFe"]
    cabecalho = [c.value for c in aba[1]]
    assert aba.cell(row=2, column=cabecalho.index("Chave") + 1).value == (
        "35260612345678000199550010000001231000001234"
    )


def test_o_nome_do_pacote_chega_a_planilha(janela, lote):
    """É por ele que se acha a origem de uma linha meses depois."""
    _enviar_e_executar(janela, lote, nome="XMLs_junho_filial_02.zip")
    _, corpo = janela.pedir("/baixar")
    aba = openpyxl.load_workbook(io.BytesIO(corpo))["NFe"]
    assert aba.cell(row=2, column=1).value == "XMLs_junho_filial_02.zip::nota.xml"


def test_dois_pacotes_de_mesmo_nome_nao_se_atropelam(janela, lote):
    for _ in range(2):
        codigo, envio = janela.pedir(
            "/enviar", corpo=lote, ferramenta="dixml", nome="lote.zip"
        )
        assert codigo == 200, envio
    codigo, dados = janela.postar("/executar", ferramenta="dixml")
    assert codigo == 200, dados
    assert {f["rotulo"]: f["valor"] for f in dados["fichas"]}["Linhas de item"] == "2"


def test_os_arquivos_enviados_somem_depois_de_rodar(janela, lote):
    _enviar_e_executar(janela, lote)
    assert janela.sessao.recebidos_de("dixml") == []


def test_o_que_nao_e_nota_aparece_na_tela_em_vez_de_sumir(janela, tmp_path):
    caminho = tmp_path / "com_evento.zip"
    with zipfile.ZipFile(caminho, "w") as zf:
        zf.writestr("evento.xml", "<procEventoNFe><evento/></procEventoNFe>")
    _, dados = _enviar_e_executar(janela, caminho.read_bytes())
    titulos = " ".join(l["titulo"] for l in dados["listas"])
    assert "Não entraram na planilha" in titulos


# -- o que a janela recusa -------------------------------------------------

def test_arquivo_do_tipo_errado_e_recusado_com_o_motivo(janela):
    codigo, dados = janela.pedir(
        "/enviar", corpo=b"conteudo", ferramenta="dixml", nome="livro.xlsx"
    )
    assert codigo == 400
    assert ".zip" in dados["erro"]


def test_arquivo_vazio_e_recusado(janela):
    codigo, dados = janela.pedir(
        "/enviar", corpo=b"", ferramenta="dixml", nome="lote.zip"
    )
    assert codigo == 400
    assert "vazio" in dados["erro"]


def test_executar_sem_arquivo_avisa_em_vez_de_estourar(janela):
    codigo, dados = janela.postar("/executar", ferramenta="dixml")
    assert codigo == 400
    assert "Nenhum arquivo" in dados["erro"]


def test_ferramenta_desconhecida_nao_derruba_a_janela(janela):
    codigo, dados = janela.postar("/executar", ferramenta="naoexiste")
    assert codigo == 404
    assert "erro" in dados


def test_ferramenta_ainda_nao_importada_nao_finge_que_roda(janela):
    codigo, dados = janela.postar("/executar", ferramenta="faturabot")
    assert codigo == 400
    assert "ainda não roda" in dados["erro"]


def test_o_nome_do_arquivo_nao_escapa_da_pasta_temporaria(janela, lote):
    """O nome vem do navegador e vai para o disco: fica só o nome."""
    codigo, _ = janela.pedir(
        "/enviar", corpo=lote, ferramenta="dixml",
        nome="..%2F..%2Fescapou.zip",
    )
    assert codigo == 200
    recebido = janela.sessao.recebidos_de("dixml")[0]
    assert recebido.name == "escapou.zip"
    assert janela.sessao.pasta in recebido.parents


def test_zip_corrompido_vira_recado_em_vez_de_traceback(janela):
    codigo, dados = janela.pedir(
        "/enviar", corpo=b"isto nao e um zip", ferramenta="dixml", nome="lote.zip"
    )
    assert codigo == 200
    codigo, dados = janela.postar("/executar", ferramenta="dixml")
    assert codigo == 200
    titulos = " ".join(l["titulo"] for l in dados["listas"])
    assert "ilegíveis" in titulos


# -- a tela de configuração ------------------------------------------------

def test_a_tela_de_configuracao_traz_secoes_e_dados(janela):
    codigo, dados = janela.pedir("/configuracao", ferramenta="fiscalbot")
    assert codigo == 200, dados
    ids = [s["id"] for s in dados["secoes"]]
    assert "regras" in ids and "parceiros_sn" in ids
    assert dados["dados"]["regras"], "a carga de fábrica tem que vir preenchida"
    for secao in dados["secoes"]:
        assert secao["explicacao"]
        assert secao["campos"]


def test_ferramenta_sem_configuracao_diz_isso(janela):
    codigo, dados = janela.pedir("/configuracao", ferramenta="dixml")
    assert codigo == 400
    assert "não tem o que configurar" in dados["erro"]


def test_gravar_configuracao_com_erro_nao_grava(janela):
    """Base inconsistente auditaria o mês inteiro errado."""
    import json
    corpo = json.dumps({"dados": {
        "regras": [{"id": "X1", "cfop": "5101", "ativa": True},
                   {"id": "X1", "cfop": "5102", "ativa": True}],
        "parametros": [{"tolerancia_carga": "0.05",
                        "cst_exigem_icms_positivo": "00"}],
    }}).encode("utf-8")
    codigo, dados = janela.pedir("/configuracao", corpo=corpo, ferramenta="fiscalbot")
    assert codigo == 200
    assert dados["gravou"] is False
    assert any(p["gravidade"] == "erro" for p in dados["problemas"])


def test_a_tela_devolve_os_problemas_para_a_pessoa_ler(janela):
    import json
    corpo = json.dumps({"dados": {"regras": [], "parametros": [{}]}}).encode("utf-8")
    codigo, dados = janela.pedir("/configuracao", corpo=corpo, ferramenta="fiscalbot")
    assert codigo == 200
    assert dados["gravou"] is False
    assert dados["problemas"], "recusar sem dizer o motivo não ajuda ninguém"
    for problema in dados["problemas"]:
        assert problema["onde"] and problema["mensagem"]


# -- o Apurabot, que tem tela própria --------------------------------------

def test_o_apurabot_abre_em_janela_propria_sem_ser_reescrito(janela):
    codigo, dados = janela.pedir("/abrir", ferramenta="apurabot")
    assert codigo == 200, dados
    assert dados["endereco"].startswith("http://127.0.0.1:")
    assert "chave=" in dados["endereco"]


def test_clicar_duas_vezes_no_apurabot_reaproveita_a_mesma_janela(janela):
    _, primeira = janela.pedir("/abrir", ferramenta="apurabot")
    _, segunda = janela.pedir("/abrir", ferramenta="apurabot")
    assert primeira["endereco"] == segunda["endereco"]


def test_ferramenta_sem_tela_propria_nao_finge_ter(janela):
    codigo, dados = janela.pedir("/abrir", ferramenta="dixml")
    assert codigo == 400
    assert "janela própria" in dados["erro"]


# -- o encerramento --------------------------------------------------------

def test_encerrar_avisa_a_janela_e_limpa_o_que_ficou(janela, lote):
    janela.pedir("/enviar", corpo=lote, ferramenta="dixml", nome="lote.zip")
    pasta = janela.sessao.pasta
    codigo, _ = janela.pedir("/encerrar")
    assert codigo == 200
    assert janela.sessao.encerrar.is_set()
    janela.sessao.limpar()
    assert not pasta.exists()


# -- o lançador ------------------------------------------------------------

def test_o_lancador_da_central_pergunta_ao_python_em_vez_de_adivinhar():
    """Escolher `py -3` por vir primeiro na lista foi o defeito de antes."""
    bat = (RAIZ / "Hinove.bat").read_text(encoding="cp1252")
    assert "verificar.py" in bat
    for candidato in ("python", "py -3"):
        assert f'"{candidato}"' in bat


def test_os_arquivos_de_entrada_estao_na_raiz():
    for arquivo in ("Hinove.bat", "Apurabot.bat", "rodar.py", "verificar.py"):
        assert (RAIZ / arquivo).is_file(), arquivo


# -- o GerarServPend de ponta a ponta --------------------------------------

def _relatorio(caminho: Path, colunas: list[str], linhas: list[list]) -> bytes:
    """Um relatório sintético, do jeito que o Sankhya entrega."""
    from openpyxl import Workbook

    livro = Workbook()
    aba = livro.active
    aba.append(["Relatorio de exemplo"])
    aba.append(["Emitido em 14/09/2026"])
    aba.append(colunas)
    for linha in linhas:
        aba.append(linha)
    livro.save(str(caminho))
    return caminho.read_bytes()


def _relatorios_de_servicos(tmp_path) -> tuple[bytes, bytes]:
    """Um ASIS e um Portal de Compras sintéticos, que casam numa nota."""
    from pendentes import parametros

    # Os nomes das colunas vêm da carga de fábrica, que é onde eles moram —
    # copiá-los para dentro do teste seria manter duas listas em sincronia.
    fabrica = parametros.carregar_fabrica()
    colunas_de = lambda fonte: [e.nome for e in parametros.colunas_de(fabrica, fonte)]

    cnpj = "99888777000166"
    asis = _relatorio(tmp_path / "ASIS.xlsx", colunas_de("asis"), [
        ["1234", "05/08/2026", "", "", "17.01", "Assessoria", "Campo Grande",
         "PRESTADOR INVENTADO LTDA", cnpj, "1500.00", "Servico", "11222333000144",
         "", 5, 10, 0.65, 1, 3, 5, 1, 2, 0, 0, "VER-1"],
        ["9999", "06/08/2026", "", "", "17.01", "Assessoria", "Campo Grande",
         "PRESTADOR INVENTADO LTDA", cnpj, "90.00", "Servico", "11222333000144",
         "", 5, 10, 0.65, 1, 3, 5, 1, 2, 0, 0, "VER-2"],
    ])
    portal = _relatorio(tmp_path / "PC27.xlsx", colunas_de("portal_de_compras"), [
        ["COMPRADOR", 1, "VENDA DE SERVICO", "1234", "4001", "PARCEIRO LTDA",
         1500.00, "REQUISITANTE", "1", "HINOVE MATRIZ", "2020", "SERVICOS",
         "CR 100", cnpj, "11222333000144", "10/08/2026"],
    ])
    return asis, portal


def test_os_relatorios_da_semana_entram_pela_janela_e_sai_planilha(
        janela, tmp_path, monkeypatch):
    """A prova da costura: os arquivos entram, o Resultado volta para a tela.

    O livro e o snapshot são desviados para uma pasta do teste — eles são
    dados da empresa e moram fora do git, e um teste não escreve na pasta de
    quem roda de verdade.
    """
    from pendentes import estado, snapshot

    monkeypatch.setattr(estado, "_raiz", lambda: tmp_path)
    monkeypatch.setattr(snapshot, "_raiz", lambda: tmp_path)

    asis, portal = _relatorios_de_servicos(tmp_path)

    for nome, corpo in (("ASIS.xlsx", asis), ("PC27.xlsx", portal)):
        codigo, envio = janela.pedir("/enviar", corpo=corpo,
                                     ferramenta="gerarservpend", nome=nome)
        assert codigo == 200, envio

    codigo, dados = janela.postar("/executar", ferramenta="gerarservpend")
    assert codigo == 200, dados

    fichas = {f["rotulo"]: f["valor"] for f in dados["fichas"]}
    assert fichas["Lançadas"] == "1"
    assert fichas["Pendentes"] == "1"
    assert dados["planilha"].startswith("Notas_Servico_Pendentes_")
    assert "Confronto por procedimento" in [l["titulo"] for l in dados["listas"]]

    codigo, corpo = janela.pedir("/baixar")
    assert codigo == 200
    livro = openpyxl.load_workbook(io.BytesIO(corpo))
    assert livro.sheetnames == ["Lancadas", "Pendentes", "Canceladas",
                                "Sem Correspondencia ASIS", "Fora do relatorio"]


def test_a_ferramenta_de_mercadorias_pede_varios_arquivos_em_qualquer_ordem(
        janela):
    """Três relatórios, uma caixa só: cada um é reconhecido pelo cabeçalho."""
    _, dados = janela.pedir("/ferramentas")
    pendentes = next(f for f in dados["ferramentas"]
                     if f["id"] == "gerarpendentes")
    assert pendentes["entrada"]["varios"] is True
    assert pendentes["entrada"]["extensoes"] == [".xls", ".xlsx", ".xlsm"]
    assert pendentes["verbo"] and pendentes["detalhe"]


def test_o_xml_e_a_conferencia_entram_pela_janela_e_sai_a_planilha(
        janela, tmp_path, monkeypatch):
    """A prova da costura do lado de mercadorias, com as sete abas."""
    from pendentes import estado, parametros, snapshot

    monkeypatch.setattr(estado, "_raiz", lambda: tmp_path)
    monkeypatch.setattr(snapshot, "_raiz", lambda: tmp_path)

    fabrica = parametros.carregar_fabrica()
    colunas_de = lambda fonte: [e.nome for e in parametros.colunas_de(fabrica, fonte)]

    chave = "35260612345678000199550010000001231000001234"
    def nota(numero, chave_da_nota, **campos):
        return [
            numero, "4", "FORNECEDOR INVENTADO LTDA", "01/07/2026", "1102",
            "1500.00", campos.get("fantasia", "HINOVE MATRIZ"), chave_da_nota,
            "30/07/2026", campos.get("manifestacao", "Ciência"),
            "NF-e autorizada", "NF-e Normal", "Autorizada", "COMPRA", "Saida",
            "Não se aplica", 12, "02/07/2026", "Importado", "fulano", "", "1",
            "12345678000199", "9001", "Sim", "Sim", "1",
        ]

    xml = _relatorio(tmp_path / "XML31.xlsx", colunas_de("xml"), [
        nota("1001", chave),
        nota("1002", chave[:-1] + "5", manifestacao="Desconhecida"),
        nota("1003", chave[:-1] + "7", fantasia=""),
    ])
    ce = _relatorio(tmp_path / "CE31.xlsx",
                    colunas_de("conferencia_de_entradas"), [
        [chave, "Sim", "não", "divergência", "05/07/2026", "7788",
         '<div><span>&#128994;</span></div>'],
    ])

    for nome, corpo in (("XML31.xlsx", xml), ("CE31.xlsx", ce)):
        codigo, envio = janela.pedir("/enviar", corpo=corpo,
                                     ferramenta="gerarpendentes", nome=nome)
        assert codigo == 200, envio

    codigo, dados = janela.postar("/executar", ferramenta="gerarpendentes")
    assert codigo == 200, dados

    fichas = {f["rotulo"]: f["valor"] for f in dados["fichas"]}
    assert fichas["Pendentes (áreas)"] == "1"
    assert dados["planilha"].startswith("Pendentes")
    assert "Descartados antes do roteamento" in [l["titulo"]
                                                 for l in dados["listas"]]

    codigo, corpo = janela.pedir("/baixar")
    assert codigo == 200
    livro = openpyxl.load_workbook(io.BytesIO(corpo))
    assert livro.sheetnames == ["Pendentes", "CTe", "Manifestados",
                                "Entradas 3os", "Lançados",
                                "PENDENTES FIS-FAT", "Descartados"]
    assert livro["Pendentes"].max_column == 39
    assert livro["Descartados"].max_row == 2


# -- anexar aos poucos, vendo o que já veio e o que falta --------------------

def test_quem_reconhece_os_relatorios_declara_que_confere(janela):
    _, dados = janela.pedir("/ferramentas")
    por_id = {f["id"]: f for f in dados["ferramentas"]}
    assert por_id["gerarservpend"]["confere"] is True
    assert por_id["gerarpendentes"]["confere"] is True
    assert por_id["dixml"]["confere"] is False


def _estados(quadro) -> dict:
    return {d["id"]: d["estado"] for d in quadro["documentos"]}


def test_antes_de_anexar_o_quadro_lista_o_que_a_ferramenta_espera(janela):
    codigo, quadro = janela.pedir("/conferir", ferramenta="gerarservpend")
    assert codigo == 200, quadro
    assert _estados(quadro) == {
        "asis": "falta", "portal_de_compras": "falta",
        "conferencia_de_servicos": "opcional",
        "semana_anterior_servicos": "opcional",
    }
    assert quadro["pronta"] is False
    assert quadro["pendencias"]


def test_anexar_aos_poucos_marca_cada_relatorio_e_libera_no_fim(
        janela, tmp_path):
    """O defeito que o time pediu para corrigir: anexar parcialmente."""
    asis, portal = _relatorios_de_servicos(tmp_path)

    janela.pedir("/enviar", corpo=asis, ferramenta="gerarservpend",
                 nome="ASIS.xlsx")
    _, quadro = janela.pedir("/conferir", ferramenta="gerarservpend")
    assert _estados(quadro)["asis"] == "ok"
    assert _estados(quadro)["portal_de_compras"] == "falta"
    assert quadro["documentos"][0]["anexos"][0]["nome"] == "ASIS.xlsx"
    assert any("Portal de Compras" in p for p in quadro["pendencias"])

    # Gerar antes da hora diz o que falta — e não joga fora o que já veio.
    codigo, dados = janela.postar("/executar", ferramenta="gerarservpend")
    assert codigo == 400
    assert "Portal de Compras" in dados["erro"]
    assert len(janela.sessao.recebidos_de("gerarservpend")) == 1

    janela.pedir("/enviar", corpo=portal, ferramenta="gerarservpend",
                 nome="PC27.xlsx")
    _, quadro = janela.pedir("/conferir", ferramenta="gerarservpend")
    assert quadro["pronta"] is True


def test_o_mesmo_relatorio_duas_vezes_trava_ate_tirar_um(janela, tmp_path):
    asis, portal = _relatorios_de_servicos(tmp_path)
    for nome, corpo in (("ASIS.xlsx", asis), ("PC27.xlsx", portal),
                        ("PC27-copia.xlsx", portal)):
        janela.pedir("/enviar", corpo=corpo, ferramenta="gerarservpend",
                     nome=nome)

    _, quadro = janela.pedir("/conferir", ferramenta="gerarservpend")
    assert _estados(quadro)["portal_de_compras"] == "repetido"
    assert quadro["pronta"] is False

    codigo, _ = janela.pedir("/remover", ferramenta="gerarservpend", indice=2)
    assert codigo == 200
    _, quadro = janela.pedir("/conferir", ferramenta="gerarservpend")
    assert _estados(quadro)["portal_de_compras"] == "ok"
    assert quadro["pronta"] is True


def test_remover_e_anexar_de_novo_nao_troca_um_arquivo_pelo_outro(
        janela, tmp_path):
    """A subpasta de cada anexo não pode ser reaproveitada pelo próximo."""
    asis, portal = _relatorios_de_servicos(tmp_path)
    janela.pedir("/enviar", corpo=asis, ferramenta="gerarservpend",
                 nome="ASIS.xlsx")
    janela.pedir("/enviar", corpo=portal, ferramenta="gerarservpend",
                 nome="PC27.xlsx")
    janela.pedir("/remover", ferramenta="gerarservpend", indice=0)
    janela.pedir("/enviar", corpo=asis, ferramenta="gerarservpend",
                 nome="ASIS.xlsx")

    recebidos = janela.sessao.recebidos_de("gerarservpend")
    assert sorted(c.name for c in recebidos) == ["ASIS.xlsx", "PC27.xlsx"]
    assert all(c.is_file() for c in recebidos)


def test_remover_o_que_nao_existe_avisa(janela):
    codigo, dados = janela.pedir("/remover", ferramenta="gerarservpend",
                                 indice=7)
    assert codigo == 400
    assert dados["erro"]
