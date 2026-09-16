"""A planilha de saída: uma aba por tipo de documento.

O que a formatação resolve, e por que ela não é enfeite:

* **Chave, CNPJ e número da nota em formato texto (`@`).** São códigos de 44 e
  14 dígitos. Em coluna numérica o Excel os exibe como `3,52604E+43`, perde os
  zeros à esquerda e faz o PROCV falhar contra qualquer outra base. O valor já
  chega como texto; o `@` é a defesa para quando alguém editar a planilha.
* **Data como data.** Permite filtrar por período e agrupar por mês.
* **Aptos Narrow 10.** Cabe mais coluna na tela, que é o que se faz o dia
  inteiro numa planilha de 150 colunas.

O pandas escrevia esta planilha antes. Ele saiu porque tem extensão compilada:
não poderia viajar embarcado em `vendor/`, e sem ele instalado a ferramenta não
abriria na máquina do time fiscal. O openpyxl faz o mesmo trabalho aqui.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

FONTE_PADRAO = Font(name="Aptos Narrow", size=10)
FONTE_CABECALHO = Font(name="Aptos Narrow", size=10, bold=True)

FORMATO_TEXTO = "@"
FORMATO_DATA = "DD/MM/YYYY"

#: Colunas que precisam continuar texto, pelo nome exato.
NOMES_TEXTO = {"Chave", "Chave CT-e", "Emitente Doc", "Destinatario Doc", "Numero NF"}

#: E pelo fim do nome — pega o que vem do achatamento, em qualquer caminho.
SUFIXOS_TEXTO = ("CNPJ", "CPF", "chNFe", "chCTe", "refNFe", "refCTe", "nProt", "chave")


def coluna_e_texto(nome: str) -> bool:
    if nome in NOMES_TEXTO:
        return True
    ultimo = nome.split("/")[-1]
    return any(
        ultimo.lower() == sufixo.lower() or ultimo.endswith(sufixo)
        for sufixo in SUFIXOS_TEXTO
    )


def _formato_de_cada_coluna(
    colunas: list[str], colunas_de_data: list[str]
) -> list[str | None]:
    de_data = set(colunas_de_data)
    formatos: list[str | None] = []
    for coluna in colunas:
        if coluna in de_data:
            formatos.append(FORMATO_DATA)
        elif coluna_e_texto(coluna):
            formatos.append(FORMATO_TEXTO)
        else:
            formatos.append(None)
    return formatos


def _escrever_aba(planilha, colunas, linhas, colunas_de_data) -> None:
    if not colunas:
        return
    formatos = _formato_de_cada_coluna(colunas, colunas_de_data)

    for indice, nome in enumerate(colunas, start=1):
        celula = planilha.cell(row=1, column=indice, value=nome)
        celula.font = FONTE_CABECALHO

    for numero_da_linha, linha in enumerate(linhas, start=2):
        for indice, coluna in enumerate(colunas, start=1):
            conteudo: Any = linha.get(coluna, "")
            if conteudo == "" or conteudo is None:
                continue                    # célula vazia não precisa de estilo
            celula = planilha.cell(row=numero_da_linha, column=indice, value=conteudo)
            celula.font = FONTE_PADRAO
            if formatos[indice - 1]:
                celula.number_format = formatos[indice - 1]


def escrever(destino: Path, abas: list[tuple[str, list[str], list[dict], list[str]]]) -> Path:
    """Grava o `.xlsx`. Cada aba é `(nome, colunas, linhas, colunas_de_data)`.

    Abas sem nenhuma linha continuam sendo criadas, vazias: a planilha tem
    sempre a mesma cara, e a ausência de CT-e no lote fica visível.
    """
    livro = Workbook()
    livro.remove(livro.active)
    for nome, colunas, linhas, colunas_de_data in abas:
        _escrever_aba(livro.create_sheet(nome), colunas, linhas, colunas_de_data)
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino
