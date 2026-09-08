"""A planilha gravada: o que o time fiscal abre no Excel."""
from __future__ import annotations

from datetime import date

import openpyxl
from conftest import CHAVE_NFE, nfe, zipar

from dixml.extracao import escrever, extrair, nome_sugerido
from dixml.planilha import coluna_e_texto


def _abrir(caminho):
    return openpyxl.load_workbook(caminho)


def test_a_planilha_tem_sempre_as_duas_abas(tmp_path, lote):
    """Sem CT-e no lote, a aba continua lá — a ausência fica visível."""
    destino = escrever(extrair([lote]), tmp_path / "saida.xlsx")
    livro = _abrir(destino)
    assert livro.sheetnames == ["NFe", "CTe"]


def test_a_aba_de_nfe_traz_cabecalho_e_uma_linha_por_item(tmp_path, lote):
    destino = escrever(extrair([lote]), tmp_path / "saida.xlsx")
    aba = _abrir(destino)["NFe"]
    cabecalho = [c.value for c in aba[1]]
    assert cabecalho[:4] == ["Arquivo", "Chave", "Numero NF", "Serie"]
    assert aba.max_row == 5, "quatro itens mais o cabeçalho"


def test_a_chave_chega_ao_excel_como_texto(tmp_path, lote):
    """Em coluna numérica ela viraria 3,52604E+43 e perderia o PROCV."""
    destino = escrever(extrair([lote]), tmp_path / "saida.xlsx")
    aba = _abrir(destino)["NFe"]
    coluna = [c.value for c in aba[1]].index("Chave") + 1
    celula = aba.cell(row=2, column=coluna)
    assert celula.value == CHAVE_NFE
    assert celula.number_format == "@"


def test_a_data_chega_como_data_e_e_exibida_no_formato_brasileiro(tmp_path, lote):
    destino = escrever(extrair([lote]), tmp_path / "saida.xlsx")
    aba = _abrir(destino)["NFe"]
    coluna = [c.value for c in aba[1]].index("Emissao Data") + 1
    celula = aba.cell(row=2, column=coluna)
    assert celula.value.date() == date(2026, 6, 1)
    assert celula.number_format == "DD/MM/YYYY"


def test_o_valor_chega_como_numero(tmp_path, lote):
    destino = escrever(extrair([lote]), tmp_path / "saida.xlsx")
    aba = _abrir(destino)["NFe"]
    coluna = [c.value for c in aba[1]].index("Valor Item") + 1
    assert aba.cell(row=2, column=coluna).value == 255.00


def test_a_fonte_e_a_mesma_da_planilha_que_o_time_ja_usa(tmp_path, lote):
    destino = escrever(extrair([lote]), tmp_path / "saida.xlsx")
    aba = _abrir(destino)["NFe"]
    assert aba.cell(row=1, column=1).font.bold
    assert aba.cell(row=2, column=1).font.name == "Aptos Narrow"
    assert aba.cell(row=2, column=1).font.sz == 10


def test_lote_vazio_gera_planilha_em_vez_de_erro(tmp_path):
    vazio = zipar(tmp_path / "vazio.zip", {"leiame.txt": "sem xml aqui"})
    destino = escrever(extrair([vazio]), tmp_path / "saida.xlsx")
    livro = _abrir(destino)
    assert livro.sheetnames == ["NFe", "CTe"]
    assert livro["NFe"].max_row == 1


def test_o_nome_sugerido_carrega_a_hora_para_nao_sobrescrever():
    from datetime import datetime

    assert nome_sugerido(datetime(2026, 9, 8, 14, 5)) == "DiXML_2026-09-08_1405.xlsx"


def test_colunas_de_documento_sao_reconhecidas_em_qualquer_caminho():
    assert coluna_e_texto("Chave")
    assert coluna_e_texto("emit/CNPJ")
    assert coluna_e_texto("protocolo/nProt")
    assert coluna_e_texto("infNFe/refNFe")
    assert not coluna_e_texto("Valor NF")
    assert not coluna_e_texto("vPrest/vTPrest")


def test_o_cte_grava_as_colunas_que_o_lote_trouxe(tmp_path, lote):
    destino = escrever(extrair([lote]), tmp_path / "saida.xlsx")
    aba = _abrir(destino)["CTe"]
    cabecalho = [c.value for c in aba[1]]
    assert cabecalho[:3] == ["Arquivo", "Chave CT-e", "Versao"]
    assert "vPrest/vTPrest" in cabecalho
    assert aba.max_row == 2


def test_nota_sem_item_nao_quebra_a_gravacao(tmp_path):
    lote = zipar(tmp_path / "sem_item.zip", {"nota.xml": nfe(itens=" ")})
    r = extrair([lote])
    assert len(r.nfe) == 0
    destino = escrever(r, tmp_path / "saida.xlsx")
    assert _abrir(destino)["NFe"].max_row == 1
