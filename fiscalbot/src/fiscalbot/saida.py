"""A planilha que o time fiscal abre: as abas `Relatório` e `Resumo`.

A ordem das 25 primeiras colunas não é estética — é a ordem de conferência.
Primeiro o que identifica o documento, depois o veredito e as seis ocorrências,
depois o que explica o enquadramento (CFOP, CST, carga, UFs) e por último os
valores. As colunas restantes do relatório original seguem atrás, na ordem em
que vieram, para nada se perder.

Três decisões de formatação que não são enfeite:

* **Carga efetiva é fórmula**, não valor gravado. Quem confere quer ver a conta
  (`ICMS ÷ valor contábil`) e poder mexer no numerador para simular.
* **Coluna com "chave" no nome vira texto.** Chave de acesso tem 44 dígitos: em
  coluna numérica o Excel exibe `3,52604E+43` e o PROCV para de casar.
* **Cor por grupo** — azul para a auditoria, amarelo para o que foi derivado,
  bege para o que veio do relatório. A pessoa sabe de onde cada coluna saiu.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .auditoria import ADVERTENCIA, CONFORME, VALIDACAO_MANUAL, Achado
from .leitura import MapaDeColunas

#: As 25 colunas fixas. `origem` é o campo do relatório, ou `None` quando a
#: coluna é produzida pela auditoria.
COLUNAS_FIXAS: tuple[tuple[str, str | None], ...] = (
    ("Nro unico Nota", "NUNOTA"),
    ("Descricao (Tipo de Operacao)", "DESCTIPO"),
    ("Descricao Parceiro/Empresa", "PARCNOME"),
    ("Nome Fantasia (Empresa)", "NOMEFANT"),
    ("Status", None),
    ("Operacao", None),
    ("Auditoria", None),
    ("! CST", None),
    ("! Valor ICMS", None),
    ("! Produto", None),
    ("! Aliquota", None),
    ("! Carga Efetiva", None),
    ("! Outros", None),
    ("CFOP", "CFOP"),
    ("Descricao da CFOP", "DESCCFOP"),
    ("Cod. da Situacao Tributaria", "CST"),
    ("Carga Efetiva", None),
    ("UF de Origem", "UFO"),
    ("UF de Destino", "UFD"),
    ("Oper. INTRA/INTER", None),
    ("Entrada/Saida", "ES"),
    ("Especie do documento", "ESP"),
    ("Vlr. contabil", "VCONT"),
    ("Aliquota ICMS", "ALIQ"),
    ("Vlr. do ICMS", "ICMS"),
)

COL_STATUS = 5
COL_CARGA = 17          # a fórmula
COL_TIPO = 20
COL_VCONT = 23
COL_ICMS = 25

FONTE = Font(name="Aptos Narrow", size=12)
FONTE_CABECALHO = Font(name="Aptos Narrow", size=12, bold=True)
CENTRO = Alignment(horizontal="center", vertical="center")
_PONTILHADA = Side(style="dotted", color="FF969696")
BORDA = Border(left=_PONTILHADA, right=_PONTILHADA,
               top=_PONTILHADA, bottom=_PONTILHADA)

AZUL = PatternFill("solid", fgColor="FFDDEBF7")         # auditoria
AZUL_FORTE = PatternFill("solid", fgColor="FFBDD7EE")
AMARELO = PatternFill("solid", fgColor="FFFFF2CC")      # derivadas
AMARELO_FORTE = PatternFill("solid", fgColor="FFFFE699")
BEGE = PatternFill("solid", fgColor="FFF5EED7")         # vindas do relatório

COR_DO_STATUS = {
    ADVERTENCIA: Font(name="Aptos Narrow", size=12, bold=True, color="FFC00000"),
    VALIDACAO_MANUAL: Font(name="Aptos Narrow", size=12, color="FFBF8F00"),
    CONFORME: Font(name="Aptos Narrow", size=12, color="FF006100"),
}

COLUNAS_DE_AUDITORIA = range(5, 14)                     # Status .. ! Outros
COLUNAS_DERIVADAS = (COL_CARGA, COL_TIPO)


def _valores_da_linha(achado: Achado) -> dict[int, Any]:
    """O que a auditoria põe em cada coluna fixa."""
    o = achado.ocorrencias
    return {
        5: achado.status, 6: achado.operacao, 7: achado.auditoria,
        8: o.cst, 9: o.icms, 10: o.produto,
        11: o.aliquota, 12: o.carga, 13: o.outros,
        COL_CARGA: None,                                # fórmula, escrita depois
        COL_TIPO: achado.tipo_operacao,
    }


def montar_relatorio(livro: Workbook, linhas: list[list[Any]], cabecalho: int,
                     ultima: int, mapa: MapaDeColunas,
                     achados: list[Achado]) -> None:
    """A aba `Relatório`: 25 colunas fixas e o restante do original atrás."""
    aba = livro.create_sheet("Relatório")
    original = linhas[cabecalho]
    n_colunas_originais = len(original)

    usadas = {mapa.posicoes[campo] for _, campo in COLUNAS_FIXAS
              if campo and mapa.posicoes.get(campo, -1) >= 0}
    restantes = [c for c in range(n_colunas_originais) if c not in usadas]
    total = len(COLUNAS_FIXAS) + len(restantes)

    # --- cabeçalho: a coluna que veio do relatório mantém o nome original
    for i, (rotulo, campo) in enumerate(COLUNAS_FIXAS, start=1):
        de_origem = mapa.posicoes.get(campo, -1) if campo else -1
        aba.cell(1, i, original[de_origem] if de_origem >= 0 else rotulo)
    for k, c in enumerate(restantes, start=len(COLUNAS_FIXAS) + 1):
        aba.cell(1, k, original[c])

    # --- corpo
    for r, indice in enumerate(range(cabecalho + 1, ultima + 1), start=2):
        origem_linha = linhas[indice]
        da_auditoria = _valores_da_linha(achados[r - 2])
        for i, (_, campo) in enumerate(COLUNAS_FIXAS, start=1):
            if i in da_auditoria:
                aba.cell(r, i, da_auditoria[i])
                continue
            de_origem = mapa.posicoes.get(campo, -1) if campo else -1
            aba.cell(r, i, origem_linha[de_origem] if de_origem >= 0 else "")
        for k, c in enumerate(restantes, start=len(COLUNAS_FIXAS) + 1):
            aba.cell(r, k, origem_linha[c] if c < len(origem_linha) else "")

    ultima_linha = len(achados) + 1
    _carga_em_formula(aba, ultima_linha)
    _chaves_como_texto(aba, ultima_linha, total)
    _formatar(aba, ultima_linha, total)


def _carga_em_formula(aba, ultima_linha: int) -> None:
    """`ICMS ÷ valor contábil`, para a conta ficar à vista."""
    icms = get_column_letter(COL_ICMS)
    contabil = get_column_letter(COL_VCONT)
    for r in range(2, ultima_linha + 1):
        celula = aba.cell(r, COL_CARGA)
        celula.value = f'=IFERROR({icms}{r}/{contabil}{r},"")'
        celula.number_format = "0.00%"


def _chaves_como_texto(aba, ultima_linha: int, total: int) -> None:
    """Chave de acesso em coluna numérica vira notação científica."""
    for c in range(1, total + 1):
        nome = str(aba.cell(1, c).value or "").lower()
        if "chave" not in nome:
            continue
        for r in range(2, ultima_linha + 1):
            celula = aba.cell(r, c)
            celula.number_format = "@"
            if celula.value is not None:
                celula.value = str(celula.value).strip()


def _formatar(aba, ultima_linha: int, total: int) -> None:
    de_auditoria = set(COLUNAS_DE_AUDITORIA)
    derivadas = set(COLUNAS_DERIVADAS)

    for linha in aba.iter_rows(min_row=1, max_row=ultima_linha, max_col=total):
        cabecalho = linha[0].row == 1
        for celula in linha:
            celula.alignment = CENTRO
            celula.border = BORDA
            coluna = celula.column
            if cabecalho:
                celula.font = FONTE_CABECALHO
                if coluna in de_auditoria:
                    celula.fill = AZUL_FORTE
                elif coluna in derivadas:
                    celula.fill = AMARELO_FORTE
                elif coluna <= len(COLUNAS_FIXAS):
                    celula.fill = BEGE
            else:
                celula.font = FONTE
                if coluna in de_auditoria:
                    celula.fill = AZUL
                elif coluna in derivadas:
                    celula.fill = AMARELO

    for r in range(2, ultima_linha + 1):
        celula = aba.cell(r, COL_STATUS)
        fonte = COR_DO_STATUS.get(str(celula.value))
        if fonte is not None:
            celula.font = fonte

    aba.freeze_panes = "A2"
    aba.auto_filter.ref = f"A1:{get_column_letter(total)}{ultima_linha}"
    for coluna, largura in ((COL_STATUS, 12), (6, 26), (7, 48),
                            (COL_CARGA, 13), (COL_TIPO, 14)):
        aba.column_dimensions[get_column_letter(coluna)].width = largura


def _percentual(parte: int, total: int) -> str:
    valor = (parte / total * 100) if total else 0.0
    return f"{valor:.1f}".replace(".", ",") + "%"


def montar_resumo(livro: Workbook, achados: list[Achado]) -> None:
    """A aba `Resumo`: os três números, as ocorrências e o que sobrou para a mão."""
    aba = livro.create_sheet("Resumo")
    total = len(achados)
    conformes = sum(1 for a in achados if a.status == CONFORME)
    advertencias = sum(1 for a in achados if a.status == ADVERTENCIA)
    manuais = total - conformes - advertencias

    aba.cell(1, 1, "FISCALBOT - RESUMO").font = Font(
        name="Aptos Narrow", size=14, bold=True)

    for linha, (rotulo, quantidade) in enumerate(
        (("Conformes", conformes), ("Advertências", advertencias),
         ("Validação manual", manuais)), start=3
    ):
        aba.cell(linha, 1, rotulo).font = Font(name="Aptos Narrow", size=11, bold=True)
        aba.cell(linha, 2, quantidade)
        aba.cell(linha, 3, _percentual(quantidade, total))
    aba.cell(7, 1, "Total").font = Font(name="Aptos Narrow", size=11, bold=True)
    aba.cell(7, 2, total)

    rotulos = ("CST", "Valor ICMS", "Produto", "Alíquota", "Carga Efetiva", "Outros")
    aba.cell(10, 1, "Ocorrência").font = Font(name="Aptos Narrow", size=11, bold=True)
    aba.cell(10, 2, "Qtde").font = Font(name="Aptos Narrow", size=11, bold=True)
    for k, rotulo in enumerate(rotulos):
        quantos = sum(1 for a in achados if a.ocorrencias.como_lista()[k])
        aba.cell(11 + k, 1, rotulo)
        aba.cell(11 + k, 2, quantos)

    aba.cell(18, 1, "Validação manual").font = Font(name="Aptos Narrow", size=11, bold=True)
    aba.cell(18, 2, "Qtde").font = Font(name="Aptos Narrow", size=11, bold=True)
    por_operacao: dict[str, int] = {}
    for achado in achados:
        if achado.status != VALIDACAO_MANUAL:
            continue
        chave = achado.operacao or "(sem operacao)"
        por_operacao[chave] = por_operacao.get(chave, 0) + 1
    ordenadas = sorted(por_operacao.items(), key=lambda item: -item[1])
    for k, (operacao, quantos) in enumerate(ordenadas, start=1):
        aba.cell(18 + k, 1, operacao)
        aba.cell(18 + k, 2, quantos)

    for coluna, largura in (("A", 42), ("B", 10), ("C", 10)):
        aba.column_dimensions[coluna].width = largura


def escrever(destino: Path, linhas: list[list[Any]], cabecalho: int, ultima: int,
             mapa: MapaDeColunas, achados: list[Achado]) -> Path:
    """Grava a planilha auditada. O arquivo de entrada não é tocado."""
    livro = Workbook()
    livro.remove(livro.active)
    montar_relatorio(livro, linhas, cabecalho, ultima, mapa, achados)
    montar_resumo(livro, achados)
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino
