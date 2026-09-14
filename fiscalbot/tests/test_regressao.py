"""A prova de que o motor Python faz o mesmo que a macro VBA.

Este é o teste que autoriza aposentar o `.xlsm`. Ele pega um relatório real já
auditado pela macro — que traz as abas `original` (a entrada) e `Relatório` (a
saída dela) — roda o motor sobre a mesma entrada e compara **linha a linha**:
status, operação, texto de auditoria, tipo da operação e as seis ocorrências.

O critério é divergência zero. Não é "quase igual": qualquer diferença é uma
nota auditada de outro jeito, e o Livro Fiscal que o Apurabot consome sai daqui.

O arquivo tem dado fiscal real e não é versionado. Sem ele, o teste é pulado —
mas então ninguém pode dizer que o porte está provado.
"""
from __future__ import annotations

import pytest
from conftest import COLUNAS  # noqa: F401  (mantém o sys.path do pacote)

from fiscalbot import leitura, planilha
from fiscalbot.auditoria import ADVERTENCIA, CONFORME, VALIDACAO_MANUAL, auditar_linha
from fiscalbot.base import carregar

#: Onde cada campo comparado está na aba `Relatório` (índice de coluna, base 0).
NA_SAIDA_DA_MACRO = {
    "status": 4, "operacao": 5, "auditoria": 6, "cst": 7, "icms": 8,
    "produto": 9, "aliquota": 10, "carga": 11, "outros": 12, "tipo": 19,
}


def _com_os_parceiros_locais(base):
    """Empresta as listas de parceiros da base viva — elas não são versionadas."""
    viva = carregar()
    if not viva.parceiros_simples_nacional and not viva.parceiros_cavaco:
        pytest.skip(
            "a base local não tem as listas de parceiros. Elas são dado da "
            "empresa, não vêm no repositório e precisam ser cadastradas na "
            "tela do Fiscalbot — sem elas as camadas de frete SN e de cavaco "
            "não rodam, e a comparação não é justa."
        )
    base.parceiros_simples_nacional = viva.parceiros_simples_nacional
    base.parceiros_cavaco = viva.parceiros_cavaco
    return base


@pytest.fixture(scope="session")
def base_viva(regras_da_macro):
    """As regras da macro, com os parceiros desta máquina."""
    return _com_os_parceiros_locais(regras_da_macro)


@pytest.fixture(scope="session")
def confronto(padrao_ouro, base_viva):
    """Roda o motor sobre a entrada da macro e junta as duas saídas."""
    import xlrd

    livro = xlrd.open_workbook(str(padrao_ouro))
    if "Relatório" not in livro.sheet_names():
        pytest.skip(
            f"{padrao_ouro.name} não tem a aba 'Relatório': é um relatório "
            "cru, não a saída já auditada pela macro."
        )
    entrada = livro.sheet_by_name("original")
    da_macro = livro.sheet_by_name("Relatório")

    linhas = [[entrada.cell_value(r, c) for c in range(entrada.ncols)]
              for r in range(entrada.nrows)]
    cabecalho = leitura.localizar_cabecalho(linhas)
    mapa = leitura.mapear(linhas[cabecalho], base_viva.parametros)

    ultima = cabecalho
    for i in range(cabecalho + 1, len(linhas)):
        if str(linhas[i][mapa.posicoes["CFOP"]]).strip():
            ultima = i

    achados = [auditar_linha(linhas[i], mapa, base_viva)
               for i in range(cabecalho + 1, ultima + 1)]
    return achados, da_macro


def _do_python(achado) -> dict[str, str]:
    o = achado.ocorrencias
    return {"status": achado.status, "operacao": achado.operacao,
            "auditoria": achado.auditoria, "tipo": achado.tipo_operacao,
            "cst": o.cst, "icms": o.icms, "produto": o.produto,
            "aliquota": o.aliquota, "carga": o.carga, "outros": o.outros}


def test_o_motor_audita_a_mesma_quantidade_de_registros(confronto):
    achados, da_macro = confronto
    assert len(achados) == da_macro.nrows - 1


def test_divergencia_zero_contra_a_macro(confronto):
    """O teste que autoriza aposentar o .xlsm."""
    achados, da_macro = confronto
    divergencias = []
    for i, achado in enumerate(achados, start=1):
        obtido = _do_python(achado)
        for campo, coluna in NA_SAIDA_DA_MACRO.items():
            esperado = str(da_macro.cell_value(i, coluna)).strip()
            atual = str(obtido[campo]).strip()
            if esperado != atual:
                divergencias.append(
                    f"linha {i}, {campo}: macro={esperado!r} python={atual!r}")
    assert not divergencias, (
        f"{len(divergencias)} divergência(s). As 5 primeiras:\n  "
        + "\n  ".join(divergencias[:5])
    )


def test_os_totais_do_resumo_batem(confronto, padrao_ouro):
    """Os três números que o time fiscal olha primeiro."""
    import xlrd

    achados, _ = confronto
    resumo = xlrd.open_workbook(str(padrao_ouro)).sheet_by_name("Resumo")
    # Só o bloco de KPIs (linhas 3 a 7). Mais abaixo o rótulo "Validação
    # manual" volta como título da lista por operação, com "Qtde" ao lado.
    da_macro = {str(resumo.cell_value(r, 0)).strip(): resumo.cell_value(r, 1)
                for r in range(2, 7)}
    do_python = {
        "Conformes": sum(1 for a in achados if a.status == CONFORME),
        "Advertências": sum(1 for a in achados if a.status == ADVERTENCIA),
        "Validação manual": sum(1 for a in achados if a.status == VALIDACAO_MANUAL),
        "Total": len(achados),
    }
    for rotulo, quantos in do_python.items():
        assert quantos == da_macro.get(rotulo), rotulo


#: O que é veredicto — o que muda a vida de quem confere. O texto da coluna
#: `Auditoria` fica de fora: ele descreve o que a regra pede, então muda
#: legitimamente quando a regra é corrigida.
VEREDICTO = ("status", "operacao", "tipo", "cst", "icms", "produto",
             "aliquota", "carga", "outros")


def test_corrigir_a_regra_e11_nao_mudou_nenhum_veredicto(padrao_ouro,
                                                          base_de_fabrica):
    """A base de hoje não é mais a da macro. Os veredictos, sim.

    A regra E11 tinha `TABELAUF` na coluna de ICMS, onde o operador é inerte,
    e conferia a alíquota contra a lista literal `7;12`. A correção moveu o
    operador para a coluna de alíquota, que passa a vir da matriz origem ×
    destino. Nos 730 registros que casam com a E11 nesta competência, as duas
    formas dão o mesmo resultado — a matriz devolve exatamente 7 ou 12 para
    todos os pares presentes. O que muda é o texto que explica a regra.

    Se um dia essa correção mudar veredicto, é aqui que se descobre.
    """
    import xlrd

    base = _com_os_parceiros_locais(base_de_fabrica)
    livro = xlrd.open_workbook(str(padrao_ouro))
    if "Relatório" not in livro.sheet_names():
        pytest.skip("o arquivo não traz a saída da macro")
    entrada = livro.sheet_by_name("original")
    da_macro = livro.sheet_by_name("Relatório")

    linhas = [[entrada.cell_value(r, c) for c in range(entrada.ncols)]
              for r in range(entrada.nrows)]
    cabecalho = leitura.localizar_cabecalho(linhas)
    mapa = leitura.mapear(linhas[cabecalho], base.parametros)
    ultima = cabecalho
    for i in range(cabecalho + 1, len(linhas)):
        if str(linhas[i][mapa.posicoes["CFOP"]]).strip():
            ultima = i

    divergencias = []
    for k, i in enumerate(range(cabecalho + 1, ultima + 1), start=1):
        obtido = _do_python(auditar_linha(linhas[i], mapa, base))
        for campo in VEREDICTO:
            esperado = str(da_macro.cell_value(k, NA_SAIDA_DA_MACRO[campo])).strip()
            if esperado != str(obtido[campo]).strip():
                divergencias.append(
                    f"linha {k}, {campo}: macro={esperado!r} python={obtido[campo]!r}")
    assert not divergencias, (
        f"a base de hoje mudou {len(divergencias)} veredicto(s) em relação à "
        f"macro. As 5 primeiras:\n  " + "\n  ".join(divergencias[:5])
    )


def test_a_planilha_gerada_repete_as_colunas_da_macro(confronto, padrao_ouro,
                                                      base_viva, tmp_path):
    """A ordem das 25 primeiras colunas é a ordem de conferência do time."""
    import openpyxl

    from fiscalbot.execucao import auditar, escrever

    resultado = auditar(padrao_ouro, base_viva)
    destino = escrever(resultado, tmp_path / "auditado.xlsx")
    gerada = openpyxl.load_workbook(destino)["Relatório"]

    import xlrd
    da_macro = xlrd.open_workbook(str(padrao_ouro)).sheet_by_name("Relatório")

    for c in range(1, 26):
        assert str(gerada.cell(1, c).value) == str(da_macro.cell_value(0, c - 1)), \
            f"cabeçalho da coluna {c}"
    assert gerada.max_column == da_macro.ncols
