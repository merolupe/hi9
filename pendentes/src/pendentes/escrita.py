"""A aba formatada — com o formato aplicado **antes** de qualquer escrita.

Esta ordem não é preferência de estilo: é a correção de um defeito que os dois
módulos VBA já carregam, e que uma reimplementação ingênua refaria.

No Excel, `Columns(n).Insert Shift:=xlToRight` faz a coluna nova **herdar o
formato da vizinha à esquerda**. Como as seis colunas de conferência nascem
logo depois de `Chave Acesso`, que é Texto, elas nasciam Texto — e gravar uma
data numa célula formatada como Texto a converte em string literal no formato
regional do Windows. Formatar depois não reverte, e não levanta exceção
nenhuma. O mesmo vale para a chave de acesso: 44 dígitos numa coluna numérica
viram `3,52604E+43`, e o PROCX para de casar.

Aqui não há Excel nem `Insert`, mas a doutrina fica: a aba **declara** o
formato de cada coluna, e só então recebe a primeira linha. O tipo do valor é
coagido junto — texto é escrito como texto, valor como número, data como data.
É a diferença entre a planilha que parece certa e a planilha que **é** certa.

Do lado do que desapareceu: `AutoFit` não existe em openpyxl. A largura é
calculada do conteúdo e tem teto, que é o que o módulo de serviços já fazia à
mão depois do AutoFit (`If ColumnWidth > 60 Then ColumnWidth = 60`).
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .texto import texto_de
from .valores import data_br, numero_br

#: O formato de célula de cada tipo de coluna.
FORMATOS = {
    "texto": "@",
    "valor": "#,##0.00",
    "data": "DD/MM/YYYY",
    # A coluna que ninguém sabe se é data ou contagem de dias: recebe o mesmo
    # formato de data, e o valor vai **como veio**. Ver `valor_formatado`.
    "data_ou_contagem": "DD/MM/YYYY",
    "geral": "General",
}

#: Teto e piso de largura de coluna, em caracteres.
LARGURA_MAXIMA = 60
LARGURA_MINIMA = 8
#: Quantas linhas são olhadas para calcular a largura. Olhar todas custa caro
#: e não muda o resultado.
LINHAS_PARA_MEDIR = 200


@dataclass(frozen=True)
class Coluna:
    """Uma coluna da aba: o rótulo, o formato e, se for o caso, a largura fixa."""

    rotulo: str
    formato: str = "geral"
    largura: int | None = None


@dataclass(frozen=True)
class Estilo:
    """Como a aba se parece. O padrão é o do módulo de serviços."""

    fonte: str = "Aptos Narrow"
    tamanho: int = 10
    cor_do_cabecalho: str = "FFD9E2F3"
    centralizar: bool = False
    bordas: bool = False


_PONTILHADA = Side(style="dotted", color="FF969696")
BORDA = Border(left=_PONTILHADA, right=_PONTILHADA,
               top=_PONTILHADA, bottom=_PONTILHADA)
CENTRO = Alignment(horizontal="center", vertical="center")


def novo_livro() -> Workbook:
    """Um `.xlsx` sem a aba vazia que o openpyxl cria sozinho."""
    livro = Workbook()
    livro.remove(livro.active)
    return livro


def preparar_aba(livro: Workbook, nome: str, colunas: Sequence[Coluna], *,
                 oculta: bool = False, estilo: Estilo | None = None):
    """Cria a aba, **formata as colunas** e escreve o cabeçalho.

    Nesta ordem. Quando esta função retorna, a aba ainda não tem uma linha de
    dado e já sabe o formato de cada coluna.
    """
    estilo = estilo or Estilo()
    aba = livro.create_sheet(nome)

    # 1. o formato, antes de tudo
    for i, coluna in enumerate(colunas, start=1):
        letra = get_column_letter(i)
        aba.column_dimensions[letra].number_format = FORMATOS.get(
            coluna.formato, FORMATOS["geral"])

    # 2. só então o cabeçalho
    negrito = Font(name=estilo.fonte, size=estilo.tamanho, bold=True)
    fundo = PatternFill("solid", fgColor=estilo.cor_do_cabecalho)
    for i, coluna in enumerate(colunas, start=1):
        celula = aba.cell(1, i, coluna.rotulo)
        celula.font = negrito
        celula.fill = fundo
        if estilo.centralizar:
            celula.alignment = CENTRO
        if estilo.bordas:
            celula.border = BORDA

    if oculta:
        aba.sheet_state = "hidden"
    return aba


def valor_formatado(valor: Any, formato: str) -> Any:
    """O valor no tipo que a coluna declarou.

    Vazio continua vazio: coluna de valor sem valor não vira zero, porque
    zero é uma afirmação e ausência não é.

    `data_ou_contagem` é a exceção, e existe por uma coluna só: o VBA escreve
    `Dias Emissão Doc` **sem conversão** e formata a coluna como `DD/MM/YYYY`.
    Se o valor for contagem de dias, o Excel exibe uma data de 1900 — defeito
    visível, e preservado até a pendência nº 2 ser respondida. Converter aqui
    trocaria o defeito por outro: o número viraria uma data de verdade, e a
    célula deixaria de ser a que a macro produz.
    """
    if formato == "texto":
        return texto_de(valor)
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return ""
    if formato == "data_ou_contagem":
        return valor
    if formato == "valor":
        return numero_br(valor)
    if formato == "data":
        convertido = data_br(valor)
        if isinstance(convertido, datetime) and not (
                convertido.hour or convertido.minute or convertido.second):
            return convertido.date()
        return convertido
    return valor


def escrever_linhas(aba, colunas: Sequence[Coluna],
                    linhas: Iterable[Sequence[Any]], *,
                    estilo: Estilo | None = None) -> int:
    """Escreve as linhas de dados. Devolve quantas escreveu."""
    estilo = estilo or Estilo()
    fonte = Font(name=estilo.fonte, size=estilo.tamanho)
    formatos = [FORMATOS.get(c.formato, FORMATOS["geral"]) for c in colunas]
    escritas = 0
    for r, linha in enumerate(linhas, start=2):
        for i, coluna in enumerate(colunas, start=1):
            bruto = linha[i - 1] if i - 1 < len(linha) else ""
            celula = aba.cell(r, i, valor_formatado(bruto, coluna.formato))
            celula.number_format = formatos[i - 1]
            celula.font = fonte
            if estilo.centralizar:
                celula.alignment = CENTRO
            if estilo.bordas:
                celula.border = BORDA
        escritas += 1
    return escritas


def _largura(coluna: Coluna, valores: Sequence[Any]) -> int:
    if coluna.largura:
        return coluna.largura
    maior = len(str(coluna.rotulo))
    for valor in valores[:LINHAS_PARA_MEDIR]:
        if isinstance(valor, (datetime, date)):
            tamanho = 10
        else:
            tamanho = len(texto_de(valor))
        maior = max(maior, tamanho)
    return max(LARGURA_MINIMA, min(LARGURA_MAXIMA, maior + 2))


def finalizar_aba(aba, colunas: Sequence[Coluna],
                  linhas: Sequence[Sequence[Any]] = ()) -> None:
    """Largura, autofiltro e congelamento — o que faz a aba ser usável."""
    for i, coluna in enumerate(colunas, start=1):
        valores = [linha[i - 1] if i - 1 < len(linha) else ""
                   for linha in linhas[:LINHAS_PARA_MEDIR]]
        aba.column_dimensions[get_column_letter(i)].width = _largura(coluna, valores)
    ultima = get_column_letter(max(1, len(colunas)))
    aba.auto_filter.ref = f"A1:{ultima}{max(1, len(linhas) + 1)}"
    aba.freeze_panes = "A2"


def escrever_aba(livro: Workbook, nome: str, colunas: Sequence[Coluna],
                 linhas: Sequence[Sequence[Any]], *, oculta: bool = False,
                 estilo: Estilo | None = None):
    """Aba completa, na ordem certa: formato, cabeçalho, dados, acabamento."""
    aba = preparar_aba(livro, nome, colunas, oculta=oculta, estilo=estilo)
    escrever_linhas(aba, colunas, linhas, estilo=estilo)
    finalizar_aba(aba, colunas, linhas)
    return aba


def ocultar(aba) -> None:
    """`xlSheetHidden` — some da vista, continua reexibível por clique direito.

    O VBA oculta as abas auxiliares e o comentário registra o motivo: elas são
    evidência para auditoria. Esconder, e não apagar, é a decisão.
    """
    aba.sheet_state = "hidden"


def salvar(livro: Workbook, caminho: Path | str) -> Path:
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    livro.save(str(caminho))
    return caminho
