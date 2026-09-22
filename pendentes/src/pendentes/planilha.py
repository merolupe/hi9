"""Leitura de planilha: `.xls`, `.xlsx` e `.xlsm` viram a mesma matriz.

O Sankhya entrega `.xls` 97-2003; o ASIS entrega `.xlsx`; alguém reabre e
salva como `.xlsm`. Os três chegam aqui e saem iguais — uma lista de abas, e
cada aba uma matriz de valores onde número é número, data é data e texto é
texto. Essa distinção importa: o CFOP vem numérico, a chave de acesso vem como
texto de 44 dígitos, e cada um é tratado de um jeito.

Duas coisas que o VBA fazia e que **desaparecem** aqui, porque não há Excel:

* `AbrirWorkbookRobusto`, com as três estratégias de abertura (reaproveitar o
  arquivo já aberto, baixar o `AutomationSecurity`, entrar na pasta e abrir
  pelo nome relativo);
* a cópia dos arquivos para `%TEMP%` e o `SetAttr` que contornavam o OneDrive
  entregando arquivo "apenas na nuvem".

A Central entrega o arquivo já gravado numa pasta temporária local. O problema
que essas duas mitigações resolviam deixa de existir.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

EXTENSOES = (".xls", ".xlsx", ".xlsm")


class PlanilhaIlegivel(Exception):
    """O arquivo não é uma planilha que este projeto leia."""


@dataclass(frozen=True)
class Aba:
    """Uma aba lida: o nome como está no arquivo e as linhas como vieram."""

    nome: str
    linhas: list[list[Any]] = field(default_factory=list)

    def __len__(self) -> int:
        return len(self.linhas)

    def linha(self, indice: int) -> list[Any]:
        return self.linhas[indice] if 0 <= indice < len(self.linhas) else []


@dataclass(frozen=True)
class Arquivo:
    """Um arquivo lido, com as abas na ordem em que estão nele."""

    caminho: Path
    abas: list[Aba] = field(default_factory=list)

    @property
    def nome(self) -> str:
        return self.caminho.name

    def aba(self, nome: str) -> Aba | None:
        """A aba pelo nome, comparado sem caixa e sem acento."""
        from .texto import chave_de_texto

        alvo = chave_de_texto(nome)
        for aba in self.abas:
            if chave_de_texto(aba.nome) == alvo:
                return aba
        return None

    @property
    def primeira(self) -> Aba:
        return self.abas[0] if self.abas else Aba("", [])


def _ler_xls(caminho: Path, limite: int | None) -> list[Aba]:
    import xlrd

    livro = xlrd.open_workbook(str(caminho), formatting_info=False)
    try:
        abas = []
        for planilha in livro.sheets():
            ultima = planilha.nrows if limite is None else min(limite, planilha.nrows)
            linhas = [
                [_valor_do_xls(planilha.cell(r, c), livro.datemode)
                 for c in range(planilha.ncols)]
                for r in range(ultima)
            ]
            abas.append(Aba(planilha.name, linhas))
        return abas
    finally:
        livro.release_resources()


def _valor_do_xls(celula: Any, datemode: int) -> Any:
    """Data do `.xls` é número com uma marca; sem a marca, vira número mesmo."""
    import xlrd

    if celula.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(celula.value, datemode)
        except (ValueError, OverflowError):              # pragma: no cover
            return celula.value
    if celula.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(celula.value)
    if celula.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if celula.ctype == xlrd.XL_CELL_ERROR:
        return ""
    return celula.value


def _ler_xlsx(caminho: Path, limite: int | None) -> list[Aba]:
    import openpyxl

    # Aberto aqui, e não pelo nome: o openpyxl recusa um `.xlsx` de verdade
    # quando ele se chama `.xls`, antes mesmo de olhar o conteúdo.
    fluxo = caminho.open("rb")
    try:
        livro = openpyxl.load_workbook(fluxo, data_only=True, read_only=True)
    except Exception:
        fluxo.close()
        raise
    try:
        abas = []
        for planilha in livro.worksheets:
            linhas = []
            for linha in planilha.iter_rows(values_only=True):
                linhas.append([c if c is not None else "" for c in linha])
                if limite is not None and len(linhas) >= limite:
                    break
            abas.append(Aba(planilha.title, linhas))
        return abas
    finally:
        livro.close()
        fluxo.close()


#: Os primeiros bytes de cada formato. O `.xlsx` é um `.zip`; o `.xls`
#: 97-2003 é um documento OLE.
_ASSINATURA_XLSX = b"PK\x03\x04"
_ASSINATURA_XLS = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _formato(caminho: Path) -> str:
    """O formato de verdade, pelo conteúdo — `""` quando não reconhece.

    A extensão mente com frequência: o arquivo da semana anterior é salvo como
    `.xlsx` e renomeado para `XMLAnterior.xls`, e o Excel abre os dois sem
    reclamar. Então quem decide o leitor é o começo do arquivo, e a extensão
    só vale quando o conteúdo não diz nada.
    """
    with caminho.open("rb") as arquivo:
        inicio = arquivo.read(8)
    if inicio.startswith(_ASSINATURA_XLSX):
        return ".xlsx"
    if inicio.startswith(_ASSINATURA_XLS):
        return ".xls"
    if inicio.lstrip(b"\xef\xbb\xbf \t\r\n").startswith(b"<"):
        return "html"
    return ""


def ler(caminho: Path | str, *, limite_de_linhas: int | None = None) -> Arquivo:
    """O arquivo inteiro, aba por aba.

    `limite_de_linhas` lê só as primeiras linhas de cada aba — é o que o
    reconhecimento de papel usa para espiar o cabeçalho sem carregar um
    relatório de sete mil linhas para descobrir que ele é o de sempre.
    """
    caminho = Path(caminho)
    if not caminho.is_file():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")
    sufixo = caminho.suffix.lower()
    try:
        if sufixo not in EXTENSOES:
            raise PlanilhaIlegivel(
                f"{caminho.name} não é uma planilha que a ferramenta leia. "
                f"Esperado {', '.join(EXTENSOES)}."
            )
        formato = _formato(caminho) or sufixo
        if formato == "html":
            raise PlanilhaIlegivel(
                f"{caminho.name} é uma página HTML com nome de planilha. "
                f"Abra no Excel e salve como .xlsx antes de anexar."
            )
        if formato == ".xls":
            abas = _ler_xls(caminho, limite_de_linhas)
        elif formato in (".xlsx", ".xlsm"):
            abas = _ler_xlsx(caminho, limite_de_linhas)
        else:
            raise PlanilhaIlegivel(
                f"{caminho.name} não é uma planilha que a ferramenta leia. "
                f"Esperado {', '.join(EXTENSOES)}."
            )
    except PlanilhaIlegivel:
        raise
    except Exception as erro:                            # noqa: BLE001
        raise PlanilhaIlegivel(
            f"Não consegui abrir {caminho.name}: {erro}"
        ) from erro
    return Arquivo(caminho, abas)
