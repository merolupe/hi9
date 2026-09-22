"""O painel desenhado: duas abas, três gráficos e nenhuma fórmula.

### Por que valores e não fórmulas

O arquivo de origem é feito de `CONT.SES` e `SOMASES` apontando para
`Pendentes!$D$2:$D$68`. Aqui as contas já foram feitas em Python, e o que vai
para a célula é o número. Três motivos, nesta ordem:

* o intervalo `$2:$68` é a semana 38 e mais nenhuma. Fórmula com intervalo
  fixo é a mesma armadilha do índice de coluna fixo que o porte tirou do VBA:
  na semana seguinte ela aponta para o lugar errado **sem errar**;
* fórmula só mostra número depois que o Excel abre e calcula. Quem recebe o
  arquivo por e-mail e olha no celular vê o painel montado;
* o valor gravado é o valor que a ferramenta afirma. Se ele estiver errado, o
  erro é reproduzível — não depende de quem abriu, com qual versão.

O que se perde é o painel se corrigir sozinho quando alguém edita uma linha
depois. É deliberado: o resumo é gerado **depois** da classificação fechar, e
editar o relatório depois disso pede rodar o resumo de novo, não confiar num
recálculo silencioso.

### O que o arquivo perde ao passar por aqui

`[FATO]` O openpyxl não preserva gráfico nem imagem das abas que ele apenas
relê. O relatório da semana, como a ferramenta o gera, não tem nenhum dos
dois — mas um gráfico que alguém tenha colado à mão em outra aba não
sobrevive. Está dito na tela, contado, em vez de sumir calado.
"""
from __future__ import annotations

from datetime import date
from typing import Sequence

from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import colunas as col
from .fontes import Pendencia
from .painel import Corte, Painel

#: A fonte do painel é a mesma do resto da planilha.
FONTE = "Aptos Narrow"
TAMANHO = 10
COR_DO_CABECALHO = "FFD9E2F3"
#: O destaque das notas acima do limite de dias — o mesmo tom de alerta que a
#: tela da Central usa para o que pede atenção.
COR_DO_DESTAQUE = "FFFFC7CE"
COR_DA_LETRA_DESTACADA = "FF9C0006"


def _titulo(aba, linha: int, coluna: int, texto: str, largura: int = 4) -> None:
    """Um título de bloco: negrito, mesclado pela largura do bloco."""
    celula = aba.cell(linha, coluna, texto)
    celula.font = Font(name=FONTE, size=TAMANHO + 1, bold=True)
    if largura > 1:
        aba.merge_cells(start_row=linha, start_column=coluna,
                        end_row=linha, end_column=coluna + largura - 1)


def _cabecalho(aba, linha: int, coluna: int, rotulos: Sequence[str]) -> None:
    negrito = Font(name=FONTE, size=TAMANHO, bold=True)
    fundo = PatternFill("solid", fgColor=COR_DO_CABECALHO)
    for i, rotulo in enumerate(rotulos):
        celula = aba.cell(linha, coluna + i, rotulo)
        celula.font = negrito
        celula.fill = fundo
        celula.alignment = Alignment(horizontal="center", vertical="center")


def _celula(aba, linha: int, coluna: int, valor, formato: str | None = None):
    celula = aba.cell(linha, coluna, valor)
    celula.font = Font(name=FONTE, size=TAMANHO)
    if formato:
        celula.number_format = formato
    return celula


# -- a aba auxiliar ---------------------------------------------------------

def escrever_auxiliar(livro, painel: Painel):
    """`_AuxResumo` — o intervalo para onde os três gráficos apontam.

    Ela existe pelo mesmo motivo que existe no arquivo de origem: gráfico de
    Excel aponta para células, não para uma lista. Nasce oculta.
    """
    aba = livro.create_sheet(col.ABA_AUXILIAR)
    aba.sheet_state = "hidden"

    # A:C — as três categorias, que é o que a pizza mostra.
    _cabecalho(aba, 1, col.BLOCO_DAS_CATEGORIAS,
               (col.TABELA_DE_CATEGORIAS[0], col.CABECALHO_DA_QUANTIDADE,
                col.CABECALHO_DO_VALOR))
    for i, linha in enumerate(painel.categorias, start=2):
        _celula(aba, i, 1, linha.categoria)
        _celula(aba, i, 2, linha.quantidade)
        _celula(aba, i, 3, linha.valor, col.FORMATO_DE_VALOR)

    # D:G e H:K — a unidade contada e a unidade somada.
    for bloco, somar in ((col.BLOCO_DA_CONTAGEM_POR_UNIDADE, False),
                         (col.BLOCO_DO_VALOR_POR_UNIDADE, True)):
        _cabecalho(aba, 1, bloco,
                   (col.CABECALHO_DA_UNIDADE, *col.CATEGORIAS_DO_GRAFICO))
        _escrever_cortes(aba, bloco, painel.unidades, somar)

    # L:O — todos os guardiões; AD:AG — só os que entram no gráfico.
    _cabecalho(aba, 1, col.BLOCO_DOS_GUARDIOES,
               (col.CABECALHO_DO_GUARDIAO, *col.CATEGORIAS_DO_GRAFICO))
    _escrever_cortes(aba, col.BLOCO_DOS_GUARDIOES, painel.guardioes, False)
    _cabecalho(aba, 1, col.BLOCO_DO_GRAFICO_DE_GUARDIOES,
               (col.CABECALHO_DO_GUARDIAO, *col.CATEGORIAS_DO_GRAFICO))
    _escrever_cortes(aba, col.BLOCO_DO_GRAFICO_DE_GUARDIOES,
                     painel.guardioes_do_grafico, False)
    return aba


def _escrever_cortes(aba, coluna: int, cortes: Sequence[Corte],
                     somar: bool) -> None:
    for i, corte in enumerate(cortes, start=2):
        _celula(aba, i, coluna, corte.nome)
        for j, categoria in enumerate(col.CATEGORIAS_DO_GRAFICO, start=1):
            valor = (corte.valor[categoria] if somar
                     else corte.quantidade[categoria])
            _celula(aba, i, coluna + j, valor,
                    col.FORMATO_DE_VALOR if somar else None)


# -- os três gráficos -------------------------------------------------------

def _pizza(auxiliar, quantas: int) -> PieChart:
    """A proporção de notas pendentes por categoria."""
    grafico = PieChart()
    grafico.height, grafico.width = 8.0, 11.5
    grafico.add_data(Reference(auxiliar, min_col=2, min_row=1,
                               max_row=1 + quantas), titles_from_data=True)
    grafico.set_categories(Reference(auxiliar, min_col=1, min_row=2,
                                     max_row=1 + quantas))
    grafico.legend.position = "b"
    return grafico


def _barras(auxiliar, coluna: int, quantas: int, *, deitada: bool,
            altura: float, largura: float) -> BarChart:
    """Uma barra empilhada por categoria — as três somam o total do corte."""
    grafico = BarChart()
    grafico.type = "bar" if deitada else "col"
    grafico.grouping = "stacked"
    grafico.overlap = 100
    grafico.height, grafico.width = altura, largura
    grafico.add_data(
        Reference(auxiliar, min_col=coluna + 1, max_col=coluna + 3,
                  min_row=1, max_row=1 + quantas),
        titles_from_data=True)
    grafico.set_categories(Reference(auxiliar, min_col=coluna, min_row=2,
                                     max_row=1 + quantas))
    grafico.legend.position = "b"
    return grafico


# -- o painel ---------------------------------------------------------------

def escrever_painel(livro, painel: Painel, *, destacar_acima_de: int,
                    posicao: int = 0):
    """`Resumo Executivo` — a aba que vai na frente do arquivo."""
    aba = livro.create_sheet(col.ABA_DO_PAINEL, posicao)
    aba.sheet_view.showGridLines = False
    for letra, largura in col.LARGURAS.items():
        aba.column_dimensions[letra].width = largura

    esquerda, direita = col.COLUNA_DA_ESQUERDA, col.COLUNA_DA_DIREITA
    _titulo(aba, col.LINHA_DO_TITULO, esquerda, col.TITULO)
    _titulo(aba, col.LINHA_DA_PIZZA, esquerda, col.TITULO_DA_PIZZA)

    _tabela_de_categorias(aba, painel, esquerda)
    _tops(aba, painel, direita, destacar_acima_de)
    _ajustes(aba, painel, destacar_acima_de)

    _titulo(aba, col.LINHA_DOS_GRAFICOS_DE_BAIXO, esquerda,
            col.TITULO_DAS_UNIDADES)
    _titulo(aba, col.LINHA_DOS_GRAFICOS_DE_BAIXO, direita,
            col.TITULO_DOS_GUARDIOES, largura=7)
    return aba


def _tabela_de_categorias(aba, painel: Painel, coluna: int) -> None:
    linha = col.LINHA_DA_TABELA
    _cabecalho(aba, linha, coluna, col.TABELA_DE_CATEGORIAS)
    negrito = Font(name=FONTE, size=TAMANHO, bold=True)
    for i, dados in enumerate([*painel.categorias, painel.total], start=1):
        if dados is None:                                # pragma: no cover
            continue
        _celula(aba, linha + i, coluna, dados.categoria)
        _celula(aba, linha + i, coluna + 1, dados.quantidade)
        _celula(aba, linha + i, coluna + 2, dados.valor, col.FORMATO_DE_VALOR)
        _celula(aba, linha + i, coluna + 3,
                "" if dados.media_dias is None else dados.media_dias,
                col.FORMATO_DE_MEDIA)
        if dados.categoria == col.TOTAL:
            for j in range(4):
                aba.cell(linha + i, coluna + j).font = negrito


def _tops(aba, painel: Painel, coluna: int, destacar_acima_de: int) -> None:
    referencia = painel.referencia.strftime("%d/%m")
    cabecalho = tuple(r.format(referencia=referencia) for r in col.TABELA_DO_TOP)
    for titulo, linha_do_titulo, notas in (
            (col.TITULO_DO_TOP_DIAS, col.LINHA_DO_TOP_DIAS, painel.top_dias),
            (col.TITULO_DO_TOP_VALOR, col.LINHA_DO_TOP_VALOR, painel.top_valor)):
        _titulo(aba, linha_do_titulo, coluna, titulo, largura=len(cabecalho))
        _cabecalho(aba, linha_do_titulo + 1, coluna, cabecalho)
        for i, nota in enumerate(notas, start=linha_do_titulo + 2):
            _escrever_nota(aba, i, coluna, nota, painel.referencia,
                           destacar_acima_de)


def _escrever_nota(aba, linha: int, coluna: int, nota: Pendencia,
                   referencia: date, destacar_acima_de: int) -> None:
    dias = nota.dias(referencia)
    _celula(aba, linha, coluna, nota.categoria)
    _celula(aba, linha, coluna + 1, nota.guardiao)
    _celula(aba, linha, coluna + 2, nota.gestor)
    _celula(aba, linha, coluna + 3, nota.parceiro)
    _celula(aba, linha, coluna + 4, nota.valor, col.FORMATO_DE_VALOR)
    _celula(aba, linha, coluna + 5, nota.emissao, col.FORMATO_DE_DATA)
    celula = _celula(aba, linha, coluna + 6,
                     "" if dias is None else dias, col.FORMATO_DE_DIAS)
    # O "destacar acima de" do painel original é um número escrito num canto,
    # e nada acontece com ele. Aqui ele pinta a célula — senão é enfeite.
    if dias is not None and destacar_acima_de > 0 and dias > destacar_acima_de:
        celula.fill = PatternFill("solid", fgColor=COR_DO_DESTAQUE)
        celula.font = Font(name=FONTE, size=TAMANHO, bold=True,
                           color=COR_DA_LETRA_DESTACADA)


def _ajustes(aba, painel: Painel, destacar_acima_de: int) -> None:
    """Os dois rótulos do canto: a data de referência e o limite de destaque."""
    coluna = col.COLUNA_DOS_AJUSTES
    _celula(aba, col.LINHA_DO_TITULO, coluna, col.ROTULO_DA_REFERENCIA)
    _celula(aba, col.LINHA_DO_TITULO, coluna + 1, painel.referencia,
            col.FORMATO_DE_DATA)
    _celula(aba, col.LINHA_DA_PIZZA, coluna, col.ROTULO_DO_DESTAQUE)
    _celula(aba, col.LINHA_DA_PIZZA, coluna + 1, destacar_acima_de,
            col.FORMATO_DE_DIAS)


def desenhar(aba, auxiliar, painel: Painel) -> None:
    """Prende os três gráficos ao painel, apontados para a aba auxiliar.

    Bloco vazio não vira gráfico vazio: sem unidade cadastrada, o gráfico de
    unidades simplesmente não existe, e a tela diz por quê.
    """
    esquerda = get_column_letter(col.COLUNA_DA_ESQUERDA)
    direita = get_column_letter(col.COLUNA_DA_DIREITA)
    abaixo = col.LINHA_DOS_GRAFICOS_DE_BAIXO + 1

    if painel.categorias:
        aba.add_chart(_pizza(auxiliar, len(painel.categorias)),
                      f"{esquerda}{col.LINHA_DA_PIZZA + 1}")
    if painel.unidades:
        aba.add_chart(
            _barras(auxiliar, col.BLOCO_DO_VALOR_POR_UNIDADE,
                    len(painel.unidades), deitada=False,
                    altura=8.5, largura=11.5),
            f"{esquerda}{abaixo}")
    if painel.guardioes_do_grafico:
        aba.add_chart(
            _barras(auxiliar, col.BLOCO_DO_GRAFICO_DE_GUARDIOES,
                    len(painel.guardioes_do_grafico), deitada=True,
                    altura=8.5, largura=20.0),
            f"{direita}{abaixo}")
