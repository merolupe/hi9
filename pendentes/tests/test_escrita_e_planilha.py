"""A leitura das três extensões e a aba formatada antes de receber dado."""
from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from conftest import CHAVE, escrever
from pendentes import escrita
from pendentes.escrita import Coluna
from pendentes.planilha import PlanilhaIlegivel, ler

COLUNAS = (
    Coluna("Chave Acesso", "texto"),
    Coluna("Valor da Nota", "valor"),
    Coluna("Dh. Emissão", "data"),
    Coluna("Guardião"),
)
LINHAS = [
    [CHAVE, "1.500,00", "01/07/2026", "Suprimentos"],
    [CHAVE[:-1] + "5", 320.5, date(2026, 7, 2), ""],
]


# -- leitura ---------------------------------------------------------------

def test_a_leitura_devolve_as_abas_na_ordem_do_arquivo(tmp_path):
    caminho = escrever(tmp_path / "duas.xlsx", {
        "Pendentes": [["Chave Acesso"], ["1"]],
        "PENDENTES FIS-FAT": [["Chave Acesso"], ["2"]],
    })
    arquivo = ler(caminho)

    assert [aba.nome for aba in arquivo.abas] == ["Pendentes", "PENDENTES FIS-FAT"]
    assert arquivo.primeira.nome == "Pendentes"


def test_a_aba_e_achada_por_nome_sem_caixa_nem_acento(tmp_path):
    caminho = escrever(tmp_path / "uma.xlsx", {"Pendentes": [["Chave Acesso"]]})
    assert ler(caminho).aba("PENDENTES") is not None
    assert ler(caminho).aba("Lançados") is None


def test_o_limite_de_linhas_espia_so_o_comeco(tmp_path):
    """Descobrir o papel de um arquivo não exige carregá-lo inteiro."""
    caminho = escrever(tmp_path / "grande.xlsx",
                       {"Plan1": [["Cabeçalho"]] + [[i] for i in range(500)]})
    assert len(ler(caminho, limite_de_linhas=5).primeira) == 5
    assert len(ler(caminho).primeira) == 501


def test_extensao_que_o_projeto_nao_le_diz_o_que_esperava(tmp_path):
    caminho = tmp_path / "relatorio.csv"
    caminho.write_text("a;b", encoding="utf-8")
    with pytest.raises(PlanilhaIlegivel) as erro:
        ler(caminho)
    assert ".xls" in str(erro.value)


def test_arquivo_inexistente_estoura_dizendo_o_caminho(tmp_path):
    with pytest.raises(FileNotFoundError):
        ler(tmp_path / "nao-existe.xlsx")


# -- escrita ---------------------------------------------------------------

def test_o_formato_da_coluna_e_aplicado_antes_de_qualquer_escrita():
    """A doutrina que uma reimplementação ingênua erraria.

    Quando `preparar_aba` retorna, a aba ainda não tem nenhuma linha de dado e
    já sabe o formato de cada coluna.
    """
    livro = escrita.novo_livro()
    aba = escrita.preparar_aba(livro, "Pendentes", COLUNAS)

    assert aba.max_row == 1                       # só o cabeçalho
    assert aba.column_dimensions["A"].number_format == "@"
    assert aba.column_dimensions["B"].number_format == "#,##0.00"
    assert aba.column_dimensions["C"].number_format == "DD/MM/YYYY"


def test_a_chave_de_44_digitos_sai_como_texto_e_nao_como_notacao_cientifica(
        tmp_path):
    livro = escrita.novo_livro()
    escrita.escrever_aba(livro, "Pendentes", COLUNAS, LINHAS)
    caminho = escrita.salvar(livro, tmp_path / "saida.xlsx")

    aba = load_workbook(str(caminho))["Pendentes"]
    assert aba.cell(2, 1).value == CHAVE
    assert isinstance(aba.cell(2, 1).value, str)
    assert aba.cell(2, 1).number_format == "@"


def test_valor_vira_numero_e_data_vira_data(tmp_path):
    livro = escrita.novo_livro()
    escrita.escrever_aba(livro, "Pendentes", COLUNAS, LINHAS)
    caminho = escrita.salvar(livro, tmp_path / "saida.xlsx")

    aba = load_workbook(str(caminho))["Pendentes"]
    assert aba.cell(2, 2).value == 1500.0
    assert aba.cell(2, 3).value.date() == date(2026, 7, 1)
    assert aba.cell(3, 2).value == 320.5


def test_celula_de_valor_vazia_continua_vazia_em_vez_de_virar_zero():
    """Zero é uma afirmação; ausência não é."""
    assert escrita.valor_formatado("", "valor") == ""
    assert escrita.valor_formatado(None, "data") == ""
    assert escrita.valor_formatado("", "texto") == ""


def test_a_aba_sai_com_autofiltro_e_com_a_primeira_linha_congelada(tmp_path):
    livro = escrita.novo_livro()
    escrita.escrever_aba(livro, "Pendentes", COLUNAS, LINHAS)
    caminho = escrita.salvar(livro, tmp_path / "saida.xlsx")

    aba = load_workbook(str(caminho))["Pendentes"]
    assert aba.auto_filter.ref == "A1:D3"
    assert aba.freeze_panes == "A2"


def test_a_largura_e_calculada_do_conteudo_e_tem_teto(tmp_path):
    longa = [Coluna("Descrição", "texto")]
    livro = escrita.novo_livro()
    escrita.escrever_aba(livro, "Larga", longa, [["x" * 300]])
    caminho = escrita.salvar(livro, tmp_path / "larga.xlsx")

    aba = load_workbook(str(caminho))["Larga"]
    assert aba.column_dimensions["A"].width == escrita.LARGURA_MAXIMA


def test_aba_auxiliar_fica_oculta_e_nao_apagada(tmp_path):
    """Ocultar, e não apagar: as auxiliares são evidência para auditoria."""
    livro = escrita.novo_livro()
    escrita.escrever_aba(livro, "Pendentes", COLUNAS, LINHAS)
    escrita.escrever_aba(livro, "CTe", COLUNAS, [], oculta=True)
    caminho = escrita.salvar(livro, tmp_path / "saida.xlsx")

    relido = load_workbook(str(caminho))
    assert "CTe" in relido.sheetnames
    assert relido["CTe"].sheet_state == "hidden"


def test_a_planilha_escrita_e_lida_de_volta_pelo_proprio_projeto(tmp_path):
    livro = escrita.novo_livro()
    escrita.escrever_aba(livro, "Pendentes", COLUNAS, LINHAS)
    caminho = escrita.salvar(livro, tmp_path / "ida-e-volta.xlsx")

    lido = ler(caminho).aba("Pendentes")
    assert lido.linhas[0][0] == "Chave Acesso"
    assert lido.linhas[1][0] == CHAVE
