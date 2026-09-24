"""Uma semana inteira de mercadorias, do arquivo arrastado à planilha gravada.

Os arquivos são sintéticos e montados no próprio teste — chave de acesso, CNPJ
e nome de fornecedor inventados, regra nº 1 do `CLAUDE.md`. O que eles provam
não é o número de uma execução real (isso depende do padrão-ouro, que ainda não
chegou): é que as sete abas saem com a largura certa, que cada nota vai para
exatamente uma delas, que o que a limpeza tirou continua existindo em
`Descartados`, que a classificação atravessa a semana pelo livro e que o que a
ferramenta não consegue explicar aparece na tela em vez de sumir.
"""
from __future__ import annotations

import openpyxl
import pytest

from pendentes import estado, parametros, snapshot
from pendentes.mercadorias import colunas as col
from pendentes.mercadorias.execucao import SemRegistros, gerar
from conftest import (
    COLUNAS_CE, COLUNAS_XML, FAROL_VERDE, FAROL_VERMELHO, chave_de, escrever,
    linha_ce, linha_xml, relatorio,
)

# Cada chave exercita um caminho do pipeline.
PENDENTE = chave_de(1)          # sobra depois das quatro condições
LANCADA = chave_de(2)           # Conf fiscal = Sim
PARA_FISCAL = chave_de(3)       # física Sim + incongruência vazia → B1 → B2
DO_CTE = chave_de(4)            # tomador de CT-e
DESCONHECIDA = chave_de(5)      # manifestação desconhecida
CANCELADA = chave_de(6)         # NF-e cancelada
DE_ENTRADA = chave_de(7)        # Entrada/Saida = Entrada
DE_TERCEIRO = chave_de(8)       # A1: sem Nome Fantasia
DE_TRANSPORTE = chave_de(9)     # A3: NF-e destinada a transporte
SEM_CADASTRO = chave_de(10)     # A2: sem Nome Parceiro


@pytest.fixture()
def semana(tmp_path) -> dict:
    """Os três relatórios da semana, mais a pasta de dados e a de saída."""
    xml = relatorio(tmp_path / "XML31.xlsx", COLUNAS_XML, [
        linha_xml("1001", PENDENTE, valor="1500.00"),
        linha_xml("1002", LANCADA, valor="200.00"),
        linha_xml("1003", PARA_FISCAL, valor="300.00"),
        linha_xml("1004", DO_CTE, tomador="TRANSPORTADORA XPTO LTDA"),
        linha_xml("1005", DESCONHECIDA, manifestacao="Desconhecida"),
        linha_xml("1006", CANCELADA, situacao="NF-e cancelada"),
        linha_xml("1007", DE_ENTRADA, entrada_saida="Entrada"),
        linha_xml("1008", DE_TERCEIRO, fantasia=""),
        linha_xml("1009", DE_TRANSPORTE, tipo="nf-e  destinada a transporte "),
        linha_xml("1010", SEM_CADASTRO, parceiro="", cnpj="98765432000188",
                  valor="50.00"),
    ])
    ce = relatorio(tmp_path / "CE31.xlsx", COLUNAS_CE, [
        linha_ce(LANCADA, fiscal="Sim", farol=FAROL_VERDE),
        linha_ce(PARA_FISCAL, fisica="Sim", incongruencia="", fiscal="não",
                 pedido="7788", farol=FAROL_VERDE),
        linha_ce(SEM_CADASTRO, fisica="", fiscal="não", pedido="",
                 farol="", data=""),
        linha_ce(PENDENTE, fisica="Sim", incongruencia="divergência de peso",
                 fiscal="não", pedido="9001", farol=FAROL_VERMELHO),
    ])
    return {
        "arquivos": [xml, ce],
        "saida": tmp_path / "saida",
        "dados": tmp_path / "dados",
        "tmp": tmp_path,
    }


def _rodar(semana, **extras):
    semana["saida"].mkdir(exist_ok=True)
    return gerar(semana["arquivos"], semana["saida"],
                 raiz_dos_dados=semana["dados"], responsavel="teste", **extras)


def _aba(resultado, nome):
    return openpyxl.load_workbook(resultado.planilha)[nome]


def _linhas(resultado, nome):
    aba = _aba(resultado, nome)
    return [[c.value for c in linha] for linha in aba.iter_rows(min_row=2)]


def planilha_anterior(caminho, semana_anterior, pendentes=(), fis_fat=()):
    """Uma planilha da semana passada, com as duas abas que a herança lê."""
    return escrever(caminho, {
        "Pendentes": [[c.rotulo for c in col.pendentes(semana_anterior)],
                      *pendentes],
        "PENDENTES FIS-FAT": [[c.rotulo for c in col.fis_fat(semana_anterior)],
                              *fis_fat],
    })


def linha_da_semana_anterior(chave, *, tipo="", guardiao="", gestor="",
                             categoria="", retorno="", fis_fat=False):
    """Uma linha devolvida pelo time, com a classificação preenchida.

    As posições saem do **layout**, pelo nome da coluna: é o que mantém a
    fixture honesta quando a largura da aba muda — como mudou quando
    `Pedido vinculado` entrou.
    """
    layout = col.fis_fat(30) if fis_fat else col.pendentes(30)
    linha = [""] * len(layout)
    if fis_fat:
        linha[0], linha[1], linha[2] = tipo, guardiao, retorno
    else:
        linha[0], linha[1], linha[2] = tipo, guardiao, gestor
        linha[3], linha[4] = categoria, retorno
    linha[col.indice(layout, col.X_CHAVE)] = chave
    linha[col.indice(layout, col.X_NRO_NOTA)] = "1001"   # para o papel casar
    return linha


def _em(linha, rotulo, *, semana=31):
    """O valor de uma coluna da `Pendentes` pelo nome dela, nunca por índice."""
    return linha[col.indice(col.pendentes(semana), rotulo)]


# -- a planilha -------------------------------------------------------------

def test_a_execucao_gera_as_sete_abas_na_ordem_e_com_a_largura_de_cada_uma(
        semana):
    resultado = _rodar(semana)
    livro = openpyxl.load_workbook(resultado.planilha)

    assert livro.sheetnames == [nome for nome, _ in col.ABAS]
    assert livro["Pendentes"].max_column == 39
    assert livro["PENDENTES FIS-FAT"].max_column == 36
    assert livro["CTe"].max_column == 27
    assert livro["Manifestados"].max_column == 27
    assert livro["Entradas 3os"].max_column == 27
    assert livro["Lançados"].max_column == 33
    assert livro["Descartados"].max_column == 28


def test_as_auxiliares_ficam_ocultas_e_reexibiveis(semana):
    """Ocultar não é apagar: elas são evidência para auditoria."""
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    for nome, visivel in col.ABAS:
        esperado = "visible" if visivel else "hidden"
        assert livro[nome].sheet_state == esperado, nome


def test_as_abas_saem_com_a_ordem_exata_de_colunas(semana):
    resultado = _rodar(semana)
    livro = openpyxl.load_workbook(resultado.planilha)
    layouts = {
        "Pendentes": col.pendentes(resultado.semana),
        "PENDENTES FIS-FAT": col.fis_fat(resultado.semana),
        "CTe": col.AUXILIARES,
        "Lançados": col.LANCADOS,
        "Descartados": col.DESCARTADOS,
    }
    for nome, layout in layouts.items():
        lido = [c.value for c in livro[nome][1]]
        assert lido == [c.rotulo for c in layout], nome


def test_cada_nota_do_xml_vai_para_exatamente_uma_aba(semana):
    resultado = _rodar(semana)
    livro = openpyxl.load_workbook(resultado.planilha)
    escritas = sum(livro[nome].max_row - 1 for nome, _ in col.ABAS)
    assert escritas == resultado.documentos == 10


def test_o_arquivo_sai_com_o_nome_que_o_time_ja_reconhece(semana):
    resultado = _rodar(semana)
    assert resultado.planilha.name == f"Pendentes{resultado.semana}.xlsx"


def test_o_roteamento_distribui_as_notas_pelas_tres_auxiliares(semana):
    resultado = _rodar(semana)
    assert resultado.por_destino == {"CTe": 1, "Manifestados": 2,
                                     "Entradas 3os": 1}
    assert [linha[0] for linha in _linhas(resultado, "CTe")] == ["1004"]
    assert [linha[0] for linha in _linhas(resultado, "Manifestados")] == \
        ["1005", "1006"]


def test_o_que_a_limpeza_tirou_esta_na_aba_descartados_com_o_motivo(semana):
    """Hoje estas linhas somem sem rastro, sem contador e sem aba."""
    resultado = _rodar(semana)
    descartados = _linhas(resultado, "Descartados")
    assert [(linha[0], linha[27]) for linha in descartados] == [
        ("1008", "XML de terceiro (sem Nome Fantasia)"),
        ("1009", "NF-e destinada a transporte"),
    ]
    assert resultado.xml_de_terceiro == 1
    assert resultado.nfe_de_transporte == 1


def test_a_nota_lancada_sai_da_principal_e_vai_para_lancados(semana):
    resultado = _rodar(semana)
    assert [linha[0] for linha in _linhas(resultado, "Lançados")] == ["1002"]
    assert resultado.lancados == 1
    numeros = [linha[5] for linha in _linhas(resultado, "Pendentes")]
    assert "1002" not in numeros


def test_b1_manda_a_nota_conferida_para_a_fis_fat(semana):
    """Física conferida e sem divergência é pendência de lançamento fiscal."""
    resultado = _rodar(semana)
    fis_fat = _linhas(resultado, "PENDENTES FIS-FAT")
    assert [linha[3] for linha in fis_fat] == ["1003"]
    assert fis_fat[0][1] == "Fiscal"
    assert resultado.reclassificadas_para_fiscal == 1
    assert resultado.fis_fat == 1


def test_o_bloco_de_conferencia_chega_na_planilha_com_os_seis_valores(semana):
    resultado = _rodar(semana)
    linha = next(l for l in _linhas(resultado, "Pendentes") if l[5] == "1001")
    assert _em(linha, col.P_CONF_FISICA) == "Sim"
    assert _em(linha, col.P_DATA_CONF_FISICA).strftime("%d/%m/%Y") == "05/07/2026"
    assert _em(linha, col.P_CONF_FISCAL) == "não"
    assert _em(linha, col.P_INCONGRUENCIA) == "divergência de peso"
    assert _em(linha, col.P_NRO_DO_PEDIDO) == "9001"
    assert _em(linha, col.P_PEDIDO_CONFIRMADO) == "Não"       # farol vermelho


def test_dias_emissao_doc_chega_como_veio_com_formato_de_data(semana):
    """O defeito da pendência nº 2, preservado **exatamente** como ele é.

    O VBA escreve essa coluna sem conversão e formata a coluna como data. Se
    o valor for contagem de dias, o Excel exibe uma data de 1900. Converter o
    número numa data de verdade aqui trocaria um defeito por outro — e a
    célula deixaria de ser a que a macro produz.
    """
    aba = _aba(_rodar(semana), "Pendentes")
    rotulos = [c.value for c in aba[1]]
    celula = aba.cell(2, rotulos.index("Dias Emissão Doc") + 1)
    assert celula.number_format == "DD/MM/YYYY"
    # A célula guarda o 12 que veio do relatório; quem lê com formato de data
    # vê 12/01/1900, que é o que o Excel exibe hoje. Se a ferramenta tivesse
    # convertido o número numa data, a célula guardaria 11/01/1900 — um dia a
    # menos, porque a conversão não reproduz o calendário de 1900 do Excel.
    assert celula.value.strftime("%d/%m/%Y") == "12/01/1900"


def test_o_farol_sai_nos_tres_estados(semana):
    resultado = _rodar(semana)
    assert resultado.farol == {"Sim": 2, "Não": 1, "sem pedido": 1}


def test_a_regra_a2_chega_na_planilha(semana):
    resultado = _rodar(semana)
    linha = next(l for l in _linhas(resultado, "Pendentes") if l[5] == "1010")
    assert linha[6] == "Sem cadastro"
    assert linha[7] == "98765432000188"


def test_a_chave_de_acesso_sai_como_texto_de_44_digitos(semana):
    """44 dígitos em coluna numérica viram notação científica e o PROCX cai."""
    resultado = _rodar(semana)
    aba = _aba(resultado, "Pendentes")
    celula = aba.cell(2, col.indice(col.pendentes(31), col.X_CHAVE) + 1)
    assert celula.value == PENDENTE
    assert celula.number_format == "@"


def test_o_pedido_vinculado_chega_na_planilha(semana):
    """A nona coluna, do jeito que o relatório da semana 37 a mostra.

    A nota 1001 foi conferida e tem pedido — sai o número. A 1010 entrou no CE
    sem conferência física — sai o rótulo. O terceiro caso, a nota que nem está
    no CE, é exercitado no teste de unidade: nesta semana sintética todas as
    notas que sobram em `Pendentes` estão no relatório de conferência.
    """
    linhas = _linhas(_rodar(semana), "Pendentes")
    por_nota = {l[5]: _em(l, col.P_PEDIDO_VINCULADO) for l in linhas}
    assert por_nota["1001"] == 9001
    assert por_nota["1010"] == "NA Conf Física"


def test_o_pedido_vinculado_nao_aparece_na_fis_fat(semana):
    aba = _aba(_rodar(semana), "PENDENTES FIS-FAT")
    cabecalho = [c.value for c in next(aba.iter_rows(max_row=1))]
    assert col.P_PEDIDO_VINCULADO not in cabecalho
    assert aba.max_column == 36


def test_as_duas_abas_visiveis_tem_autofiltro_e_cabecalho_congelado(semana):
    aba = _aba(_rodar(semana), "Pendentes")
    assert aba.auto_filter.ref.startswith("A1:")
    assert aba.freeze_panes == "A2"


# -- o que a tela mostra ----------------------------------------------------

def test_o_titulo_traz_a_semana_a_contagem_e_o_valor(semana):
    resultado = _rodar(semana)
    assert f"Semana {resultado.semana}" in resultado.titulo()
    assert "3 notas pendentes" in resultado.titulo()
    assert "R$ 1.850,00" in resultado.titulo()


def test_as_quatro_fichas_sao_as_que_o_time_olha_primeiro(semana):
    fichas = dict(_rodar(semana).fichas())
    assert list(fichas) == ["Pendentes (áreas)", "Pendentes Fiscal/Fat.",
                            "Lançados na semana", "Sem classificação"]
    assert fichas["Pendentes (áreas)"] == "2"
    assert fichas["Pendentes Fiscal/Fat."] == "1"


def test_o_descarte_aparece_na_tela_contado(semana):
    listas = dict((titulo, itens) for titulo, itens, _ in _rodar(semana).listas())
    itens = listas["Descartados antes do roteamento"]
    assert "XML de terceiro (sem Nome Fantasia): 1" in itens
    assert "NF-e destinada a transporte: 1" in itens


def test_o_roteamento_aparece_na_tela_como_teste_de_regressao_visivel(semana):
    titulos = [titulo for titulo, _, _ in _rodar(semana).listas()]
    assert "Para onde as notas foram" in titulos


def test_sem_nada_a_bloquear_a_semana_sai_encerravel(semana):
    resultado = _rodar(semana)
    assert resultado.bloqueios() == []
    assert resultado.encerravel


def test_a_unidade_nao_reconhecida_bloqueia_e_a_lista_vermelha_vem_primeiro(
        semana):
    dados = parametros.carregar_fabrica()
    dados["unidades"] = [{"ordem": 1, "trecho": "MATRIZ", "unidade": "Matriz"}]
    resultado = _rodar(semana, dados=dados)

    # A nota 1010 tem fantasia "HINOVE MATRIZ"; a de nº 1001 também. Nenhuma
    # cai fora — o que cai fora é a que a fixture não nomeia como matriz.
    assert resultado.encerravel

    dados["unidades"] = [{"ordem": 1, "trecho": "CORUMB", "unidade": "Corumbá"}]
    resultado = _rodar(semana, dados=dados)
    assert not resultado.encerravel
    assert resultado.listas()[0][2] == "erro"
    assert "HINOVE MATRIZ" in resultado.listas()[0][1][0]


def test_a_chave_duplicada_que_diverge_no_ce_bloqueia_e_a_igual_so_conta(
        semana, tmp_path):
    ce = relatorio(tmp_path / "CE-duplicado.xlsx", COLUNAS_CE, [
        linha_ce(PENDENTE, pedido="7788"),
        linha_ce(PENDENTE, pedido="7788"),
        linha_ce(LANCADA, pedido="1"),
        linha_ce(LANCADA, pedido="2"),
    ])
    semana["arquivos"] = [semana["arquivos"][0], ce]
    resultado = _rodar(semana)

    assert resultado.chaves_duplicadas == 2
    assert resultado.chaves_divergentes == [LANCADA]
    assert not resultado.encerravel
    assert any("divergem" in item for item in resultado.bloqueios())


def test_arquivo_que_nao_casa_com_papel_nenhum_aparece_na_tela(semana, tmp_path):
    estranho = escrever(tmp_path / "Planilha da Ana.xlsx",
                        {"Plan1": [["Coisa", "Outra"], ["a", "b"]]})
    semana["arquivos"].append(estranho)
    resultado = _rodar(semana)
    assert any("Planilha da Ana.xlsx" in aviso for aviso in resultado.atencoes())


# -- o livro e a semana -----------------------------------------------------

def test_a_classificacao_da_semana_passada_volta_pela_planilha(semana):
    resultado = _rodar(semana)
    anterior = planilha_anterior(
        semana["tmp"] / "Pendentes30.xlsx", resultado.semana - 1,
        pendentes=[linha_da_semana_anterior(
            PENDENTE, tipo="Compra direta", guardiao="Suprimentos",
            gestor="Maria", categoria="Diretos", retorno="aguardando")])
    semana["arquivos"].append(anterior)

    segunda = _rodar(semana)
    linha = next(l for l in _linhas(segunda, "Pendentes") if l[5] == "1001")
    assert linha[:5] == ["Compra direta", "Suprimentos", "Maria", "Diretos",
                         "aguardando"]
    assert segunda.herdadas == 1
    assert segunda.ingestao is not None and segunda.ingestao.novas == 1


def test_a_classificacao_sobrevive_a_semana_seguinte_sem_o_arquivo(semana):
    """O livro guarda todo mundo: perder o anexo do e-mail deixa de apagar."""
    primeira = _rodar(semana)
    anterior = planilha_anterior(
        semana["tmp"] / "Pendentes30.xlsx", primeira.semana - 1,
        pendentes=[linha_da_semana_anterior(PENDENTE, guardiao="Suprimentos")])
    semana["arquivos"].append(anterior)
    _rodar(semana)

    semana["arquivos"].remove(anterior)
    terceira = _rodar(semana)
    linha = next(l for l in _linhas(terceira, "Pendentes") if l[5] == "1001")
    assert linha[1] == "Suprimentos"
    assert terceira.herdadas == 1


def test_a_classificacao_sobrevive_a_semana_em_que_a_nota_sai_de_pendentes(
        semana, tmp_path):
    """A nota lançada não some do livro — e volta classificada se reaparecer.

    `[FATO]` No VBA a herança lê **só** as abas `Pendentes` e
    `PENDENTES FIS-FAT` da semana anterior. Uma nota que foi para `Lançados`
    ou para `CTe` perde a classificação, e se voltar, volta vazia.
    """
    primeira = _rodar(semana)
    anterior = planilha_anterior(
        semana["tmp"] / "Pendentes30.xlsx", primeira.semana - 1,
        pendentes=[linha_da_semana_anterior(LANCADA, guardiao="Suprimentos")])
    semana["arquivos"].append(anterior)
    _rodar(semana)                       # a nota está em `Lançados` nesta
    semana["arquivos"].remove(anterior)

    # Na semana seguinte ela volta a ser pendente: o CE não a traz mais.
    ce = relatorio(tmp_path / "CE32.xlsx", COLUNAS_CE, [linha_ce(PENDENTE)])
    semana["arquivos"] = [semana["arquivos"][0], ce]
    terceira = _rodar(semana)

    linha = next(l for l in _linhas(terceira, "Pendentes") if l[5] == "1002")
    assert linha[1] == "Suprimentos"


def test_a_heranca_le_tambem_a_aba_fis_fat_da_semana_anterior(semana):
    """A `PENDENTES FIS-FAT` é a segunda fonte, e o VBA a lê por cabeçalho."""
    primeira = _rodar(semana)
    anterior = planilha_anterior(
        semana["tmp"] / "Pendentes30.xlsx", primeira.semana - 1,
        fis_fat=[linha_da_semana_anterior(
            SEM_CADASTRO, tipo="Compra direta", guardiao="Faturamento",
            retorno="em análise", fis_fat=True)])
    semana["arquivos"].append(anterior)

    segunda = _rodar(semana)
    fis_fat = _linhas(segunda, "PENDENTES FIS-FAT")
    assert [linha[3] for linha in fis_fat] == ["1003", "1010"]
    da_nota = next(l for l in fis_fat if l[3] == "1010")
    assert da_nota[:3] == ["Compra direta", "Faturamento", "em análise"]


def test_a_planilha_da_semana_anterior_vence_a_fis_fat_onde_as_duas_divergem(
        semana):
    primeira = _rodar(semana)
    anterior = planilha_anterior(
        semana["tmp"] / "Pendentes30.xlsx", primeira.semana - 1,
        pendentes=[linha_da_semana_anterior(SEM_CADASTRO,
                                            guardiao="Suprimentos")],
        fis_fat=[linha_da_semana_anterior(SEM_CADASTRO, guardiao="Faturamento",
                                          fis_fat=True)])
    semana["arquivos"].append(anterior)

    segunda = _rodar(semana)
    linha = next(l for l in _linhas(segunda, "Pendentes") if l[5] == "1010")
    assert linha[1] == "Suprimentos"


def test_a_ingestao_e_idempotente(semana):
    primeira = _rodar(semana)
    anterior = planilha_anterior(
        semana["tmp"] / "Pendentes30.xlsx", primeira.semana - 1,
        pendentes=[linha_da_semana_anterior(PENDENTE, guardiao="Suprimentos")])
    semana["arquivos"].append(anterior)

    _rodar(semana)
    segunda = _rodar(semana)
    assert segunda.ingestao.novas == 0
    assert segunda.ingestao.alteradas == 0


def test_o_livro_e_gravado_fora_do_git_com_carimbo(semana):
    _rodar(semana)
    caminho = estado.caminho_do_livro("mercadorias", raiz=semana["dados"])
    assert caminho.is_file()
    livro = estado.carregar("mercadorias", raiz=semana["dados"])
    assert livro.atualizado_por == "teste"


def test_a_semana_vira_snapshot_com_a_planilha_e_as_impressoes(semana):
    resultado = _rodar(semana)
    pasta = resultado.pasta_do_snapshot
    assert (pasta / "classificacao.yaml").is_file()
    assert (pasta / "entradas.json").is_file()
    assert (pasta / resultado.planilha.name).is_file()
    ultima = snapshot.ultima_semana("mercadorias", raiz=semana["dados"])
    assert ultima["encerravel"] is True
    assert ultima["roteamento"] == {"CTe": 1, "Manifestados": 2,
                                    "Entradas 3os": 1}
    assert ultima["valor_pendente"] == 1850.0


def test_rodar_a_mesma_semana_duas_vezes_nao_sobrescreve_a_evidencia(semana):
    primeira = _rodar(semana)
    segunda = _rodar(semana)
    assert segunda.pasta_do_snapshot.name.endswith("-2")
    assert primeira.pasta_do_snapshot != segunda.pasta_do_snapshot


def test_a_semana_cadastrada_vence_a_deduzida(semana):
    dados = parametros.carregar_fabrica()
    dados["semana"] = {"numero": "31", "data_de_referencia": "10/08/2026"}
    resultado = _rodar(semana, dados=dados)
    assert (resultado.ano, resultado.semana) == (2026, 31)
    rotulos = [c.value for c in _aba(resultado, "Pendentes")[1]]
    assert rotulos[4] == "Retorno semana 30"


# -- o que aborta, e o que só degrada --------------------------------------

def test_sem_o_xml_a_execucao_para_nomeando_o_papel(semana):
    from pendentes.papeis import PapelAusente

    with pytest.raises(PapelAusente, match="XML"):
        gerar([semana["arquivos"][1]], semana["tmp"],
              raiz_dos_dados=semana["dados"])


def test_sem_a_conferencia_de_entradas_a_execucao_para(semana):
    from pendentes.papeis import PapelAusente

    with pytest.raises(PapelAusente, match="Conferência de Entradas"):
        gerar([semana["arquivos"][0]], semana["tmp"],
              raiz_dos_dados=semana["dados"])


def test_coluna_obrigatoria_ausente_lista_todas_de_uma_vez(tmp_path):
    from pendentes.cabecalho import ColunasFaltando

    mutiladas = [c for c in COLUNAS_XML if c not in ("Série Doc", "Status")]
    xml = relatorio(tmp_path / "XML.xlsx", mutiladas,
                    [linha_xml("1", chave_de(1))[:len(mutiladas)]])
    ce = relatorio(tmp_path / "CE.xlsx", COLUNAS_CE, [linha_ce(chave_de(1))])
    with pytest.raises(ColunasFaltando) as erro:
        gerar([xml, ce], tmp_path, raiz_dos_dados=tmp_path / "dados")
    assert "Série Doc" in str(erro.value) and "Status" in str(erro.value)


def test_o_rodape_do_relatorio_nao_vira_nota(tmp_path):
    xml = relatorio(tmp_path / "XML.xlsx", COLUNAS_XML, [
        linha_xml("1", chave_de(1)),
        ["", "", "", "", "", "Total: 1.500,00"],
    ])
    ce = relatorio(tmp_path / "CE.xlsx", COLUNAS_CE, [linha_ce(chave_de(1))])
    (tmp_path / "saida").mkdir()
    resultado = gerar([xml, ce], tmp_path / "saida",
                      raiz_dos_dados=tmp_path / "dados")
    assert resultado.documentos == 1


def test_xml_sem_nenhuma_linha_de_dado_para_dizendo_isso(tmp_path):
    xml = relatorio(tmp_path / "XML.xlsx", COLUNAS_XML, [])
    ce = relatorio(tmp_path / "CE.xlsx", COLUNAS_CE, [linha_ce(chave_de(1))])
    with pytest.raises(SemRegistros):
        gerar([xml, ce], tmp_path, raiz_dos_dados=tmp_path / "dados")


def test_a_planilha_da_semana_anterior_nao_e_confundida_com_o_xml(semana):
    """As 27 colunas do XML estão contidas nela; só a âncora ausente separa."""
    primeira = _rodar(semana)
    anterior = planilha_anterior(
        semana["tmp"] / "Pendentes30.xlsx", primeira.semana - 1,
        pendentes=[linha_da_semana_anterior(PENDENTE, guardiao="Suprimentos")])
    semana["arquivos"].append(anterior)
    segunda = _rodar(semana)
    assert segunda.documentos == 10


def test_a_saida_da_semana_passada_nao_e_tomada_pelo_xml(semana):
    """`CTe`, `Manifestados` e `Descartados` são o XML cru, sem `Guardião`.

    Anexada de volta, a planilha gerada casava com o papel do XML por uma
    dessas abas — o mesmo defeito que serviços tinha com o Portal de Compras.
    """
    primeira = _rodar(semana)
    semana["arquivos"].append(primeira.planilha)
    segunda = _rodar(semana)
    assert segunda.documentos == 10


# =========================================================================
# B1.5 — o pedido de compra, na planilha inteira
# =========================================================================

def test_b1_5_escreve_suprimentos_na_planilha_e_deixa_a_nota_na_pendentes(semana):
    """A nota `PENDENTE` do fixture tem o farol vermelho: pedido não confirmado.

    O teste é sobre o arquivo entregue, não sobre a função: ela sai com
    `Suprimentos` no guardião, **na `Pendentes`** — sem aba própria.
    """
    resultado = _rodar(semana)

    linha = next(l for l in _linhas(resultado, "Pendentes") if l[5] == "1001")
    assert linha[1] == "Suprimentos"
    assert resultado.suprimentos.nao_confirmado == 1
    assert "PENDENTES FIS-FAT" in _aba(resultado, "Pendentes").parent.sheetnames
    assert all(l[1] != "Suprimentos"
               for l in _linhas(resultado, "PENDENTES FIS-FAT"))


def test_b1_5_nao_toca_na_nota_sem_pedido_vinculado(semana):
    """`SEM_CADASTRO` está no CE com o farol vazio: não há pedido a confirmar."""
    resultado = _rodar(semana)

    assert resultado.suprimentos.sem_pedido >= 1
    linha = next((l for l in _linhas(resultado, "Pendentes") if l[5] == "1010"),
                 None)
    assert linha is not None
    assert linha[1] in (None, "")


def test_b1_5_aparece_no_relato_da_tela(semana):
    resultado = _rodar(semana)
    texto = " ".join(resultado.heranca())

    assert "regra B1.5" in texto
    assert "Suprimentos" in texto
