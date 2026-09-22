"""A base preenche o que está vazio, marca o que preencheu, e não mente.

`[FATO]` Ligada por decisão do Compliance Tributário de 22/09/2026, depois da
medição contra a semana 38 — categoria acertando 44 de 44 com evidência firme,
guardião acertando 77% no grau de sugestão. O que estes testes protegem são as
três travas que a decisão veio acompanhada de: **só preenche o que está
vazio**, **a célula preenchida sai marcada**, e o preenchimento roda **entre
B1 e B2**.
"""
from __future__ import annotations

import openpyxl
import pytest

from pendentes import parametros
from pendentes.conhecimento import base as bc
from pendentes.conhecimento.base import Conhecimento
from pendentes.mercadorias import colunas as col
from pendentes.mercadorias import precategorizacao as pre
from pendentes.mercadorias.execucao import gerar
from pendentes.mercadorias.fontes import Documento
from conftest import (COLUNAS_CE, COLUNAS_XML, chave_de, linha_ce, linha_xml,
                      relatorio)

FIRME, SUGESTAO = bc.FIRME, bc.SUGESTAO
TUDO = {col.C_CATEGORIA: SUGESTAO, col.C_GUARDIAO: SUGESTAO}


def conhecimento(**parceiros) -> Conhecimento:
    """Uma base pronta, sem passar pela importação — é a consulta que importa."""
    def campo(valor, notas, semanas):
        return {"proposto": valor,
                "evidencia": {"valor": valor, "notas": notas, "apoio": notas,
                              "semanas": semanas,
                              "alternativas": {valor: notas}}}

    base = Conhecimento()
    for codigo, dados in parceiros.items():
        codigo = codigo.lstrip("p")
        base.parceiros[codigo] = {"nome": f"PARCEIRO {codigo}"}
        for campo_nome, (valor, notas, semanas) in dados.items():
            base.parceiros[codigo][campo_nome] = campo(valor, notas, semanas)
    return base


def documento(codigo="4", categorizacao=None, fantasia="HINOVE MATRIZ"):
    valores = linha_xml("1001", chave_de(1), cod_parceiro=codigo,
                        fantasia=fantasia)
    documento = Documento(posicao=0, valores=valores)
    documento.categorizacao = categorizacao or ["" for _ in range(5)]
    return documento


# -- só preenche o que está vazio ------------------------------------------

def test_preenche_a_celula_vazia_com_o_que_a_base_propoe():
    base = conhecimento(p4={"guardiao": ("Almoxarifado", 9, 4),
                            "categoria": ("Indiretos", 9, 4)})
    doc = documento()
    relato = pre.preencher([doc], base, minimo_por_coluna=TUDO)

    assert doc.categorizacao[col.POSICAO_DA_CATEGORIZACAO[col.C_GUARDIAO]] == \
        "Almoxarifado"
    assert doc.categorizacao[col.POSICAO_DA_CATEGORIZACAO[col.C_CATEGORIA]] == \
        "Indiretos"
    assert relato.total == 2


def test_nao_sobrescreve_o_que_veio_da_heranca_ou_da_regra_b1():
    """O que uma pessoa classificou vale mais do que qualquer histórico."""
    base = conhecimento(p4={"guardiao": ("Almoxarifado", 9, 4)})
    ja_classificado = ["" for _ in range(5)]
    ja_classificado[col.POSICAO_DA_CATEGORIZACAO[col.C_GUARDIAO]] = "Fiscal"
    doc = documento(categorizacao=ja_classificado)

    relato = pre.preencher([doc], base, minimo_por_coluna=TUDO)
    assert doc.categorizacao[col.POSICAO_DA_CATEGORIZACAO[col.C_GUARDIAO]] == \
        "Fiscal"
    assert relato.total == 0
    assert doc.propostas == {}


def test_celula_com_espaco_em_branco_conta_como_vazia():
    base = conhecimento(p4={"guardiao": ("Almoxarifado", 9, 4)})
    so_espacos = ["" for _ in range(5)]
    so_espacos[col.POSICAO_DA_CATEGORIZACAO[col.C_GUARDIAO]] = "   "
    doc = documento(categorizacao=so_espacos)

    pre.preencher([doc], base, minimo_por_coluna=TUDO)
    assert doc.categorizacao[col.POSICAO_DA_CATEGORIZACAO[col.C_GUARDIAO]] == \
        "Almoxarifado"


# -- o grau mínimo, por coluna ---------------------------------------------

def test_o_grau_minimo_de_cada_coluna_e_parametro():
    """Evidência curta preenche em `sugestao` e não preenche em `firme`."""
    curta = conhecimento(p4={"guardiao": ("Almoxarifado", 1, 1)})

    com_sugestao = documento()
    pre.preencher([com_sugestao], curta,
                  minimo_por_coluna={col.C_GUARDIAO: SUGESTAO})
    assert com_sugestao.propostas == {col.C_GUARDIAO: SUGESTAO}

    so_firme = documento()
    pre.preencher([so_firme], curta, minimo_por_coluna={col.C_GUARDIAO: FIRME})
    assert so_firme.propostas == {}


def test_coluna_desligada_nao_preenche_e_a_tela_diz(capsys):
    base = conhecimento(p4={"guardiao": ("Almoxarifado", 9, 4)})
    doc = documento()
    relato = pre.preencher([doc], base,
                           minimo_por_coluna={col.C_GUARDIAO: "nao"})
    assert doc.propostas == {}
    assert col.C_GUARDIAO in relato.desligadas
    assert any("desligado no parâmetro" in linha for linha in relato.linhas())


def test_base_vazia_nao_preenche_nada():
    doc = documento()
    relato = pre.preencher([doc], Conhecimento(), minimo_por_coluna=TUDO)
    assert relato.total == 0
    assert doc.categorizacao == ["" for _ in range(5)]


def test_parceiro_fora_da_base_continua_vazio():
    base = conhecimento(p999={"guardiao": ("Almoxarifado", 9, 4)})
    doc = documento(codigo="4")
    assert pre.preencher([doc], base, minimo_por_coluna=TUDO).total == 0


def test_a_unidade_da_nota_escolhe_a_regra_mais_especifica():
    base = conhecimento(p4={"guardiao": ("Almoxarifado", 9, 4)})
    base.parceiros["4"]["unidades"] = {
        "HINOVE REGISTRO": {
            "rotulo": "HINOVE REGISTRO",
            "guardiao": {"proposto": "Balança Registro",
                         "evidencia": {"valor": "Balança Registro", "notas": 9,
                                       "apoio": 9, "semanas": 4,
                                       "alternativas": {"Balança Registro": 9}}},
        }
    }
    do_registro = documento(fantasia="Hinove  Registro")
    pre.preencher([do_registro], base, minimo_por_coluna=TUDO)
    assert do_registro.categorizacao[
        col.POSICAO_DA_CATEGORIZACAO[col.C_GUARDIAO]] == "Balança Registro"


# -- a marca ---------------------------------------------------------------

def test_a_marca_sai_na_posicao_que_o_layout_daquela_aba_manda():
    doc = documento()
    doc.propostas = {col.C_GUARDIAO: FIRME, col.C_CATEGORIA: FIRME}

    na_pendentes = pre.marcar(doc, col.pendentes(38))
    assert na_pendentes == {col.indice(col.pendentes(38), col.C_GUARDIAO),
                            col.indice(col.pendentes(38), col.C_CATEGORIA)}

    # A FIS-FAT não tem `Categoria`: coluna que não existe não tem o que marcar.
    na_fis_fat = pre.marcar(doc, col.fis_fat(38))
    assert na_fis_fat == {col.indice(col.fis_fat(38), col.C_GUARDIAO)}


def test_a_contagem_da_tela_separa_por_coluna_e_por_grau():
    relato = pre.Preenchimento()
    relato.contar(col.C_GUARDIAO, SUGESTAO)
    relato.contar(col.C_GUARDIAO, SUGESTAO)
    relato.contar(col.C_CATEGORIA, FIRME)
    linhas = relato.linhas()
    assert f"{col.C_GUARDIAO}: 2 com evidência sugestao" in linhas
    assert f"{col.C_CATEGORIA}: 1 com evidência firme" in linhas
    assert relato.total == 3


# -- a operação, e o conectivo que não é divergência ------------------------

def test_o_conectivo_nao_separa_dois_tipos_de_operacao_iguais():
    """`Compra Uso e Consumo` e `Compra Uso Consumo` são a mesma coisa."""
    assert bc.chave_de_operacao("Compra Uso e Consumo") == \
        bc.chave_de_operacao("Compra Uso Consumo")
    assert bc.chave_de_operacao("Retorno de Conserto") == \
        bc.chave_de_operacao("Retorno Conserto")


def test_normalizar_tira_ligação_e_nao_aproxima_palavras():
    """O que separa `Compra MP` de `Compra Embalagem` é substantivo."""
    assert bc.chave_de_operacao("Compra MP") != \
        bc.chave_de_operacao("Compra Embalagem")
    assert bc.chave_de_operacao("Compra Insumo") != \
        bc.chave_de_operacao("Compra Uso Consumo")


def test_a_grafia_escrita_e_a_que_o_livro_mais_usa(tmp_path):
    from pendentes.estado import Classificacao, Livro

    livro = Livro("mercadorias", registros={
        "a": Classificacao("a", tipo_de_operacao="Compra Uso Consumo"),
        "b": Classificacao("b", tipo_de_operacao="Compra Uso Consumo"),
        "c": Classificacao("c", tipo_de_operacao="Compra Uso e Consumo"),
    })
    grafias = pre.grafias_do_livro(livro)
    assert grafias[bc.chave_de_operacao("Compra Uso e Consumo")] == \
        "Compra Uso Consumo"


def test_a_operacao_sai_na_grafia_do_livro_e_nao_na_da_base():
    """A base aprendeu uma redação; o time usa outra. Quem manda é o time."""
    from pendentes.estado import Classificacao, Livro

    base = Conhecimento()
    base.operacoes["1102"] = {
        "|": {"proposto": "Compra Uso e Consumo",
              "evidencia": {"valor": "Compra Uso e Consumo", "notas": 20,
                            "apoio": 20, "semanas": 9,
                            "alternativas": {"Compra Uso e Consumo": 20}}},
    }
    livro = Livro("mercadorias", registros={
        "a": Classificacao("a", tipo_de_operacao="Compra Uso Consumo"),
    })
    doc = documento()
    pre.preencher([doc], base, livro=livro,
                  minimo_por_coluna={col.C_TIPO_DE_OPERACAO: FIRME})
    posicao = col.POSICAO_DA_CATEGORIZACAO[col.C_TIPO_DE_OPERACAO]
    assert doc.categorizacao[posicao] == "Compra Uso Consumo"


def test_sem_livro_a_operacao_sai_na_grafia_da_base():
    base = Conhecimento()
    base.operacoes["1102"] = {
        "|": {"proposto": "Compra Uso e Consumo",
              "evidencia": {"valor": "Compra Uso e Consumo", "notas": 20,
                            "apoio": 20, "semanas": 9,
                            "alternativas": {"Compra Uso e Consumo": 20}}},
    }
    doc = documento()
    pre.preencher([doc], base,
                  minimo_por_coluna={col.C_TIPO_DE_OPERACAO: FIRME})
    posicao = col.POSICAO_DA_CATEGORIZACAO[col.C_TIPO_DE_OPERACAO]
    assert doc.categorizacao[posicao] == "Compra Uso e Consumo"


def test_a_regra_de_operacao_mais_especifica_usa_a_categoria_recem_preenchida():
    """Categoria é preenchida antes, e a regra por CFOP + parceiro +
    categoria depende dela — a ordem dentro da passagem é a regra."""
    base = conhecimento(p4={"categoria": ("Indiretos", 9, 4)})
    base.operacoes["1102"] = {
        "4|INDIRETOS": {"proposto": "Compra Uso Consumo",
                        "evidencia": {"valor": "Compra Uso Consumo",
                                      "notas": 9, "apoio": 9, "semanas": 4,
                                      "alternativas": {"Compra Uso Consumo": 9}}},
    }
    doc = documento()
    pre.preencher([doc], base, minimo_por_coluna={
        col.C_CATEGORIA: SUGESTAO, col.C_TIPO_DE_OPERACAO: FIRME})
    posicao = col.POSICAO_DA_CATEGORIZACAO[col.C_TIPO_DE_OPERACAO]
    assert doc.categorizacao[posicao] == "Compra Uso Consumo"


# -- a carga de fábrica ----------------------------------------------------

def test_a_fabrica_liga_as_tres_colunas():
    ligadas = parametros.pre_categorizacao(parametros.carregar_fabrica())
    assert ligadas == {col.C_CATEGORIA: "sugestao", col.C_GUARDIAO: "sugestao",
                       col.C_TIPO_DE_OPERACAO: "firme"}


# -- de ponta a ponta ------------------------------------------------------

PENDENTE = chave_de(1)


@pytest.fixture()
def semana(tmp_path) -> dict:
    """Uma nota pendente, sem classificação nenhuma no livro."""
    xml = relatorio(tmp_path / "XML38.xlsx", COLUNAS_XML, [
        linha_xml("1001", PENDENTE, cod_parceiro="4", valor="1500.00"),
    ])
    ce = relatorio(tmp_path / "CE38.xlsx", COLUNAS_CE, [
        linha_ce(PENDENTE, fisica="", fiscal="não", pedido="", farol=""),
    ])
    return {"arquivos": [xml, ce], "saida": tmp_path / "saida",
            "dados": tmp_path / "dados"}


def _semear(raiz, **parceiros):
    bc.gravar(conhecimento(**parceiros), raiz=raiz)


def _rodar(semana, **extras):
    semana["saida"].mkdir(parents=True, exist_ok=True)
    return gerar(semana["arquivos"], semana["saida"],
                 raiz_dos_dados=semana["dados"], responsavel="teste", **extras)


def test_a_planilha_sai_com_a_celula_preenchida_e_realcada(semana):
    _semear(semana["dados"], p4={"guardiao": ("Almoxarifado", 9, 4),
                                 "categoria": ("Indiretos", 9, 4)})
    resultado = _rodar(semana)

    aba = openpyxl.load_workbook(resultado.planilha)["Pendentes"]
    guardiao = col.indice(col.pendentes(resultado.semana), col.C_GUARDIAO) + 1
    categoria = col.indice(col.pendentes(resultado.semana), col.C_CATEGORIA) + 1
    assert aba.cell(2, guardiao).value == "Almoxarifado"
    assert aba.cell(2, categoria).value == "Indiretos"
    assert aba.cell(2, guardiao).fill.fgColor.rgb == "FFFFF2CC"
    assert aba.cell(2, categoria).fill.fgColor.rgb == "FFFFF2CC"
    # A célula que ninguém propôs continua sem realce.
    nota = col.indice(col.pendentes(resultado.semana), col.X_NRO_NOTA) + 1
    assert aba.cell(2, nota).fill.fgColor.rgb != "FFFFF2CC"


def test_a_tela_conta_o_que_a_base_preencheu_e_pede_conferencia(semana):
    _semear(semana["dados"], p4={"guardiao": ("Almoxarifado", 9, 4)})
    resultado = _rodar(semana)

    titulos = [titulo for titulo, _, _ in resultado.listas()]
    assert any("base de conhecimento" in titulo for titulo in titulos)
    tom = next(tom for titulo, _, tom in resultado.listas()
               if "base de conhecimento" in titulo)
    assert tom == "atencao"
    assert resultado.pre_categorizacao.total == 1


def test_sem_base_no_disco_a_planilha_sai_como_sempre_saiu(semana):
    resultado = _rodar(semana)
    aba = openpyxl.load_workbook(resultado.planilha)["Pendentes"]
    guardiao = col.indice(col.pendentes(resultado.semana), col.C_GUARDIAO) + 1
    assert aba.cell(2, guardiao).value in (None, "")
    assert resultado.pre_categorizacao.total == 0


def test_o_guardiao_proposto_encaminha_a_nota_como_o_escrito_a_mao(semana):
    """Roda antes de B2: `Faturamento` proposto manda a nota para a FIS-FAT.

    É a consequência declarada da ordem. A alternativa produziria um arquivo
    que se contradiz — nota na `Pendentes` com guardião de outra população.
    """
    _semear(semana["dados"], p4={"guardiao": ("Faturamento", 9, 4)})
    resultado = _rodar(semana)

    livro = openpyxl.load_workbook(resultado.planilha)
    assert livro["Pendentes"].max_row == 1          # só o cabeçalho
    assert livro["PENDENTES FIS-FAT"].max_row == 2
    assert resultado.fis_fat == 1


def test_a_regra_b1_continua_mandando_no_guardiao(tmp_path):
    """B1 roda antes: a nota conferida e sem incongruência é do Fiscal, e a
    base não tem nada a dizer sobre uma célula que já foi preenchida."""
    chave = chave_de(3)
    xml = relatorio(tmp_path / "XML38.xlsx", COLUNAS_XML,
                    [linha_xml("1003", chave, cod_parceiro="4")])
    ce = relatorio(tmp_path / "CE38.xlsx", COLUNAS_CE,
                   [linha_ce(chave, fisica="Sim", incongruencia="",
                             fiscal="não", pedido="7788", farol="")])
    semana = {"arquivos": [xml, ce], "saida": tmp_path / "saida",
              "dados": tmp_path / "dados"}
    _semear(semana["dados"], p4={"guardiao": ("Almoxarifado", 9, 4)})
    resultado = _rodar(semana)

    assert resultado.reclassificadas_para_fiscal == 1
    aba = openpyxl.load_workbook(resultado.planilha)["PENDENTES FIS-FAT"]
    guardiao = col.indice(col.fis_fat(resultado.semana), col.C_GUARDIAO) + 1
    assert aba.cell(2, guardiao).value == "Fiscal"
    assert aba.cell(2, guardiao).fill.fgColor.rgb != "FFFFF2CC"
