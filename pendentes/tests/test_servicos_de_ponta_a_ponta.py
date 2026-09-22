"""Uma semana inteira de serviços, do arquivo arrastado à planilha gravada.

Os arquivos são sintéticos e montados no próprio teste — CNPJ, número de nota
e nome de prestador inventados, regra nº 1 do `CLAUDE.md`. O que eles provam
não é o número da execução real (isso depende do padrão-ouro, que ainda não
chegou): é que as quatro abas saem com a largura certa, que as identidades
aritméticas fecham, que a classificação atravessa a semana pelo livro e que o
que a ferramenta não consegue explicar aparece na tela em vez de sumir.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from pendentes import estado, parametros, snapshot
from pendentes.servicos import colunas
from pendentes.servicos.execucao import SemRegistros, gerar
from conftest import (
    CNPJ_FILIAL, COLUNAS_ANTERIOR_SERVICOS, COLUNAS_ASIS, COLUNAS_CONFERENCIA,
    COLUNAS_PORTAL, escrever, linha_asis, linha_conferencia, linha_portal,
    relatorio,
)

CNPJ_A = "99888777000166"        # casa pelo procedimento 1
CNPJ_B = "99888777000155"        # casa pelo procedimento 2 (RPS)
CNPJ_C = "99888777000144"        # casa pelo procedimento 3 (CNPJ + valor)
CNPJ_D = "99888777000133"        # casa pelo procedimento 4 (chave fraca)
CNPJ_E = "99888777000122"        # cancelada
CNPJ_F = "99888777000111"        # pendente, com pedido e vínculo
CNPJ_SOLTO = "99888777000177"    # lançamento sem contraparte no ASIS


@pytest.fixture()
def semana(tmp_path) -> dict:
    """Os quatro relatórios da semana, mais a pasta de dados e a de saída."""
    asis = relatorio(tmp_path / "ASIS.xlsx", COLUNAS_ASIS, [
        linha_asis("202600000001234", CNPJ_A, "1500.00"),
        linha_asis("777", CNPJ_B, "320.50", rps="0555"),
        linha_asis("888", CNPJ_C, "999.99"),
        linha_asis("999", CNPJ_D, "111.11"),
        linha_asis("1000", CNPJ_E, "50.00", cancelamento="06/08/2026",
                   motivo_do_cancelamento="erro de emissao"),
        linha_asis("1001", CNPJ_F, "77.00", prestador="OFICINA SEM CADASTRO"),
    ])
    portal = relatorio(tmp_path / "PC27.xlsx", COLUNAS_PORTAL, [
        linha_portal(1, "1234", CNPJ_A, 1500.00),
        linha_portal(2, "555", CNPJ_B, 320.50),
        linha_portal(3, "123", CNPJ_C, 999.99),
        linha_portal(4, "999", "99888777000199", 111.11),
        linha_portal(5, "42", CNPJ_SOLTO, 4321.00, negociacao="12/08/2026"),
        linha_portal(6, "", CNPJ_F, 0, top="1102",
                     descricao_do_top="PC COMPRA DE SERVICO"),
        linha_portal(7, "", "00000000000000", 10.00),
    ])
    conferencia = relatorio(tmp_path / "Conferencia.xlsx", COLUNAS_CONFERENCIA,
                            [linha_conferencia(900, 77.00)])
    return {
        "arquivos": [asis, portal, conferencia],
        "saida": tmp_path / "saida",
        "dados": tmp_path / "dados",
    }


def _rodar(semana, **extras):
    semana["saida"].mkdir(exist_ok=True)
    return gerar(semana["arquivos"], semana["saida"],
                 raiz_dos_dados=semana["dados"], responsavel="teste", **extras)


# -- a planilha -------------------------------------------------------------

def test_a_execucao_gera_as_quatro_abas_com_a_largura_de_cada_uma(semana):
    resultado = _rodar(semana)
    livro = openpyxl.load_workbook(resultado.planilha)

    assert livro.sheetnames == list(colunas.ABAS)
    assert livro["Lancadas"].max_column == 12
    assert livro["Pendentes"].max_column == 36
    assert livro["Canceladas"].max_column == 16
    # 7 de análise + as 17 colunas do relatório do Portal de Compras.
    assert livro["Sem Correspondencia ASIS"].max_column == 7 + len(COLUNAS_PORTAL)


def test_as_abas_saem_com_a_ordem_exata_de_colunas(semana):
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    for nome, layout in (("Lancadas", colunas.LANCADAS),
                         ("Pendentes", colunas.PENDENTES),
                         ("Canceladas", colunas.CANCELADAS)):
        lido = [c.value for c in livro[nome][1]]
        assert lido == [c.rotulo for c in layout], nome


def test_o_arquivo_sai_com_o_nome_que_o_time_ja_reconhece(semana):
    assert _rodar(semana).planilha.name.startswith("Notas_Servico_Pendentes_")


def test_cada_nota_do_asis_vai_para_exatamente_uma_aba(semana):
    resultado = _rodar(semana)
    livro = openpyxl.load_workbook(resultado.planilha)
    escritas = sum(livro[nome].max_row - 1
                   for nome in ("Lancadas", "Pendentes", "Canceladas"))
    assert escritas == resultado.notas == 6


def test_as_identidades_aritmeticas_fecham(semana):
    """As três somas que a execução de referência confirmou, em miniatura."""
    r = _rodar(semana)
    assert r.lancadas + r.pendentes + r.canceladas == r.notas
    assert r.lancadas + r.sem_correspondencia == r.lancamentos
    assert sum(r.por_procedimento.values()) == r.lancadas


def test_os_quatro_procedimentos_sao_exercitados(semana):
    assert _rodar(semana).por_procedimento == {1: 1, 2: 1, 3: 1, 4: 1}


def test_a_cancelada_leva_data_e_motivo_e_nao_entra_no_confronto(semana):
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    aba = livro["Canceladas"]
    assert aba.max_row == 2
    linha = [c.value for c in aba[2]]
    assert linha[0] == "1000"
    assert linha[15] == "erro de emissao"


def test_o_lancamento_sem_contraparte_vai_para_a_aba_inversa(semana):
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    aba = livro["Sem Correspondencia ASIS"]
    assert aba.max_row == 2
    linha = [c.value for c in aba[2]]
    assert linha[0] == "Nao" and linha[2] == "Nao"     # o CNPJ não está no ASIS
    assert linha[7] == "42"                            # Nro. Nota, como texto


def test_a_coluna_de_cidade_e_promovida_para_depois_das_predefinidas(semana):
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    rotulos = [c.value for c in livro["Sem Correspondencia ASIS"][1]]
    assert rotulos[7:12] == ["Nro. Nota", "Parceiro",
                             "Nome Parceiro (Parceiro)", "Vlr. Nota", "Empresa"]
    assert rotulos[12] == "Cidade do Parceiro"


def test_a_pendente_recebe_pedido_comprador_e_vinculo(semana):
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    linha = [c.value for c in livro["Pendentes"][2]]
    assert linha[0] == "1001"
    assert linha[10] == 6                       # Pedido de compra mais recente
    assert linha[11] == "COMPRADOR INVENTADO"
    assert linha[28] == 900                     # Pedido conferencia (NU)
    assert linha[31] == "1x"
    assert linha[33] == "Sim"                   # o semáforo decodificado
    assert linha[35] == "Exato"


def test_a_discriminacao_perde_o_artefato_de_exportacao(semana):
    """`_x000D_` é o retorno de carro escapado do XML, e atravessa como texto."""
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    observacao = [c.value for c in livro["Pendentes"][2]][13]
    assert "_x000D_" not in observacao
    assert observacao == "Servico prestado em agosto"


# -- o que a tela mostra ----------------------------------------------------

def test_o_confronto_por_chave_fraca_bloqueia_o_encerramento(semana):
    resultado = _rodar(semana)
    assert resultado.por_chave_fraca == 1
    assert not resultado.encerravel
    assert any("procedimento 4" in item for item in resultado.bloqueios())


def test_o_lancamento_com_cnpj_zerado_e_contado_em_vez_de_sumir(semana):
    assert any("CNPJ vazio ou zerado" in item
               for item in _rodar(semana).atencoes())


def test_o_corte_do_prefixo_de_ano_e_contado(semana):
    """A suposição de que o Sankhya não tem número de 10+ dígitos fica visível."""
    resultado = _rodar(semana)
    assert resultado.numeros_com_corte_de_ano == 1
    assert any("prefixo de ano" in item for item in resultado.atencoes())


def test_o_confronto_por_procedimento_aparece_na_tela(semana):
    titulos = [titulo for titulo, _, _ in _rodar(semana).listas()]
    assert "Confronto por procedimento" in titulos


def test_a_lista_vermelha_vem_antes_de_todas(semana):
    listas = _rodar(semana).listas()
    assert listas[0][2] == "erro"


def test_a_filial_nao_mapeada_bloqueia_e_nomeia_o_cnpj(tmp_path):
    asis = relatorio(tmp_path / "ASIS.xlsx", COLUNAS_ASIS, [
        linha_asis("1", CNPJ_A, "10.00", tomador="77777777000177"),
    ])
    portal = relatorio(tmp_path / "PC.xlsx", COLUNAS_PORTAL, [
        linha_portal(1, "9", CNPJ_A, 90.00),
    ])
    (tmp_path / "saida").mkdir()
    resultado = gerar([asis, portal], tmp_path / "saida",
                      raiz_dos_dados=tmp_path / "dados")
    assert resultado.filiais_nao_mapeadas == ["77777777000177"]
    assert not resultado.encerravel
    assert any("77777777000177" in item for item in resultado.bloqueios())


def test_o_vinculo_ambiguo_bloqueia_o_encerramento(tmp_path):
    asis = relatorio(tmp_path / "ASIS.xlsx", COLUNAS_ASIS,
                     [linha_asis("1", CNPJ_A, "100.00")])
    portal = relatorio(tmp_path / "PC.xlsx", COLUNAS_PORTAL,
                       [linha_portal(1, "9", CNPJ_A, 90.00)])
    conferencia = relatorio(tmp_path / "Conf.xlsx", COLUNAS_CONFERENCIA, [
        linha_conferencia(900, 100.00),
        linha_conferencia(901, 100.00),
    ])
    (tmp_path / "saida").mkdir()
    resultado = gerar([asis, portal, conferencia], tmp_path / "saida",
                      raiz_dos_dados=tmp_path / "dados")
    assert resultado.vinculos_ambiguos == 1
    assert not resultado.encerravel


def test_a_chave_duplicada_no_sankhya_bloqueia_e_vai_para_a_coluna_obs(tmp_path):
    asis = relatorio(tmp_path / "ASIS.xlsx", COLUNAS_ASIS,
                     [linha_asis("500", CNPJ_A, "100.00")])
    portal = relatorio(tmp_path / "PC.xlsx", COLUNAS_PORTAL, [
        linha_portal(1, "500", CNPJ_A, 100.00),
        linha_portal(2, "500", CNPJ_A, 100.00),
    ])
    (tmp_path / "saida").mkdir()
    resultado = gerar([asis, portal], tmp_path / "saida",
                      raiz_dos_dados=tmp_path / "dados")
    assert resultado.chaves_duplicadas == 1
    assert not resultado.encerravel
    aba = openpyxl.load_workbook(resultado.planilha)["Lancadas"]
    assert [c.value for c in aba[2]][11] == "Chave nota+CNPJ duplicada no Sankhya"


# -- o livro e a semana -----------------------------------------------------

def test_a_classificacao_da_semana_passada_volta_pela_planilha(semana, tmp_path):
    """A ida e volta: sai do livro, volta para o livro, regera a planilha."""
    anterior = escrever(tmp_path / "Pendentes37.xlsx", {"Pendentes": [
        COLUNAS_ANTERIOR_SERVICOS,
        ["1001", "4001", "Suprimentos", "Maria", "aguardando fornecedor"],
    ]})
    semana["arquivos"].append(anterior)
    resultado = _rodar(semana)

    linha = [c.value for c in
             openpyxl.load_workbook(resultado.planilha)["Pendentes"][2]]
    assert linha[4] == "Suprimentos"
    assert linha[5] == "Maria"
    assert linha[6] == "aguardando fornecedor"
    assert resultado.herdadas == 1


def test_a_classificacao_sobrevive_a_execucao_seguinte_sem_o_arquivo(semana,
                                                                     tmp_path):
    """O livro guarda todo mundo: perder o anexo do e-mail deixa de apagar."""
    anterior = escrever(tmp_path / "Pendentes37.xlsx", {"Pendentes": [
        COLUNAS_ANTERIOR_SERVICOS,
        ["1001", "4001", "Suprimentos", "Maria", "aguardando fornecedor"],
    ]})
    semana["arquivos"].append(anterior)
    _rodar(semana)

    semana["arquivos"].remove(anterior)
    segunda = _rodar(semana)
    linha = [c.value for c in
             openpyxl.load_workbook(segunda.planilha)["Pendentes"][2]]
    assert linha[4] == "Suprimentos"
    assert segunda.herdadas == 1


def test_o_livro_e_gravado_fora_do_git_com_carimbo(semana):
    _rodar(semana)
    caminho = estado.caminho_do_livro("servicos", raiz=semana["dados"])
    assert caminho.is_file()
    livro = estado.carregar("servicos", raiz=semana["dados"])
    assert livro.atualizado_por == "teste"
    assert livro.atualizado_em


def test_a_semana_vira_snapshot_com_a_planilha_e_as_impressoes(semana):
    resultado = _rodar(semana)
    pasta = resultado.pasta_do_snapshot
    assert pasta.is_dir()
    assert (pasta / "classificacao.yaml").is_file()
    assert (pasta / "entradas.json").is_file()
    assert (pasta / resultado.planilha.name).is_file()
    ultima = snapshot.ultima_semana("servicos", raiz=semana["dados"])
    assert ultima["encerravel"] is False
    assert ultima["confronto_por_procedimento"]
    # O valor total pendente não vai para a tela, mas vai para a evidência.
    assert ultima["valor_pendente"] == 77.0


def test_rodar_a_mesma_semana_duas_vezes_nao_sobrescreve_a_evidencia(semana):
    primeira = _rodar(semana)
    segunda = _rodar(semana)
    assert primeira.pasta_do_snapshot != segunda.pasta_do_snapshot
    assert segunda.pasta_do_snapshot.name.endswith("-2")


def test_a_semana_deduzida_aparece_no_titulo(semana):
    resultado = _rodar(semana)
    assert f"Semana {resultado.semana}" in resultado.titulo()


def test_a_semana_cadastrada_vence_a_deduzida(semana):
    dados = parametros.carregar_fabrica()
    dados["semana"] = {"numero": "31", "data_de_referencia": "10/08/2026"}
    resultado = _rodar(semana, dados=dados)
    assert (resultado.ano, resultado.semana) == (2026, 31)


# -- o que aborta, e o que só degrada --------------------------------------

def test_sem_o_asis_a_execucao_para_nomeando_o_papel(tmp_path):
    portal = relatorio(tmp_path / "PC.xlsx", COLUNAS_PORTAL,
                       [linha_portal(1, "9", CNPJ_A, 90.00)])
    from pendentes.papeis import PapelAusente

    with pytest.raises(PapelAusente, match="ASIS"):
        gerar([portal], tmp_path, raiz_dos_dados=tmp_path / "dados")


def test_coluna_obrigatoria_ausente_lista_todas_de_uma_vez(tmp_path):
    from pendentes.cabecalho import ColunasFaltando

    colunas_mutiladas = [c for c in COLUNAS_ASIS
                         if c not in ("Valor ISS", "Valor PIS")]
    asis = relatorio(tmp_path / "ASIS.xlsx", colunas_mutiladas,
                     [linha_asis("1", CNPJ_A, "10.00")[:len(colunas_mutiladas)]])
    portal = relatorio(tmp_path / "PC.xlsx", COLUNAS_PORTAL,
                       [linha_portal(1, "9", CNPJ_A, 90.00)])
    with pytest.raises(ColunasFaltando) as erro:
        gerar([asis, portal], tmp_path, raiz_dos_dados=tmp_path / "dados")
    assert "Valor ISS" in str(erro.value) and "Valor PIS" in str(erro.value)


def test_sem_lancamento_de_servico_a_execucao_para_dizendo_o_top(tmp_path):
    asis = relatorio(tmp_path / "ASIS.xlsx", COLUNAS_ASIS,
                     [linha_asis("1", CNPJ_A, "10.00")])
    portal = relatorio(tmp_path / "PC.xlsx", COLUNAS_PORTAL,
                       [linha_portal(1, "9", CNPJ_A, 90.00, top="1102")])
    with pytest.raises(SemRegistros, match="2020"):
        gerar([asis, portal], tmp_path, raiz_dos_dados=tmp_path / "dados")


def test_sem_a_conferencia_a_planilha_sai_com_as_colunas_de_vinculo_vazias(
        semana):
    """Degrada com aviso contado — não aborta, como o VBA já fazia."""
    semana["arquivos"] = semana["arquivos"][:2]
    resultado = _rodar(semana)
    linha = [c.value for c in
             openpyxl.load_workbook(resultado.planilha)["Pendentes"][2]]
    assert linha[28:36] == [None] * 8
    assert any("Conferência de Serviços" in aviso
               for aviso in resultado.atencoes())


def test_o_rodape_do_relatorio_nao_vira_nota(tmp_path):
    """Total e linha em branco no fim do export não são documentos."""
    asis = relatorio(tmp_path / "ASIS.xlsx", COLUNAS_ASIS, [
        linha_asis("1", CNPJ_A, "10.00"),
        ["", "", "", "", "", "", "", "", "", "Total: 10,00"],
    ])
    portal = relatorio(tmp_path / "PC.xlsx", COLUNAS_PORTAL,
                       [linha_portal(1, "1", CNPJ_A, 10.00)])
    (tmp_path / "saida").mkdir()
    resultado = gerar([asis, portal], tmp_path / "saida",
                      raiz_dos_dados=tmp_path / "dados")
    assert resultado.notas == 1


def test_semana_anterior_sem_a_outra_metade_da_chave_degrada_com_aviso(
        semana, tmp_path):
    """Sem `Cod Parceiro` não há identidade: nada é ingerido, e a tela diz."""
    anterior = escrever(tmp_path / "Pendentes37.xlsx", {"Pendentes": [
        ["Nro Nota", "Guardiao", "Gestor de apoio", "Retorno"],
        ["1001", "Suprimentos", "Maria", "aguardando"],
    ]})
    semana["arquivos"].append(anterior)
    resultado = _rodar(semana)

    assert resultado.ingestao is None
    assert any("Cod Parceiro" in aviso for aviso in resultado.atencoes())
    assert len(estado.carregar("servicos", raiz=semana["dados"])) == 0


def test_a_planilha_da_semana_anterior_sem_guardiao_nao_e_reconhecida(
        semana, tmp_path):
    """A âncora do papel é `Nro Nota` **e** `Guardiao`; sem ela, é outro arquivo."""
    estranho = escrever(tmp_path / "Outro.xlsx", {"Pendentes": [
        ["Nro Nota", "Cod Parceiro", "Valor NFSe (Valor Bruto)"],
        ["1001", "4001", 77.00],
    ]})
    semana["arquivos"].append(estranho)
    resultado = _rodar(semana)

    assert resultado.ingestao is None
    assert len(estado.carregar("servicos", raiz=semana["dados"])) == 0


def test_arquivo_que_nao_casa_com_papel_nenhum_aparece_na_tela(semana, tmp_path):
    """Roda com os demais, mas nomeado — nada some em silêncio."""
    estranho = escrever(tmp_path / "Planilha da Ana.xlsx",
                        {"Plan1": [["Coisa", "Outra"], ["a", "b"]]})
    semana["arquivos"].append(estranho)
    resultado = _rodar(semana)
    assert any("Planilha da Ana.xlsx" in aviso for aviso in resultado.atencoes())


# -- a saída da semana passada volta como anexo -----------------------------

def test_a_saida_da_semana_passada_nao_e_tomada_pelo_portal_de_compras(semana):
    """A `Sem Correspondencia ASIS` traz as colunas do Portal de Compras.

    Era o defeito: a planilha gerada na semana anterior, anexada de volta,
    casava com o papel do Portal pela aba inversa e a execução abortava com
    `PapelDuplicado` ao lado do Portal verdadeiro.
    """
    primeira = _rodar(semana)
    livro = openpyxl.load_workbook(primeira.planilha)
    assert livro["Sem Correspondencia ASIS"].max_row > 1

    semana["arquivos"].append(primeira.planilha)
    segunda = _rodar(semana)

    assert segunda.ingestao is not None


def test_a_saida_sem_a_aba_pendentes_fica_de_fora_com_o_motivo(semana, tmp_path):
    """Sem `Pendentes` ela não é a semana anterior — e continua não sendo Portal."""
    from pendentes.servicos.execucao import conferir

    primeira = _rodar(semana)
    livro = openpyxl.load_workbook(primeira.planilha)
    del livro["Pendentes"]
    podada = tmp_path / "Notas_Servico_Pendentes_podada.xlsx"
    livro.save(podada)

    semana["arquivos"].append(podada)
    segunda = _rodar(semana)
    assert segunda.ingestao is None
    assert any(podada.name in aviso for aviso in segunda.atencoes())

    quadro = conferir(semana["arquivos"], dados=parametros.carregar_fabrica())
    ultimo = quadro["arquivos"][-1]
    assert ultimo["papel"] == ""
    assert not ultimo["bloqueia"]
    assert "'Pendentes'" in ultimo["problema"]
