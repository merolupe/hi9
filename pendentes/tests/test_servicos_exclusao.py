"""O que sai do relatório de serviços depois da cascata, e por quê.

Duas regras do time fiscal, de 22/09/2026, e as duas têm a mesma forma: a
nota sai da `Pendentes`, vai para uma aba própria **com o motivo**, e a semana
seguinte sabe que ela saiu. Nada some.

Os arquivos são sintéticos — CNPJ, número de nota e nome de prestador
inventados, regra nº 1.
"""
from __future__ import annotations

import openpyxl
import pytest

from pendentes import parametros
from pendentes.estado import Classificacao
from pendentes.servicos import colunas
from pendentes.servicos import exclusao
from pendentes.servicos.execucao import gerar
from conftest import (COLUNAS_ANTERIOR_SERVICOS, COLUNAS_ASIS, COLUNAS_PORTAL,
                      escrever, linha_asis, linha_portal, relatorio)

CNPJ_PENDENTE = "99888777000111"
CNPJ_OUTRO = "99888777000122"
MARCA = "cancelada"


def classificacao(guardiao="", gestor="", retorno="") -> Classificacao:
    return Classificacao("1001|4001", guardiao=guardiao, gestor_de_apoio=gestor,
                         retornos={37: retorno} if retorno else {})


# -- a nota cancelada na prefeitura ----------------------------------------

def test_os_tres_campos_marcados_declaram_a_nota_cancelada():
    marcada = classificacao(MARCA, MARCA, MARCA)
    assert exclusao.cancelada_na_prefeitura(marcada, MARCA)


def test_um_campo_sozinho_nao_tira_a_nota_do_relatorio():
    """Guardião com `cancelada` e o resto vazio é alguém começando a escrever."""
    for parcial in (classificacao(MARCA), classificacao("", MARCA),
                    classificacao("", "", MARCA),
                    classificacao(MARCA, MARCA)):
        assert not exclusao.cancelada_na_prefeitura(parcial, MARCA)


def test_a_marca_e_comparada_sem_caixa_e_sem_acento():
    escrita_a_mao = classificacao("CANCELADA", " Cancelada ", "cancelada")
    assert exclusao.cancelada_na_prefeitura(escrita_a_mao, MARCA)


def test_o_retorno_que_conta_e_o_da_semana_mais_recente():
    varias = Classificacao("1001|4001", guardiao=MARCA, gestor_de_apoio=MARCA,
                           retornos={35: "aguardando", 38: MARCA})
    assert exclusao.cancelada_na_prefeitura(varias, MARCA)
    voltou_atras = Classificacao("1001|4001", guardiao=MARCA,
                                 gestor_de_apoio=MARCA,
                                 retornos={35: MARCA, 38: "reaberta"})
    assert not exclusao.cancelada_na_prefeitura(voltou_atras, MARCA)


def test_marca_vazia_no_parametro_nao_exclui_nada():
    tudo_vazio = classificacao("", "", "")
    assert not exclusao.cancelada_na_prefeitura(tudo_vazio, "")


# -- a exceção cadastrada ---------------------------------------------------

def test_a_excecao_casa_por_parceiro_e_valor():
    excecao = exclusao.Excecao("7001", 50.0, nome="BETA CONNECT LTDA")
    assert excecao.casa("7001", 50)
    assert excecao.casa(" 7001 ", "50,00")
    assert not excecao.casa("7001", 50.01)
    assert not excecao.casa("7002", 50)


def test_a_excecao_sem_motivo_escrito_explica_a_si_mesma():
    excecao = exclusao.Excecao("7001", 50.0, nome="BETA CONNECT LTDA")
    assert "BETA CONNECT LTDA" in excecao.como_motivo()
    assert "50,00" in excecao.como_motivo()
    com_motivo = exclusao.Excecao("7001", 50.0, motivo="contrato de rateio")
    assert com_motivo.como_motivo() == "contrato de rateio"


def test_linha_de_cadastro_sem_parceiro_ou_sem_valor_e_descartada():
    lidas = exclusao.excecoes_de([
        {"codigo_parceiro": "7001", "valor": 50},
        {"codigo_parceiro": "", "valor": 50},
        {"codigo_parceiro": "7001"},
    ])
    assert len(lidas) == 1


def test_o_cancelamento_vem_antes_da_excecao():
    """Nota que não existe mais não precisa casar com exceção nenhuma."""
    motivo = exclusao.motivo_da_exclusao(
        classificacao(MARCA, MARCA, MARCA), marca=MARCA,
        codigo_do_parceiro="7001", valor=50.0,
        excecoes=[exclusao.Excecao("7001", 50.0)])
    assert motivo == exclusao.MOTIVO_CANCELADA


def test_cadastro_vazio_e_nota_limpa_nao_produzem_motivo():
    assert not exclusao.motivo_da_exclusao(
        classificacao("Suprimentos", "Ana", "aguardando"), marca=MARCA,
        codigo_do_parceiro="7001", valor=50.0, excecoes=())


def test_a_fabrica_nasce_sem_excecao_nenhuma():
    """É cadastro de dado da empresa: numa máquina nova não exclui nada."""
    assert parametros.excecoes_de_servicos(parametros.carregar_fabrica()) == ()


def test_a_marca_de_cancelada_e_parametro_e_tem_padrao():
    ajuste = parametros.confronto_de_servicos(parametros.carregar_fabrica())
    assert ajuste["marca_de_cancelada"] == "cancelada"


# -- de ponta a ponta -------------------------------------------------------

@pytest.fixture()
def semana(tmp_path) -> dict:
    """Duas notas pendentes: uma para marcar, outra para a exceção."""
    asis = relatorio(tmp_path / "ASIS.xlsx", COLUNAS_ASIS, [
        linha_asis("1001", CNPJ_PENDENTE, "77.00", prestador="OFICINA ALFA"),
        linha_asis("1002", CNPJ_OUTRO, "50.00", prestador="BETA CONNECT"),
    ])
    portal = relatorio(tmp_path / "PC38.xlsx", COLUNAS_PORTAL, [
        # Um lançamento de serviço de verdade: sem nenhum, a execução aborta
        # antes de chegar à cascata, e não é disso que estes testes tratam.
        linha_portal(1, "9999", "99888777000199", 10.00),
        linha_portal(6, "", CNPJ_PENDENTE, 0, top="1102",
                     descricao_do_top="PC COMPRA DE SERVICO"),
        linha_portal(7, "", CNPJ_OUTRO, 0, top="1102",
                     descricao_do_top="PC COMPRA DE SERVICO"),
    ])
    return {"arquivos": [asis, portal], "saida": tmp_path / "saida",
            "dados": tmp_path / "dados", "tmp": tmp_path}


def _rodar(semana, **extras):
    semana["saida"].mkdir(exist_ok=True)
    return gerar(semana["arquivos"], semana["saida"],
                 raiz_dos_dados=semana["dados"], responsavel="teste", **extras)


def _codigo_do_parceiro(semana, numero: str) -> str:
    """O código que o Portal deu ao parceiro daquela nota, como a execução o vê."""
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    for linha in livro["Pendentes"].iter_rows(min_row=2, values_only=True):
        if str(linha[0]) == numero:
            return str(linha[2])
    raise AssertionError(f"nota {numero} não saiu em Pendentes")


def test_a_aba_nova_sai_com_as_36_da_pendentes_mais_o_motivo(semana):
    livro = openpyxl.load_workbook(_rodar(semana).planilha)
    aba = livro[colunas.ABA_FORA_DO_RELATORIO]
    assert aba.max_column == 37
    assert [c.value for c in aba[1]] == [c.rotulo
                                         for c in colunas.FORA_DO_RELATORIO]
    assert [c.value for c in aba[1]][-1] == "Motivo"


def test_a_nota_marcada_na_semana_passada_sai_da_pendentes(semana):
    codigo = _codigo_do_parceiro(semana, "1001")
    anterior = escrever(semana["tmp"] / "Pendentes37.xlsx", {"Pendentes": [
        COLUNAS_ANTERIOR_SERVICOS,
        ["1001", codigo, "cancelada", "cancelada", "cancelada"],
    ]})
    semana["arquivos"].append(anterior)
    resultado = _rodar(semana)

    livro = openpyxl.load_workbook(resultado.planilha)
    pendentes = [str(l[0]) for l in
                 livro["Pendentes"].iter_rows(min_row=2, values_only=True)]
    fora = list(livro[colunas.ABA_FORA_DO_RELATORIO].iter_rows(
        min_row=2, values_only=True))
    assert "1001" not in pendentes
    assert [str(l[0]) for l in fora] == ["1001"]
    assert fora[0][-1] == exclusao.MOTIVO_CANCELADA
    assert resultado.fora_do_relatorio == {exclusao.MOTIVO_CANCELADA: 1}


def test_a_nota_marcada_so_no_guardiao_continua_pendente(semana):
    codigo = _codigo_do_parceiro(semana, "1001")
    anterior = escrever(semana["tmp"] / "Pendentes37.xlsx", {"Pendentes": [
        COLUNAS_ANTERIOR_SERVICOS,
        ["1001", codigo, "cancelada", "", ""],
    ]})
    semana["arquivos"].append(anterior)
    resultado = _rodar(semana)
    assert resultado.fora_do_relatorio == {}
    livro = openpyxl.load_workbook(resultado.planilha)
    assert livro[colunas.ABA_FORA_DO_RELATORIO].max_row == 1


def test_a_aba_de_exclusoes_da_semana_passada_e_lida_de_volta(semana):
    """Perder o livro não faz a nota cancelada voltar: a aba também ensina."""
    codigo = _codigo_do_parceiro(semana, "1001")
    anterior = escrever(semana["tmp"] / "Pendentes37.xlsx", {
        "Pendentes": [COLUNAS_ANTERIOR_SERVICOS],
        colunas.ABA_FORA_DO_RELATORIO: [
            COLUNAS_ANTERIOR_SERVICOS,
            ["1001", codigo, "cancelada", "cancelada", "cancelada"],
        ],
    })
    semana["arquivos"].append(anterior)
    resultado = _rodar(semana)
    assert resultado.fora_do_relatorio == {exclusao.MOTIVO_CANCELADA: 1}


def test_a_excecao_cadastrada_tira_a_nota_daquele_parceiro_e_valor(semana):
    codigo = _codigo_do_parceiro(semana, "1002")
    dados = parametros.carregar_fabrica()
    dados["excecoes_servicos"] = [
        {"codigo_parceiro": codigo, "nome": "BETA CONNECT", "valor": 50},
    ]
    resultado = _rodar(semana, dados=dados)

    livro = openpyxl.load_workbook(resultado.planilha)
    pendentes = [str(l[0]) for l in
                 livro["Pendentes"].iter_rows(min_row=2, values_only=True)]
    fora = list(livro[colunas.ABA_FORA_DO_RELATORIO].iter_rows(
        min_row=2, values_only=True))
    assert pendentes == ["1001"]
    assert [str(l[0]) for l in fora] == ["1002"]
    assert "BETA CONNECT" in fora[0][-1]


def test_a_excecao_de_outro_valor_nao_tira_a_nota(semana):
    codigo = _codigo_do_parceiro(semana, "1002")
    dados = parametros.carregar_fabrica()
    dados["excecoes_servicos"] = [{"codigo_parceiro": codigo, "valor": 51}]
    assert _rodar(semana, dados=dados).fora_do_relatorio == {}


def test_o_que_saiu_e_contado_na_tela_com_o_motivo(semana):
    codigo = _codigo_do_parceiro(semana, "1002")
    dados = parametros.carregar_fabrica()
    dados["excecoes_servicos"] = [
        {"codigo_parceiro": codigo, "valor": 50, "motivo": "contrato de rateio"},
    ]
    resultado = _rodar(semana, dados=dados)
    fichas = dict(resultado.fichas())
    assert fichas["Fora do relatório"] == "1"
    titulos = {titulo for titulo, _, _ in resultado.listas()}
    assert "Fora do relatório, com o motivo" in titulos
    itens = next(itens for titulo, itens, _ in resultado.listas()
                 if titulo == "Fora do relatório, com o motivo")
    assert itens == ["1 — contrato de rateio"]
